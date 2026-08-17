#!/usr/bin/env python3
"""假设登记簿 Excel 导出/回灌工具（主会话调用，子 agent 不得导入）。

子命令：
    export <id> [--xlsx path]
        workspace/<id>/assumptions.md → assumptions.xlsx（幂等覆盖写）。
    apply <id> [--xlsx path] [--dry-run] [--each]
        读回客户改过的 xlsx → 校验分流：参数型（有「参数绑定」且填新值）由本工具
        直接改 src/<id>/config.py 字面值（备份 + diff + 实例化冒烟）并对
        assumptions.md 做确定性追记；文本型（填新口径）只输出 text_revisions
        JSON，交主会话派 quant-planner 落实。--each 模式只解析校验输出计划，
        不落盘（逐条敏感性由主会话用 set-param/restore 驱动）。
    set-param <id> --field <f> --value <v>
        --each 脚手架：只改 config.py 单字段（备份 + 冒烟），不动 md。
    restore <id> [--backup path]
        回滚 config.py 到备份（默认 workspace/<id>/audit/config_backups/ 最新）。
    sensitivity <id> --baseline <comparison.json> --runs-dir <dir>
        聚合 --each 各轮快照（AS<n>_comparison.json + AS<n>_meta.json）→
        output/<id>/results/assumption_sensitivity.xlsx（明细长表 + 汇总）。

约定：
- 「参数绑定」是 assumptions.md 条目的可选字段（`- 参数绑定: cost_bps`），
  仅 config.py 类体单行字面量字段可绑（field(default_factory=...) / 跨行 /
  运行期开关不可绑）。
- 参数型 md 追记：状态→revised、验证后回看→重置占位符 [verify 后填]、
  插入「修订记录」行——与 state.py 记账同属确定性工具改写，不违反编排/生产分离。
- 所有子命令 stdout 只打印 JSON（日志走 stderr），exit 0=成功/至少一条生效。

路径解析：REPORT_REPRODUCE_ROOT 环境变量覆盖，否则仓库根推导（同 state.default_root）。
"""

from __future__ import annotations

import argparse
import ast
import difflib
import json
import re
import shutil
import subprocess
import sys
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Any, Optional

_REPO_ROOT_FOR_IMPORT = Path(__file__).resolve().parents[1]
if str(_REPO_ROOT_FOR_IMPORT) not in sys.path:
    sys.path.insert(0, str(_REPO_ROOT_FOR_IMPORT))

from openpyxl import Workbook, load_workbook  # noqa: E402
from openpyxl.styles import Alignment, Font, PatternFill, Protection  # noqa: E402
from openpyxl.worksheet.datavalidation import DataValidation  # noqa: E402
from openpyxl.utils import get_column_letter  # noqa: E402

from tools import state as st  # noqa: E402
from tools.check_gates import parse_id_blocks  # noqa: E402

# ---------------------------------------------------------------------------
# 常量
# ---------------------------------------------------------------------------

HEADERS = [
    "AS ID", "标题", "来源", "影响面", "影响 milestone", "影响指标",
    "当前口径（假设内容）", "验证后回看", "状态", "高亮等级", "性质",
    "参数绑定 config 字段", "当前值", "新值", "新口径（文本说明）", "操作",
]
COL_AS_ID, COL_NEW_VALUE, COL_NEW_TEXT, COL_OP = "AS ID", "新值", "新口径（文本说明）", "操作"
CORE_COLUMNS = [COL_AS_ID, COL_NEW_VALUE, COL_NEW_TEXT, COL_OP]

EDITABLE_FILL = PatternFill("solid", fgColor="FFF2CC")
HEADER_FONT = Font(bold=True)
WIDE_TEXT_COLUMNS = {"当前口径（假设内容）", "验证后回看", "新口径（文本说明）"}

# 影响面 → 重跑范围排序（data/method 同级最重，output 最轻），取 max 即合并范围
IMPACT_ORDER = {"data": 4, "method": 3, "param": 2, "trading": 1, "output": 0}

UNPARSEABLE = object()  # ConfigField.value 哨兵：不可字面量求值 / 跨行，不可绑定

_HEADING_LINE_RE = re.compile(r"^### \[([A-Z]+)(\d+(?:\.\d+)?)\]", re.MULTILINE)
_MILESTONE_RE = re.compile(r"m\d+")


# ---------------------------------------------------------------------------
# config.py 静态解析
# ---------------------------------------------------------------------------


@dataclass
class ConfigField:
    name: str
    annotation: str
    value: Any           # UNPARSEABLE 哨兵或 ast.literal_eval 结果
    value_repr: str      # 展示用 repr
    lineno: int
    col_offset: int
    end_lineno: int
    end_col_offset: int

    @property
    def bindable(self) -> bool:
        return self.value is not UNPARSEABLE


def _find_config_class(tree: ast.Module) -> Optional[ast.ClassDef]:
    """定位配置 dataclass：优先模块级 DEFAULT_CONFIG = Xxx() 指向的类，否则第一个 @dataclass 类。"""
    default_cls_name: Optional[str] = None
    for node in tree.body:
        if isinstance(node, ast.Assign) and len(node.targets) == 1:
            tgt = node.targets[0]
            if isinstance(tgt, ast.Name) and tgt.id == "DEFAULT_CONFIG":
                if isinstance(node.value, ast.Call) and isinstance(node.value.func, ast.Name):
                    default_cls_name = node.value.func.id

    def is_dataclass(cls: ast.ClassDef) -> bool:
        for dec in cls.decorator_list:
            func = dec.func if isinstance(dec, ast.Call) else dec
            if isinstance(func, ast.Name) and func.id == "dataclass":
                return True
        return False

    classes = [n for n in tree.body if isinstance(n, ast.ClassDef) and is_dataclass(n)]
    if default_cls_name:
        for cls in classes:
            if cls.name == default_cls_name:
                return cls
    return classes[0] if classes else None


def parse_config_fields(config_path: Path) -> dict[str, ConfigField]:
    """静态解析 config.py 的 dataclass 字段。单行字面量可绑，其余标 UNPARSEABLE。"""
    if not config_path.exists():
        return {}
    try:
        tree = ast.parse(config_path.read_text(encoding="utf-8"))
    except SyntaxError:
        return {}
    cls = _find_config_class(tree)
    if cls is None:
        return {}
    fields: dict[str, ConfigField] = {}
    for node in cls.body:
        if not isinstance(node, ast.AnnAssign) or not isinstance(node.target, ast.Name):
            continue
        name = node.target.id
        annotation = ast.unparse(node.annotation) if node.annotation is not None else ""
        value: Any = UNPARSEABLE
        value_repr = "(不可绑定：无默认值)"
        lineno = col = end_lineno = end_col = 0
        if node.value is not None:
            lineno, col = node.value.lineno, node.value.col_offset
            end_lineno = node.value.end_lineno or lineno
            end_col = node.value.end_col_offset or col
            if end_lineno == lineno:
                try:
                    value = ast.literal_eval(node.value)
                    value_repr = repr(value)
                except (ValueError, SyntaxError):
                    value_repr = "(不可绑定：非字面量/default_factory)"
            else:
                value_repr = "(不可绑定：跨行字面量)"
        fields[name] = ConfigField(name, annotation, value, value_repr, lineno, col, end_lineno, end_col)
    return fields


# ---------------------------------------------------------------------------
# 新值类型校验（以当前字面量的 Python 类型为准）
# ---------------------------------------------------------------------------


def coerce_new_value(raw: Any, current: Any) -> tuple[bool, Any, str]:
    """把 Excel 单元格/CLI 原始值归一到当前字段类型。返回 (ok, value, reason)。"""
    if isinstance(current, bool):
        if isinstance(raw, bool):
            return True, raw, ""
        if isinstance(raw, str) and raw.strip().upper() in ("TRUE", "FALSE"):
            return True, raw.strip().upper() == "TRUE", ""
        return False, None, f"布尔字段应填 TRUE/FALSE，收到 {raw!r}"
    if isinstance(current, int):
        if isinstance(raw, bool):
            return False, None, "整数字段不接受布尔值"
        if isinstance(raw, int):
            return True, raw, ""
        if isinstance(raw, float) and raw.is_integer():
            return True, int(raw), ""
        if isinstance(raw, str):
            try:
                v = ast.literal_eval(raw.strip())
            except (ValueError, SyntaxError):
                return False, None, f"无法解析为整数：{raw!r}"
            if isinstance(v, int) and not isinstance(v, bool):
                return True, v, ""
        return False, None, f"整数字段收到不兼容值 {raw!r}"
    if isinstance(current, float):
        if isinstance(raw, bool):
            return False, None, "浮点字段不接受布尔值"
        if isinstance(raw, (int, float)):
            return True, float(raw), ""
        if isinstance(raw, str):
            try:
                v = ast.literal_eval(raw.strip())
            except (ValueError, SyntaxError):
                return False, None, f"无法解析为数值：{raw!r}"
            if isinstance(v, (int, float)) and not isinstance(v, bool):
                return True, float(v), ""
        return False, None, f"浮点字段收到不兼容值 {raw!r}"
    if isinstance(current, str):
        if isinstance(raw, str):
            stripped = raw.strip()
            try:
                v = ast.literal_eval(stripped)
                if isinstance(v, str):
                    return True, v, ""
            except (ValueError, SyntaxError):
                pass
            return True, stripped, ""  # 裸文本按字符串本身接受
        return False, None, f"字符串字段收到非文本值 {raw!r}"
    if isinstance(current, tuple):
        v = raw
        if isinstance(raw, str):
            try:
                v = ast.literal_eval(raw.strip())
            except (ValueError, SyntaxError):
                return False, None, f"无法解析为元组（写成 (20,30) 形式）：{raw!r}"
        if isinstance(v, (tuple, list)):
            return True, tuple(v), ""
        return False, None, f"元组字段收到不兼容值 {raw!r}"
    return False, None, f"字段类型 {type(current).__name__} 不在可注入白名单（int/float/str/bool/tuple）"


# ---------------------------------------------------------------------------
# config.py 安全改写 / 冒烟 / 备份 / 回滚
# ---------------------------------------------------------------------------


def apply_literal_to_text(text: str, field: ConfigField, new_value: Any) -> str:
    """行级精确替换值节点（单行），尾随注释天然保留。返回新全文。"""
    lines = text.splitlines(keepends=True)
    idx = field.lineno - 1
    line = lines[idx]
    lines[idx] = line[: field.col_offset] + repr(new_value) + line[field.end_col_offset :]
    new_text = "".join(lines)
    ast.parse(new_text)  # 语法校验，失败抛 SyntaxError
    return new_text


def instantiate_smoke(config_path: Path) -> tuple[bool, str]:
    """子进程 import config 并取 DEFAULT_CONFIG（触发 __post_init__ 自洽校验）。"""
    code = (
        "import importlib.util;"
        f"spec=importlib.util.spec_from_file_location('_cfg_smoke', r'''{config_path}''');"
        "m=importlib.util.module_from_spec(spec);spec.loader.exec_module(m);"
        "assert getattr(m, 'DEFAULT_CONFIG', None) is not None, 'DEFAULT_CONFIG 缺失';"
        "print('SMOKE_OK')"
    )
    try:
        proc = subprocess.run(
            [sys.executable, "-c", code], capture_output=True, text=True, timeout=120
        )
    except subprocess.TimeoutExpired:
        return False, "实例化冒烟超时（120s）"
    if proc.returncode != 0 or "SMOKE_OK" not in proc.stdout:
        return False, (proc.stderr or proc.stdout).strip()[-500:]
    return True, ""


def backup_dir(root: Path, report_id: str) -> Path:
    d = st.workspace_dir(root, report_id) / "audit" / "config_backups"
    d.mkdir(parents=True, exist_ok=True)
    return d


def backup_config(root: Path, report_id: str, config_path: Path, tag: str) -> Path:
    dst = backup_dir(root, report_id) / f"config_{tag}.py"
    shutil.copy2(config_path, dst)
    return dst


def write_diff_patch(root: Path, report_id: str, tag: str, before: str, after: str, config_path: Path) -> Optional[Path]:
    diff = "".join(
        difflib.unified_diff(
            before.splitlines(keepends=True),
            after.splitlines(keepends=True),
            fromfile=f"a/{config_path.name}",
            tofile=f"b/{config_path.name}",
        )
    )
    if not diff.strip():
        return None
    dst = backup_dir(root, report_id) / f"diff_{tag}.patch"
    dst.write_text(diff, encoding="utf-8")
    return dst


def restore_config(root: Path, report_id: str, backup: Optional[Path] = None) -> Path:
    """回滚 config.py：--backup 指定，否则取 config_backups 下最新 config_*.py。"""
    bdir = backup_dir(root, report_id)
    if backup is None:
        candidates = sorted(bdir.glob("config_*.py"), key=lambda p: p.stat().st_mtime)
        if not candidates:
            raise FileNotFoundError(f"无可用备份：{bdir}")
        backup = candidates[-1]
    dst = st.src_dir(root, report_id) / "config.py"
    shutil.copy2(backup, dst)
    return backup


# ---------------------------------------------------------------------------
# assumptions.md 读取 / 统计 / 确定性追记
# ---------------------------------------------------------------------------


def load_as_blocks(md_path: Path) -> list:
    if not md_path.exists():
        raise FileNotFoundError(f"假设登记簿不存在：{md_path}")
    blocks = [b for b in parse_id_blocks(md_path.read_text(encoding="utf-8")) if b.prefix == "AS"]
    if not blocks:
        raise ValueError(f"假设登记簿无 AS 条目：{md_path}")
    return blocks


def assumptions_stats(blocks: list) -> dict[str, int]:
    stats = {"total": len(blocks), "assumed": 0, "confirmed": 0, "revised": 0}
    for b in blocks:
        status = b.fields.get("状态", "assumed")
        stats[status if status in ("assumed", "confirmed", "revised") else "assumed"] += 1
    return stats


def _block_line_span(lines: list[str], as_id: str) -> Optional[tuple[int, int]]:
    """定位 `### [AS<n>]` 块的行区间 [start, end)（end 为下一标题行或全文尾）。"""
    start: Optional[int] = None
    for i, line in enumerate(lines):
        m = _HEADING_LINE_RE.match(line)
        if m:
            cur = f"{m.group(1)}{m.group(2)}"
            if cur == as_id:
                start = i
            elif start is not None:
                return start, i
    if start is None:
        return None
    return start, len(lines)


def apply_param_md_edits(
    md_path: Path, edits: list[dict], batch_id: str, ts: str
) -> tuple[list[dict], str]:
    """参数型条目的确定性三处行级编辑。返回 (md_edits_failed, 改写后全文)。

    edits 每项：{as_id, field, old_repr, new_repr}
    """
    text = md_path.read_text(encoding="utf-8")
    lines = text.splitlines()
    failed: list[dict] = []
    for edit in edits:
        span = _block_line_span(lines, edit["as_id"])
        if span is None:
            failed.append({**edit, "reason": "条目标题定位失败"})
            continue
        start, end = span
        status_idx = next(
            (i for i in range(start, end) if re.match(r"^-\s*状态\s*[:：]", lines[i])), None
        )
        review_idx = next(
            (i for i in range(start, end) if re.match(r"^-\s*验证后回看\s*[:：]", lines[i])), None
        )
        if status_idx is None or review_idx is None:
            failed.append({**edit, "reason": "条目缺「状态」或「验证后回看」行"})
            continue
        record = (
            f"- 修订记录: {ts} 批量 Excel revise（批次 {batch_id}）："
            f"{edit['field']} {edit['old_repr']} → {edit['new_repr']}；"
            "工具直接注入 config.py，未经过 agent"
        )
        lines[status_idx] = "- 状态: revised"
        lines[review_idx] = "- 验证后回看: [verify 后填]"
        lines.insert(review_idx, record)  # 状态/回看行号不变（insert 在 review_idx 处，其后行顺延但不复用）
    return failed, "\n".join(lines) + ("\n" if text.endswith("\n") else "")


def _atomic_write_text(path: Path, text: str) -> None:
    tmp = path.with_suffix(path.suffix + ".tmp")
    tmp.write_text(text, encoding="utf-8")
    tmp.replace(path)


# ---------------------------------------------------------------------------
# Excel 导出
# ---------------------------------------------------------------------------


def _auto_width(ws, max_col: int, min_w: int = 8, cap: int = 40) -> None:
    for col_idx in range(1, max_col + 1):
        letter = get_column_letter(col_idx)
        header = ws.cell(row=1, column=col_idx).value or ""
        if header in WIDE_TEXT_COLUMNS:
            ws.column_dimensions[letter].width = 60
            continue
        longest = len(str(header))
        for row in ws.iter_rows(min_row=2, max_row=min(ws.max_row, 200), min_col=col_idx, max_col=col_idx):
            cell = row[0]
            if cell.value is not None:
                longest = max(longest, min(len(str(cell.value)), 80))
        ws.column_dimensions[letter].width = min(max(longest + 2, min_w), cap)


def _write_instructions(ws, report_id: str) -> None:
    lines = [
        "假设修改工作表 · 填写说明",
        "",
        "一、两种修改通道",
        "1. 参数直改（推荐，无需等待 agent）：「参数绑定 config 字段」列非空的行，在「新值」列直接填新值，",
        "   提交后由工具自动注入 src/<id>/config.py 并重跑。类型须与「当前值」一致：",
        "   整数填 12、浮点填 5.0、字符串填 month_end（可不带引号）、布尔填 TRUE/FALSE、元组填 (20,30)。",
        "2. 口径文本（走完整审计链）：无参数绑定的行、或想修改口径叙述时，在「新口径（文本说明）」列填文字描述，",
        "   提交后由 quant-planner 改写登记簿、quant-coder 落实代码，走 code_audit/verify 审计链。",
        "3. 「新值」「新口径」可同填：新值立即注入，新口径文字交 planner 同步改写条目叙述。",
        "",
        "二、生效条件",
        "- 「操作」列必须选 revise 才生效；填了新值/新口径但没选 revise 的行会被拒绝并列出原因。",
        "- 新值与当前值相同的行会被拒绝（无变化不产生修订）。",
        "",
        "三、不可参数直改的情形（请改用「新口径」文本通道）",
        "- 「当前值」列显示「不可绑定」的字段：field(default_factory=...)（如日历周期）、跨行字面量；",
        "- 运行期开关类假设（代码内开关而非 config 字段，如消融开关）；",
        "- 无「参数绑定」的口径/方法/结构类假设。",
        "",
        "四、提交命令",
        f"- 合并重跑（默认，所有修改一次跑完）：/reproduce revise {report_id} --from-excel",
        f"- 逐条敏感性（每条独立重跑，产对比表，成本=修改条数次全量回测）：/reproduce revise {report_id} --from-excel --each",
        f"- 重新导出干净工作表：/reproduce export-assumptions {report_id}",
        "",
        "五、存档语义",
        "- 登记簿只增不删：被修改的假设状态改为 revised，原口径保留在条目内并追加「修订记录」行；",
        "- 修改后「验证后回看」会重置为待回填，由重跑的 verify 阶段用实测数据重新验证。",
    ]
    for i, line in enumerate(lines, start=1):
        ws.cell(row=i, column=1, value=line)
    ws.column_dimensions["A"].width = 110


def cmd_export(args: argparse.Namespace) -> int:
    root = st.default_root()
    md_path = st.workspace_dir(root, args.report_id) / "assumptions.md"
    try:
        blocks = load_as_blocks(md_path)
    except (FileNotFoundError, ValueError) as e:
        print(str(e), file=sys.stderr)
        return 1
    config_path = st.src_dir(root, args.report_id) / "config.py"
    fields = parse_config_fields(config_path)

    xlsx_path = Path(args.xlsx) if args.xlsx else st.workspace_dir(root, args.report_id) / "assumptions.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "假设登记簿"
    ws.append(HEADERS)
    for cell in ws[1]:
        cell.font = HEADER_FONT

    bound: list[dict] = []
    unbound_param_candidates: list[dict] = []
    for b in blocks:
        f = b.fields
        bind = f.get("参数绑定", "").strip()
        if bind:
            cfg = fields.get(bind)
            cur_repr = cfg.value_repr if cfg else "(字段不存在)"
            bound.append({"as_id": b.id, "field": bind})
        else:
            cur_repr = ""
            if f.get("影响面", "") in ("param", "trading"):
                unbound_param_candidates.append({"as_id": b.id, "impact": f.get("影响面", ""), "title": b.title})
        ws.append([
            b.id, b.title, f.get("来源", ""), f.get("影响面", ""), f.get("影响 milestone", ""),
            f.get("影响指标", ""), f.get("假设内容", ""), f.get("验证后回看", ""), f.get("状态", ""),
            f.get("高亮等级", ""), f.get("性质", ""), bind, cur_repr, None, None, None,
        ])

    editable_cols = [HEADERS.index(COL_NEW_VALUE) + 1, HEADERS.index(COL_NEW_TEXT) + 1, HEADERS.index(COL_OP) + 1]
    wrap_cols = [i + 1 for i, h in enumerate(HEADERS) if h in WIDE_TEXT_COLUMNS]
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
        for cell in row:
            if cell.column in editable_cols:
                cell.fill = EDITABLE_FILL
                cell.protection = Protection(locked=False)
            if cell.column in wrap_cols:
                cell.alignment = Alignment(wrap_text=True, vertical="top")

    dv = DataValidation(type="list", formula1='"revise"', allow_blank=True, showDropDown=False)
    dv.error = "请选择 revise（或留空跳过本行）"
    dv.errorTitle = "操作列取值"
    op_letter = get_column_letter(HEADERS.index(COL_OP) + 1)
    dv.add(f"{op_letter}2:{op_letter}10000")
    ws.add_data_validation(dv)

    ws.freeze_panes = "A2"
    ws.protection.sheet = True
    ws.protection.formatColumns = False
    _auto_width(ws, len(HEADERS))

    ws2 = wb.create_sheet("填写说明")
    _write_instructions(ws2, args.report_id)

    xlsx_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(xlsx_path)
    print(json.dumps({
        "xlsx": str(xlsx_path.relative_to(root)),
        "total": len(blocks),
        "bound": bound,
        "unbound_param_candidates": unbound_param_candidates,
        "assumptions_stats": assumptions_stats(blocks),
    }, ensure_ascii=False, indent=2))
    return 0


# ---------------------------------------------------------------------------
# Excel 读回与行校验
# ---------------------------------------------------------------------------


def _normalize_header(text: str) -> str:
    return re.sub(r"[\s（(].*$", "", str(text)).strip().lower()


def _map_headers(ws) -> dict[str, int]:
    """表头 → 列号（1 基）。先精确匹配，缺失列用归一化前缀模糊匹配。"""
    raw: dict[int, str] = {}
    for cell in ws[1]:
        if cell.value is not None:
            raw[cell.column] = str(cell.value).strip()
    mapping: dict[str, int] = {}
    for h in HEADERS:
        for col, val in raw.items():
            if val == h:
                mapping[h] = col
                break
    for h in HEADERS:
        if h in mapping:
            continue
        norm = _normalize_header(h)
        for col, val in raw.items():
            if _normalize_header(val).startswith(norm[:4]) and norm[:4]:
                mapping[h] = col
                break
    return mapping


def _cell_str(value: Any) -> str:
    return "" if value is None else str(value).strip()


def read_row_edits(xlsx_path: Path, blocks: list, fields: dict[str, ConfigField]) -> tuple[list[dict], list[dict], Optional[str]]:
    """读取 xlsx，返回 (row_edits, hard_rejected, hard_error)。

    row_edits 每项：{as_id, new_value_raw, has_new_value, new_text, impact, milestones, bind}
    hard_rejected 每项：{row, as_id, reason}
    """
    try:
        wb = load_workbook(xlsx_path, data_only=True)
    except Exception as e:  # openpyxl 各类解析异常
        return [], [], f"无法读取 xlsx（{e}）；请用 export 重新导出干净工作表"
    main = next((wb[n] for n in wb.sheetnames if "假设" in n), wb[wb.sheetnames[0]])
    mapping = _map_headers(main)
    missing = [c for c in CORE_COLUMNS if c not in mapping]
    if missing:
        return [], [], f"工作表缺少核心列 {missing}（表头被改名？）；请用 export 重新导出干净工作表"

    by_id = {b.id: b for b in blocks}
    edits: list[dict] = []
    rejected: list[dict] = []
    for row in main.iter_rows(min_row=2):
        def get(col: str) -> Any:
            idx = mapping.get(col)
            return row[idx - 1].value if idx else None

        as_id = _cell_str(get(COL_AS_ID)).upper()
        new_value_raw = get(COL_NEW_VALUE)
        new_text = _cell_str(get(COL_NEW_TEXT))
        op = _cell_str(get(COL_OP)).lower()
        has_new_value = new_value_raw is not None and _cell_str(new_value_raw) != ""
        if not as_id and not has_new_value and not new_text and not op:
            continue  # 全空行
        row_no = row[0].row
        if not as_id or as_id not in by_id:
            rejected.append({"row": row_no, "as_id": as_id or "(空)", "reason": "AS ID 不在登记簿中"})
            continue
        block = by_id[as_id]
        bind = block.fields.get("参数绑定", "").strip()
        if op != "revise":
            if op:
                rejected.append({"row": row_no, "as_id": as_id, "reason": f"操作列取值无效：{op!r}（应为 revise 或留空）"})
            elif has_new_value or new_text:
                rejected.append({"row": row_no, "as_id": as_id, "reason": "已填新值/新口径但操作列未选 revise"})
            continue
        if not has_new_value and not new_text:
            rejected.append({"row": row_no, "as_id": as_id, "reason": "已选 revise 但新值与新口径均为空"})
            continue
        if has_new_value and not bind:
            rejected.append({"row": row_no, "as_id": as_id, "reason": "本条无参数绑定，新值无效；请改填新口径列走文本通道"})
            has_new_value = False  # 新口径若同填仍生效
            if not new_text:
                continue
        edit = {
            "as_id": as_id,
            "title": block.title,
            "new_value_raw": new_value_raw,
            "has_new_value": has_new_value,
            "new_text": new_text,
            "impact": block.fields.get("影响面", ""),
            "milestones": _MILESTONE_RE.findall(block.fields.get("影响 milestone", "")),
            "bind": bind,
        }
        if has_new_value:
            cfg = fields.get(bind)
            fail_reason: Optional[str] = None
            if cfg is None:
                fail_reason = f"绑定字段 {bind} 在 config.py 中不存在"
            elif not cfg.bindable:
                fail_reason = f"绑定字段 {bind} 不可注入（{cfg.value_repr}）；请改走新口径文本通道"
            else:
                ok, value, reason = coerce_new_value(new_value_raw, cfg.value)
                if not ok:
                    fail_reason = f"字段 {bind}：{reason}"
                elif value == cfg.value:
                    fail_reason = f"字段 {bind} 新值与当前值相同（{cfg.value_repr}），无变化"
            if fail_reason is not None:
                rejected.append({"row": row_no, "as_id": as_id, "reason": fail_reason})
                edit["has_new_value"] = False  # 新口径若同填，降级为纯文本行仍生效
                if new_text:
                    edits.append(edit)
                continue
            edit["field"] = bind
            edit["old_repr"] = cfg.value_repr
            edit["new_value"] = value
            edit["new_repr"] = repr(value)
        edits.append(edit)
    return edits, rejected, None


def split_edits(edits: list[dict]) -> tuple[list[dict], list[dict]]:
    """分流 → (config_edits, text_revisions)；N/O 同填的行两边都进。"""
    config_edits = [e for e in edits if e.get("has_new_value")]
    text_revisions = [
        {"as_id": e["as_id"], "title": e["title"], "new_text": e["new_text"], "impact": e["impact"], "milestones": e["milestones"]}
        for e in edits if e.get("new_text")
    ]
    return config_edits, text_revisions


def merged_scope(edits: list[dict]) -> str:
    best = "output"
    for e in edits:
        cur = e.get("impact", "")
        if IMPACT_ORDER.get(cur, -1) > IMPACT_ORDER.get(best, -1):
            best = cur
    return best


def _check_field_conflicts(config_edits: list[dict], rejected: list[dict]) -> list[dict]:
    """同批次两行绑同一 config 字段 → 两行都拒绝（不猜优先级）。"""
    by_field: dict[str, list[dict]] = {}
    for e in config_edits:
        by_field.setdefault(e["field"], []).append(e)
    conflicts = {f for f, es in by_field.items() if len(es) > 1}
    if not conflicts:
        return config_edits
    kept = []
    for e in config_edits:
        if e["field"] in conflicts:
            rejected.append({
                "row": "-", "as_id": e["as_id"],
                "reason": f"同批次多条修改绑到同一字段 {e['field']}，请保留一条后重新提交",
            })
        else:
            kept.append(e)
    return kept


# ---------------------------------------------------------------------------
# apply（merged 形态唯一入口）/ set-param
# ---------------------------------------------------------------------------


def _apply_config_edits(
    root: Path, report_id: str, config_path: Path, config_edits: list[dict], tag: str
) -> tuple[list[dict], list[dict], dict]:
    """逐字段串行：内存替换→ast 校验→写盘→实例化冒烟；失败单字段回滚降级。

    返回 (applied, degraded, artifacts{backup_path, diff_path})。
    """
    original = config_path.read_text(encoding="utf-8")
    backup = backup_config(root, report_id, config_path, tag)
    baseline_ok, baseline_err = instantiate_smoke(config_path)
    smoke_available = baseline_ok
    if not baseline_ok:
        print(f"[warn] 原 config 实例化冒烟失败，本次仅做语法校验（跳过冒烟）：{baseline_err}", file=sys.stderr)

    applied: list[dict] = []
    degraded: list[dict] = []
    for edit in config_edits:
        field = parse_config_fields(config_path)[edit["field"]]
        try:
            new_text = apply_literal_to_text(config_path.read_text(encoding="utf-8"), field, edit["new_value"])
        except SyntaxError as e:
            degraded.append({**edit, "reason": f"改写后语法错误：{e}"})
            continue
        before_edit = config_path.read_text(encoding="utf-8")
        _atomic_write_text(config_path, new_text)
        if smoke_available:
            ok, err = instantiate_smoke(config_path)
            if not ok:
                _atomic_write_text(config_path, before_edit)
                degraded.append({**edit, "reason": f"实例化冒烟失败（已回滚）：{err}"})
                continue
        applied.append(edit)

    final_text = config_path.read_text(encoding="utf-8")
    diff_path = write_diff_patch(root, report_id, tag, original, final_text, config_path)
    return applied, degraded, {"backup_path": str(backup), "diff_path": str(diff_path) if diff_path else None}


def cmd_apply(args: argparse.Namespace) -> int:
    root = st.default_root()
    md_path = st.workspace_dir(root, args.report_id) / "assumptions.md"
    xlsx_path = Path(args.xlsx) if args.xlsx else st.workspace_dir(root, args.report_id) / "assumptions.xlsx"
    try:
        blocks = load_as_blocks(md_path)
    except (FileNotFoundError, ValueError) as e:
        print(str(e), file=sys.stderr)
        return 1
    if not xlsx_path.exists():
        print(f"工作表不存在：{xlsx_path}（先跑 export）", file=sys.stderr)
        return 1
    config_path = st.src_dir(root, args.report_id) / "config.py"
    fields = parse_config_fields(config_path)

    edits, rejected, hard_error = read_row_edits(xlsx_path, blocks, fields)
    if hard_error:
        print(hard_error, file=sys.stderr)
        return 1
    config_edits, text_revisions = split_edits(edits)
    config_edits = _check_field_conflicts(config_edits, rejected)
    # 字段冲突行整体失效（含其新口径），从 edits/text_revisions 同步剔除
    kept_ids = {id(e) for e in config_edits}
    dropped_as = {e["as_id"] for e in edits if e.get("has_new_value") and id(e) not in kept_ids}
    if dropped_as:
        edits = [e for e in edits if e["as_id"] not in dropped_as]
        text_revisions = [t for t in text_revisions if t["as_id"] not in dropped_as]

    if args.each:
        if text_revisions:
            print(json.dumps({
                "mode": "each", "error": "--each 仅支持纯参数直改行（文本修改无法自动回滚），"
                "请改用合并模式（--from-excel 不加 --each）或清空下列行的新口径："
                f"{sorted({t['as_id'] for t in text_revisions})}",
            }, ensure_ascii=False, indent=2))
            return 1

    batch_id = datetime.now().strftime("%Y%m%d_%H%M%S")
    scope = merged_scope(edits) if edits else None
    milestones = sorted({m for e in edits for m in e["milestones"]})

    def plan_json(mode: str) -> dict:
        return {
            "mode": mode,
            "batch_id": batch_id,
            "config_edits": [
                {"as_id": e["as_id"], "field": e["field"], "old": e["old_repr"], "new": e["new_repr"], "impact": e["impact"]}
                for e in config_edits
            ],
            "text_revisions": text_revisions,
            "rejected": rejected,
            "merged_impact_scope": scope,
            "milestones_affected": milestones,
        }

    if args.dry_run or args.each:
        plan = plan_json("dry-run" if args.dry_run else "each")
        plan["assumptions_stats_after"] = {
            **assumptions_stats(blocks),
            "revised": assumptions_stats(blocks)["revised"] + len({e["as_id"] for e in edits}),
        }
        print(json.dumps(plan, ensure_ascii=False, indent=2))
        return 0 if edits else 1

    if not edits:
        print(json.dumps(plan_json("applied-nothing"), ensure_ascii=False, indent=2))
        return 1

    # ---- 正式落地 ----
    applied: list[dict] = []
    degraded: list[dict] = []
    artifacts: dict = {}
    if config_edits:
        applied, degraded, artifacts = _apply_config_edits(root, args.report_id, config_path, config_edits, batch_id)
        # 降级编辑的文本通道：若该行本来就填了新口径则并入 text_revisions 已由 split 覆盖；
        # 未填新口径的降级编辑补一条说明性 text_revision，保证口径变化不丢失
        for d in degraded:
            if not any(t["as_id"] == d["as_id"] for t in text_revisions):
                text_revisions.append({
                    "as_id": d["as_id"], "title": d["title"],
                    "new_text": f"（参数直改降级）拟将 {d['field']} 由 {d['old_repr']} 改为 {d['new_repr']}；自动注入失败原因：{d['reason']}",
                    "impact": d["impact"], "milestones": d["milestones"],
                })

    md_failed: list[dict] = []
    if applied:
        md_edits = [
            {"as_id": e["as_id"], "field": e["field"], "old_repr": e["old_repr"], "new_repr": e["new_repr"]}
            for e in applied
        ]
        md_failed, new_md = apply_param_md_edits(md_path, md_edits, batch_id, st.now_iso())
        if not md_failed or len(md_failed) < len(md_edits):
            _atomic_write_text(md_path, new_md)
        # 全部失败则不落盘（保持原 md）
        elif md_failed:
            print("[warn] 全部 md 追记失败，登记簿未改写，请派 quant-planner 补登记", file=sys.stderr)

    stats = assumptions_stats(load_as_blocks(md_path)) if applied else assumptions_stats(blocks)
    result = plan_json("applied")
    result.update({
        "applied_config_edits": [
            {"as_id": e["as_id"], "field": e["field"], "old": e["old_repr"], "new": e["new_repr"]} for e in applied
        ],
        "config_edits_degraded": [
            {"as_id": d["as_id"], "field": d["field"], "old": d["old_repr"], "new": d["new_repr"], "reason": d["reason"]}
            for d in degraded
        ],
        "md_edits_failed": md_failed,
        "text_revisions": text_revisions,
        **artifacts,
        "assumptions_stats": stats,
    })
    print(json.dumps(result, ensure_ascii=False, indent=2))
    return 0 if (applied or text_revisions) else 1


def cmd_set_param(args: argparse.Namespace) -> int:
    root = st.default_root()
    config_path = st.src_dir(root, args.report_id) / "config.py"
    fields = parse_config_fields(config_path)
    cfg = fields.get(args.field)
    if cfg is None:
        print(f"字段 {args.field} 在 {config_path} 中不存在", file=sys.stderr)
        return 1
    if not cfg.bindable:
        print(f"字段 {args.field} 不可注入（{cfg.value_repr}）", file=sys.stderr)
        return 1
    ok, value, reason = coerce_new_value(args.value, cfg.value)
    if not ok:
        print(f"字段 {args.field}：{reason}", file=sys.stderr)
        return 1
    if value == cfg.value:
        print(f"字段 {args.field} 新值与当前值相同（{cfg.value_repr}），无变化", file=sys.stderr)
        return 1
    tag = datetime.now().strftime("%Y%m%d_%H%M%S") + f"_{args.field}"
    backup = backup_config(root, args.report_id, config_path, tag)
    before = config_path.read_text(encoding="utf-8")
    new_text = apply_literal_to_text(before, cfg, value)
    _atomic_write_text(config_path, new_text)
    smoke_ok, smoke_err = instantiate_smoke(config_path)
    if not smoke_ok:
        _atomic_write_text(config_path, before)
        print(f"实例化冒烟失败（已回滚）：{smoke_err}", file=sys.stderr)
        return 1
    print(json.dumps({
        "field": args.field, "old_repr": cfg.value_repr, "new_repr": repr(value),
        "value": value, "backup_path": str(backup),
    }, ensure_ascii=False))
    return 0


def cmd_restore(args: argparse.Namespace) -> int:
    root = st.default_root()
    try:
        used = restore_config(root, args.report_id, Path(args.backup) if args.backup else None)
    except FileNotFoundError as e:
        print(str(e), file=sys.stderr)
        return 1
    print(json.dumps({"restored_from": str(used)}, ensure_ascii=False))
    return 0


# ---------------------------------------------------------------------------
# sensitivity（--each 聚合）
# ---------------------------------------------------------------------------


def _load_metrics(path: Path) -> dict[str, dict]:
    data = json.loads(path.read_text(encoding="utf-8"))
    out: dict[str, dict] = {}
    for m in data.get("metrics", []):
        key = m.get("key") or m.get("name")
        if key:
            out[key] = m
    return out


def cmd_sensitivity(args: argparse.Namespace) -> int:
    root = st.default_root()
    baseline_path = Path(args.baseline)
    runs_dir = Path(args.runs_dir)
    if not baseline_path.exists():
        print(f"基线 comparison.json 不存在：{baseline_path}", file=sys.stderr)
        return 1
    baseline = _load_metrics(baseline_path)
    run_files = sorted(runs_dir.glob("AS*_comparison.json"))
    if not run_files:
        print(f"无敏感性快照：{runs_dir}/AS*_comparison.json", file=sys.stderr)
        return 1

    wb = Workbook()
    ws = wb.active
    ws.title = "敏感性明细"
    detail_headers = ["AS ID", "参数字段", "旧值", "新值", "指标 key", "指标名", "基线值", "变体值", "Δ", "Δ%"]
    ws.append(detail_headers)
    for cell in ws[1]:
        cell.font = HEADER_FONT

    summary_rows: list[list[Any]] = []
    for rf in run_files:
        as_id = rf.name.split("_comparison.json")[0]
        meta_path = runs_dir / f"{as_id}_meta.json"
        meta = json.loads(meta_path.read_text(encoding="utf-8")) if meta_path.exists() else {}
        variant = _load_metrics(rf)
        n_changed = 0
        max_pct: tuple[float, str] = (0.0, "")
        for key, bm in baseline.items():
            vm = variant.get(key)
            if vm is None:
                continue
            bv, vv = bm.get("reproduced_value"), vm.get("reproduced_value")
            if not isinstance(bv, (int, float)) or not isinstance(vv, (int, float)):
                continue
            delta = vv - bv
            pct = delta / abs(bv) if bv else None
            if delta != 0:
                n_changed += 1
            if pct is not None and abs(pct) > max_pct[0]:
                max_pct = (abs(pct), key)
            ws.append([
                as_id, meta.get("field", ""), meta.get("old_repr", ""), meta.get("new_repr", ""),
                key, bm.get("name", key), bv, vv, round(delta, 6), round(pct, 6) if pct is not None else None,
            ])
        summary_rows.append([
            as_id, meta.get("field", ""), f"{meta.get('old_repr', '?')} → {meta.get('new_repr', '?')}",
            len(baseline), n_changed, max_pct[1] or "-", round(max_pct[0], 6) if max_pct[1] else None,
        ])

    _auto_width(ws, len(detail_headers))
    ws.freeze_panes = "A2"

    ws2 = wb.create_sheet("汇总")
    summary_headers = ["AS ID", "参数字段", "旧值 → 新值", "指标总数", "变动指标数", "最大|Δ%|指标", "最大|Δ%|"]
    ws2.append(summary_headers)
    for cell in ws2[1]:
        cell.font = HEADER_FONT
    for row in summary_rows:
        ws2.append(row)
    _auto_width(ws2, len(summary_headers))
    ws2.freeze_panes = "A2"

    out_path = root / "output" / args.report_id / "results" / "assumption_sensitivity.xlsx"
    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)
    print(json.dumps({
        "xlsx": str(out_path.relative_to(root)),
        "runs": [rf.name.split("_comparison.json")[0] for rf in run_files],
        "baseline_metrics": len(baseline),
    }, ensure_ascii=False, indent=2))
    return 0


# ---------------------------------------------------------------------------
# CLI
# ---------------------------------------------------------------------------


def build_parser() -> argparse.ArgumentParser:
    p = argparse.ArgumentParser(description="假设登记簿 Excel 导出/回灌工具")
    sub = p.add_subparsers(dest="command", required=True)

    pe = sub.add_parser("export", help="assumptions.md → assumptions.xlsx")
    pe.add_argument("report_id")
    pe.add_argument("--xlsx", default=None)
    pe.set_defaults(func=cmd_export)

    pa = sub.add_parser("apply", help="读回 xlsx → 分流落地（merged）")
    pa.add_argument("report_id")
    pa.add_argument("--xlsx", default=None)
    pa.add_argument("--dry-run", action="store_true")
    pa.add_argument("--each", action="store_true", help="只解析校验输出计划（逐条敏感性模式）")
    pa.set_defaults(func=cmd_apply)

    ps = sub.add_parser("set-param", help="只改 config.py 单字段（--each 脚手架）")
    ps.add_argument("report_id")
    ps.add_argument("--field", required=True)
    ps.add_argument("--value", required=True)
    ps.set_defaults(func=cmd_set_param)

    pr = sub.add_parser("restore", help="回滚 config.py 到备份")
    pr.add_argument("report_id")
    pr.add_argument("--backup", default=None)
    pr.set_defaults(func=cmd_restore)

    pn = sub.add_parser("sensitivity", help="聚合 --each 快照 → assumption_sensitivity.xlsx")
    pn.add_argument("report_id")
    pn.add_argument("--baseline", required=True)
    pn.add_argument("--runs-dir", required=True)
    pn.set_defaults(func=cmd_sensitivity)
    return p


def main() -> int:
    args = build_parser().parse_args()
    return args.func(args)


if __name__ == "__main__":
    sys.exit(main())
