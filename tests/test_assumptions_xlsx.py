"""tools/assumptions_xlsx.py 的测试：Excel 导出/回灌、config 安全改写与降级、敏感性聚合。

fixture 策略：tmp_path 伪造最小 workspace（assumptions.md + src/<id>/config.py），
经 REPORT_REPRODUCE_ROOT 环境变量隔离，不依赖任何真实案例数据。
"""

from __future__ import annotations

import json
import sys
from pathlib import Path

import pytest
from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from tools.assumptions_xlsx import main  # noqa: E402

REPORT_ID = "demo"

ASSUMPTIONS_MD = """# 假设登记簿 assumptions.md

### [AS1] 主绩效表按零成本复现
- 来源: A1（auto 裁决）
- 假设内容: 主绩效表设定为单边 0 bps（零成本）
- 行业惯例依据: 业内常见默认
- 影响面: trading
- 影响 milestone: m4
- 影响指标: R1.Sharpe
- 预期影响: 若口径错误，复现值系统性偏高
- 验证后回看: 实测通过，假设获数据支持
- 状态: assumed
- 高亮等级: major-auto
- 参数绑定: cost_bps

### [AS2] 日历周期编码
- 来源: A2（auto 裁决）
- 假设内容: 日历特征按 sin/cos 周期编码
- 行业惯例依据: 业内常见默认
- 影响面: param
- 影响 milestone: m1
- 影响指标: 无
- 预期影响: 无
- 验证后回看: 实测通过
- 状态: assumed
- 高亮等级: -
- 参数绑定: calendar_cycles

### [AS3] 训练起点迁移
- 来源: 数据限制
- 性质: market-transplant
- 假设内容: train 起点 2000 迁移为 2007
- 行业惯例依据: 本地数据约束
- 影响面: data
- 影响 milestone: m1
- 影响指标: 全部 R
- 预期影响: 数值豁免
- 验证后回看: 实测通过
- 状态: confirmed
- 高亮等级: major-auto

### [AS4] 回看窗口 60 日
- 来源: A4（auto 裁决）
- 假设内容: 回看窗口 H=60
- 行业惯例依据: 业内常见默认
- 影响面: param
- 影响 milestone: m2
- 影响指标: R2
- 预期影响: 无
- 验证后回看: 实测通过
- 状态: assumed
- 高亮等级: -
- 参数绑定: lookback_h

### [AS5] 换手率口径
- 来源: A5（auto 裁决）
- 假设内容: 换手率按多头组单边口径
- 行业惯例依据: 业内常见默认
- 影响面: trading
- 影响 milestone: m1
- 影响指标: R1.Turnover
- 预期影响: 偏差约 2 倍
- 验证后回看: 实测通过
- 状态: assumed
- 高亮等级: -
"""

CONFIG_PY = '''"""demo config."""

from dataclasses import dataclass, field


@dataclass(frozen=True)
class DemoConfig:
    """测试配置。"""

    cost_bps: float = 0.0  # 主表单边成本（AS1）
    lookback_h: int = 60  # 回看窗口（AS4）
    mode_name: str = "month_end"
    seeds: tuple = (0, 1, 2)
    calendar_cycles: tuple = field(default_factory=lambda: (7.0, 12.0))

    def __post_init__(self) -> None:
        if self.lookback_h > 100:
            raise ValueError("lookback_h 超出可支撑历史范围")


DEFAULT_CONFIG = DemoConfig()
'''


@pytest.fixture()
def case(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> Path:
    monkeypatch.setenv("REPORT_REPRODUCE_ROOT", str(tmp_path))
    ws = tmp_path / "workspace" / REPORT_ID
    ws.mkdir(parents=True)
    (ws / "assumptions.md").write_text(ASSUMPTIONS_MD, encoding="utf-8")
    src = tmp_path / "src" / REPORT_ID
    src.mkdir(parents=True)
    (src / "config.py").write_text(CONFIG_PY, encoding="utf-8")
    return tmp_path


def run_cli(monkeypatch: pytest.MonkeyPatch, capsys, argv: list[str]) -> tuple[int, dict | None, str]:
    monkeypatch.setattr(sys, "argv", ["assumptions_xlsx.py", *argv])
    code = main()
    out = capsys.readouterr()
    try:
        payload = json.loads(out.out)
    except json.JSONDecodeError:
        payload = None
    return code, payload, out.err


def _edit_xlsx(xlsx_path: Path, edits: dict[str, dict]) -> None:
    """按 AS ID 定位行写入 新值/新口径/操作；AS ID 不存在则追加新行（模拟客户自行加行）。"""
    wb = load_workbook(xlsx_path)
    ws = wb["假设登记簿"]
    headers = {c.value: c.column for c in ws[1]}
    found: set[str] = set()
    for row in ws.iter_rows(min_row=2):
        as_id = row[headers["AS ID"] - 1].value
        if as_id in edits:
            found.add(as_id)
            e = edits[as_id]
            if "new_value" in e:
                row[headers["新值"] - 1].value = e["new_value"]
            if "new_text" in e:
                row[headers["新口径（文本说明）"] - 1].value = e["new_text"]
            if "op" in e:
                row[headers["操作"] - 1].value = e["op"]
    for as_id, e in edits.items():
        if as_id not in found:
            new_row = [None] * len(headers)
            new_row[headers["AS ID"] - 1] = as_id
            if "new_value" in e:
                new_row[headers["新值"] - 1] = e["new_value"]
            if "new_text" in e:
                new_row[headers["新口径（文本说明）"] - 1] = e["new_text"]
            if "op" in e:
                new_row[headers["操作"] - 1] = e["op"]
            ws.append(new_row)
    wb.save(xlsx_path)


# ---------------------------------------------------------------------------
# export
# ---------------------------------------------------------------------------


def test_export(case: Path, monkeypatch, capsys) -> None:
    code, payload, _ = run_cli(monkeypatch, capsys, ["export", REPORT_ID])
    assert code == 0 and payload is not None
    xlsx = case / payload["xlsx"]
    assert xlsx.exists()
    assert payload["total"] == 5
    assert payload["assumptions_stats"] == {"total": 5, "assumed": 4, "confirmed": 1, "revised": 0}
    assert {b["as_id"] for b in payload["bound"]} == {"AS1", "AS2", "AS4"}
    # AS5（trading 无绑定）进入待回填提示；AS3 影响面 data 不进
    assert [c["as_id"] for c in payload["unbound_param_candidates"]] == ["AS5"]

    wb = load_workbook(xlsx)
    assert wb.sheetnames == ["假设登记簿", "填写说明"]
    ws = wb["假设登记簿"]
    assert ws.freeze_panes == "A2"
    assert ws.max_row == 6  # 表头 + 5 条
    headers = {c.value: c.column for c in ws[1]}
    assert len(headers) == 16
    rows = {row[headers["AS ID"] - 1].value: row for row in ws.iter_rows(min_row=2)}
    # 当前值列：字面量 repr / 不可绑定标注 / 无绑定留空
    assert rows["AS1"][headers["当前值"] - 1].value == "0.0"
    assert "不可绑定" in rows["AS2"][headers["当前值"] - 1].value
    assert rows["AS3"][headers["当前值"] - 1].value is None
    assert rows["AS3"][headers["性质"] - 1].value == "market-transplant"
    # 可编辑三列：黄底 + 解锁；只读列保持锁定
    for r in rows.values():
        for name in ("新值", "新口径（文本说明）", "操作"):
            cell = r[headers[name] - 1]
            assert cell.protection.locked is False
            assert cell.fill.fgColor.rgb == "00FFF2CC"
        assert r[headers["标题"] - 1].protection.locked is True
    assert ws.protection.sheet is True
    assert len(ws.data_validations.dataValidation) == 1


# ---------------------------------------------------------------------------
# apply 全链路
# ---------------------------------------------------------------------------


def test_apply_full(case: Path, monkeypatch, capsys) -> None:
    run_cli(monkeypatch, capsys, ["export", REPORT_ID])
    xlsx = case / "workspace" / REPORT_ID / "assumptions.xlsx"
    _edit_xlsx(xlsx, {
        "AS1": {"new_value": 5.0, "op": "revise"},
        "AS3": {"new_text": "起点改用 2010 年", "op": "revise"},
    })

    code, plan, _ = run_cli(monkeypatch, capsys, ["apply", REPORT_ID, "--dry-run"])
    assert code == 0 and plan["mode"] == "dry-run"
    assert [e["as_id"] for e in plan["config_edits"]] == ["AS1"]
    assert [t["as_id"] for t in plan["text_revisions"]] == ["AS3"]
    assert plan["merged_impact_scope"] == "data"  # trading 与 data 取最重
    assert plan["rejected"] == []
    # dry-run 不落盘
    assert "cost_bps: float = 0.0" in (case / "src" / REPORT_ID / "config.py").read_text()

    code, result, _ = run_cli(monkeypatch, capsys, ["apply", REPORT_ID])
    assert code == 0 and result["mode"] == "applied"
    assert [e["as_id"] for e in result["applied_config_edits"]] == ["AS1"]
    assert result["config_edits_degraded"] == []
    assert result["md_edits_failed"] == []

    config_text = (case / "src" / REPORT_ID / "config.py").read_text()
    assert "cost_bps: float = 5.0  # 主表单边成本（AS1）" in config_text  # 注释保留
    assert Path(result["backup_path"]).exists()
    assert result["diff_path"] and Path(result["diff_path"]).exists()

    md = (case / "workspace" / REPORT_ID / "assumptions.md").read_text()
    as1_block = md.split("### [AS1]")[1].split("### [AS2]")[0]
    assert "- 状态: revised" in as1_block
    assert "- 验证后回看: [verify 后填]" in as1_block
    assert "修订记录" in as1_block and "cost_bps 0.0 → 5.0" in as1_block
    assert "工具直接注入 config.py" in as1_block
    # AS3 是文本通道，md 不由工具改写
    as3_block = md.split("### [AS3]")[1].split("### [AS4]")[0]
    assert "- 状态: confirmed" in as3_block
    assert result["assumptions_stats"] == {"total": 5, "assumed": 3, "confirmed": 1, "revised": 1}


def test_apply_smoke_degrade(case: Path, monkeypatch, capsys) -> None:
    """新值触发 __post_init__ 校验失败 → 回滚 + 降级为文本通道。"""
    run_cli(monkeypatch, capsys, ["export", REPORT_ID])
    xlsx = case / "workspace" / REPORT_ID / "assumptions.xlsx"
    _edit_xlsx(xlsx, {"AS4": {"new_value": 200, "op": "revise"}})
    code, result, _ = run_cli(monkeypatch, capsys, ["apply", REPORT_ID])
    assert code == 0
    assert result["applied_config_edits"] == []
    assert [d["as_id"] for d in result["config_edits_degraded"]] == ["AS4"]
    assert "冒烟" in result["config_edits_degraded"][0]["reason"]
    # config 回滚原样
    assert "lookback_h: int = 60" in (case / "src" / REPORT_ID / "config.py").read_text()
    # 降级补一条说明性 text_revision，口径变化不丢失
    assert any(t["as_id"] == "AS4" and "参数直改降级" in t["new_text"] for t in result["text_revisions"])
    # md 未被改写（无成功注入）
    md = (case / "workspace" / REPORT_ID / "assumptions.md").read_text()
    assert "状态: revised" not in md


# ---------------------------------------------------------------------------
# 容错与拒绝路径
# ---------------------------------------------------------------------------


@pytest.mark.parametrize(
    "edits, expect_as, expect_kw",
    [
        ({"AS99": {"new_text": "x", "op": "revise"}}, "AS99", "不在登记簿"),
        ({"AS1": {"new_value": 5.0}}, "AS1", "未选 revise"),
        ({"AS2": {"new_value": "(7,)", "op": "revise"}}, "AS2", "不可注入"),
        ({"AS1": {"new_value": 0.0, "op": "revise"}}, "AS1", "无变化"),
        ({"AS3": {"new_value": 5, "op": "revise"}}, "AS3", "无参数绑定"),
        ({"AS1": {"op": "revise"}}, "AS1", "均为空"),
        ({"AS1": {"new_text": "x", "op": "bogus"}}, "AS1", "取值无效"),
    ],
)
def test_apply_rejections(case: Path, monkeypatch, capsys, edits, expect_as, expect_kw) -> None:
    run_cli(monkeypatch, capsys, ["export", REPORT_ID])
    xlsx = case / "workspace" / REPORT_ID / "assumptions.xlsx"
    _edit_xlsx(xlsx, edits)
    code, result, _ = run_cli(monkeypatch, capsys, ["apply", REPORT_ID, "--dry-run"])
    assert code == 1  # 全部被拒 → 无生效行
    hit = [r for r in result["rejected"] if r["as_id"] == expect_as and expect_kw in r["reason"]]
    assert hit, f"未找到预期拒绝：{expect_as} / {expect_kw}；实际 {result['rejected']}"
    assert result["config_edits"] == [] and result["text_revisions"] == []


def test_apply_field_conflict(case: Path, monkeypatch, capsys) -> None:
    """同批次两行绑同一字段 → 两行都拒绝。"""
    md_path = case / "workspace" / REPORT_ID / "assumptions.md"
    md_path.write_text(md_path.read_text() + "\n### [AS6] 成本第二条\n- 来源: coder 主动简化\n- 假设内容: 第二条成本\n- 行业惯例依据: 无\n- 影响面: trading\n- 影响 milestone: m4\n- 影响指标: R1\n- 预期影响: 无\n- 验证后回看: 实测通过\n- 状态: assumed\n- 高亮等级: -\n- 参数绑定: cost_bps\n", encoding="utf-8")
    run_cli(monkeypatch, capsys, ["export", REPORT_ID])
    xlsx = case / "workspace" / REPORT_ID / "assumptions.xlsx"
    _edit_xlsx(xlsx, {
        "AS1": {"new_value": 5.0, "op": "revise"},
        "AS6": {"new_value": 10.0, "op": "revise"},
    })
    code, result, _ = run_cli(monkeypatch, capsys, ["apply", REPORT_ID, "--dry-run"])
    assert code == 1
    assert result["config_edits"] == []
    assert {r["as_id"] for r in result["rejected"]} == {"AS1", "AS6"}
    assert all("同一字段" in r["reason"] for r in result["rejected"])


def test_apply_broken_header(case: Path, monkeypatch, capsys) -> None:
    """核心列被改名 → 硬错误提示重新 export；垃圾附加列 → 忽略。"""
    run_cli(monkeypatch, capsys, ["export", REPORT_ID])
    xlsx = case / "workspace" / REPORT_ID / "assumptions.xlsx"
    wb = load_workbook(xlsx)
    ws = wb["假设登记簿"]
    headers = {c.value: c.column for c in ws[1]}
    ws.cell(row=1, column=headers["操作"], value="随便改名")
    wb.save(xlsx)
    code, _, err = run_cli(monkeypatch, capsys, ["apply", REPORT_ID, "--dry-run"])
    assert code == 1 and "缺少核心列" in err


def test_apply_value_falls_back_to_text(case: Path, monkeypatch, capsys) -> None:
    """新值校验失败但同填新口径 → 新值拒绝、新口径仍生效（降级纯文本行）。"""
    run_cli(monkeypatch, capsys, ["export", REPORT_ID])
    xlsx = case / "workspace" / REPORT_ID / "assumptions.xlsx"
    _edit_xlsx(xlsx, {"AS1": {"new_value": "not_a_number", "new_text": "成本改 5bps", "op": "revise"}})
    code, result, _ = run_cli(monkeypatch, capsys, ["apply", REPORT_ID, "--dry-run"])
    assert code == 0
    assert result["config_edits"] == []
    assert [t["as_id"] for t in result["text_revisions"]] == ["AS1"]
    assert any("无法解析" in r["reason"] for r in result["rejected"])


# ---------------------------------------------------------------------------
# --each 守卫 / set-param / restore
# ---------------------------------------------------------------------------


def test_apply_each_guard(case: Path, monkeypatch, capsys) -> None:
    run_cli(monkeypatch, capsys, ["export", REPORT_ID])
    xlsx = case / "workspace" / REPORT_ID / "assumptions.xlsx"
    _edit_xlsx(xlsx, {
        "AS1": {"new_value": 5.0, "op": "revise"},
        "AS3": {"new_text": "文本行", "op": "revise"},
    })
    code, result, _ = run_cli(monkeypatch, capsys, ["apply", REPORT_ID, "--each"])
    assert code == 1 and "--each 仅支持纯参数" in result["error"]
    assert "cost_bps: float = 0.0" in (case / "src" / REPORT_ID / "config.py").read_text()

    _edit_xlsx(xlsx, {"AS3": {"new_text": None}})
    wb = load_workbook(xlsx)
    ws = wb["假设登记簿"]
    headers = {c.value: c.column for c in ws[1]}
    for row in ws.iter_rows(min_row=2):
        if row[headers["AS ID"] - 1].value == "AS3":
            row[headers["新口径（文本说明）"] - 1].value = None
    wb.save(xlsx)
    code, result, _ = run_cli(monkeypatch, capsys, ["apply", REPORT_ID, "--each"])
    assert code == 0 and result["mode"] == "each"
    assert [e["as_id"] for e in result["config_edits"]] == ["AS1"]
    assert "cost_bps: float = 0.0" in (case / "src" / REPORT_ID / "config.py").read_text()  # 不落盘


def test_set_param_and_restore(case: Path, monkeypatch, capsys) -> None:
    config_path = case / "src" / REPORT_ID / "config.py"
    original = config_path.read_bytes()

    code, payload, _ = run_cli(monkeypatch, capsys, ["set-param", REPORT_ID, "--field", "cost_bps", "--value", "7.5"])
    assert code == 0 and payload["new_repr"] == "7.5"
    assert "cost_bps: float = 7.5" in config_path.read_text()
    assert Path(payload["backup_path"]).exists()

    # 非法值不留残改
    code, _, err = run_cli(monkeypatch, capsys, ["set-param", REPORT_ID, "--field", "lookback_h", "--value", "abc"])
    assert code == 1 and "无法解析" in err
    assert "lookback_h: int = 60" in config_path.read_text()

    code, payload, _ = run_cli(monkeypatch, capsys, ["restore", REPORT_ID, "--backup", str(case / "workspace" / REPORT_ID / "audit" / "config_backups" / Path(payload["backup_path"]).name)])
    assert code == 0
    # --backup 指定的是 set-param 备份（改动前原文）→ 字节级复原
    assert config_path.read_bytes() == original


def test_restore_default_latest(case: Path, monkeypatch, capsys) -> None:
    run_cli(monkeypatch, capsys, ["set-param", REPORT_ID, "--field", "cost_bps", "--value", "7.5"])
    code, payload, _ = run_cli(monkeypatch, capsys, ["restore", REPORT_ID])
    assert code == 0 and "config_" in payload["restored_from"]
    assert "cost_bps: float = 0.0" in (case / "src" / REPORT_ID / "config.py").read_text()


# ---------------------------------------------------------------------------
# md 追记失败降级
# ---------------------------------------------------------------------------


def test_md_edit_failure_keeps_config(case: Path, monkeypatch, capsys) -> None:
    """条目缺「状态」行 → md_edits_failed 非空，config 改动保留。"""
    md_path = case / "workspace" / REPORT_ID / "assumptions.md"
    md_path.write_text(md_path.read_text().replace("- 状态: assumed\n- 高亮等级: major-auto\n- 参数绑定: cost_bps", "- 高亮等级: major-auto\n- 参数绑定: cost_bps"), encoding="utf-8")
    run_cli(monkeypatch, capsys, ["export", REPORT_ID])
    xlsx = case / "workspace" / REPORT_ID / "assumptions.xlsx"
    _edit_xlsx(xlsx, {"AS1": {"new_value": 5.0, "op": "revise"}})
    code, result, _ = run_cli(monkeypatch, capsys, ["apply", REPORT_ID])
    assert code == 0
    assert [e["as_id"] for e in result["applied_config_edits"]] == ["AS1"]
    assert [f["as_id"] for f in result["md_edits_failed"]] == ["AS1"]
    assert "cost_bps: float = 5.0" in (case / "src" / REPORT_ID / "config.py").read_text()


# ---------------------------------------------------------------------------
# sensitivity 聚合
# ---------------------------------------------------------------------------


def _comparison(metrics: dict[str, float]) -> str:
    return json.dumps({
        "metrics": [
            {"key": k, "name": k.upper(), "report_value": 1.0, "reproduced_value": v, "pass": True}
            for k, v in metrics.items()
        ]
    })


def test_sensitivity(case: Path, monkeypatch, capsys) -> None:
    runs = case / "workspace" / REPORT_ID / "audit" / "sensitivity"
    runs.mkdir(parents=True)
    baseline = runs / "_baseline.json"
    baseline.write_text(_comparison({"sharpe": 0.7, "ret": 0.1}), encoding="utf-8")
    (runs / "AS1_comparison.json").write_text(_comparison({"sharpe": 0.6, "ret": 0.1}), encoding="utf-8")
    (runs / "AS1_meta.json").write_text(json.dumps({"field": "cost_bps", "old_repr": "0.0", "new_repr": "5.0"}), encoding="utf-8")
    (runs / "AS4_comparison.json").write_text(_comparison({"sharpe": 0.7, "ret": 0.12}), encoding="utf-8")
    (runs / "AS4_meta.json").write_text(json.dumps({"field": "lookback_h", "old_repr": "60", "new_repr": "40"}), encoding="utf-8")

    code, payload, _ = run_cli(monkeypatch, capsys, [
        "sensitivity", REPORT_ID, "--baseline", str(baseline), "--runs-dir", str(runs),
    ])
    assert code == 0 and payload["runs"] == ["AS1", "AS4"]
    wb = load_workbook(case / payload["xlsx"])
    detail = wb["敏感性明细"]
    assert detail.max_row == 1 + 2 * 2  # 表头 + 2 AS × 2 指标
    rows = list(detail.iter_rows(min_row=2, values_only=True))
    as1_sharpe = next(r for r in rows if r[0] == "AS1" and r[4] == "sharpe")
    assert as1_sharpe[1:4] == ("cost_bps", "0.0", "5.0")
    assert as1_sharpe[6] == 0.7 and as1_sharpe[7] == 0.6
    assert as1_sharpe[8] == pytest.approx(-0.1)
    summary = wb["汇总"]
    assert summary.max_row == 1 + 2
    s1 = next(r for r in summary.iter_rows(min_row=2, values_only=True) if r[0] == "AS1")
    assert s1[3] == 2 and s1[4] == 1  # 指标总数 2、变动 1
    assert s1[5] == "sharpe"
