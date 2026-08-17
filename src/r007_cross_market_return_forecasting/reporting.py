"""r007 正式结果的 Excel、PNG、JSON 与运行日志输出。

图表只绘制传入的底层表；数据缺失或方法不适用时生成明确的
``not_applicable`` 证据页，不插值、不从研报图像估读，也不伪造曲线。
"""

from __future__ import annotations

import hashlib
import json
from collections.abc import Mapping, Sequence
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Any

import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
from matplotlib.figure import Figure
from openpyxl import load_workbook
from openpyxl.drawing.image import Image as ExcelImage
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from .config import DEFAULT_M5_CONFIG, M5Config


@dataclass
class ReportingResult:
    """正式输出文件、图表状态与哈希清单。"""

    files: tuple[Path, ...]
    chart_status: pd.DataFrame
    manifest: pd.DataFrame


def _jsonable(value: Any) -> Any:
    if isinstance(value, Mapping):
        return {str(key): _jsonable(item) for key, item in value.items()}
    if isinstance(value, (list, tuple, set)):
        return [_jsonable(item) for item in value]
    if isinstance(value, Path):
        return str(value)
    if isinstance(value, (pd.Timestamp, datetime)):
        return value.isoformat()
    if isinstance(value, np.generic):
        return _jsonable(value.item())
    if isinstance(value, float) and not np.isfinite(value):
        return None
    if pd.isna(value) if not isinstance(value, (dict, list, tuple)) else False:
        return None
    return value


def _write_json(path: Path, payload: Any) -> None:
    path.write_text(
        json.dumps(
            _jsonable(payload),
            ensure_ascii=False,
            indent=2,
            allow_nan=False,
            sort_keys=True,
        )
        + "\n",
        encoding="utf-8",
    )


def _require_columns(
    frame: pd.DataFrame,
    columns: Sequence[str],
    context: str,
) -> None:
    missing = sorted(set(columns) - set(frame.columns))
    if missing:
        raise KeyError(f"{context} 缺少字段: {missing}")


def _setup_plot_theme() -> None:
    try:
        plt.style.use("seaborn-v0_8-whitegrid")
    except OSError:
        plt.style.use("default")
    plt.rcParams.update(
        {
            "font.sans-serif": [
                "Arial Unicode MS",
                "PingFang SC",
                "Heiti SC",
                "DejaVu Sans",
            ],
            "axes.unicode_minus": False,
        }
    )


def _save_figure(figure: Figure, path: Path, *, config: M5Config) -> None:
    figure.tight_layout()
    figure.savefig(
        path,
        dpi=config.reporting.dpi,
        bbox_inches="tight",
        facecolor="white",
    )
    plt.close(figure)


def _not_applicable_chart(
    path: Path,
    *,
    title: str,
    reason: str,
    config: M5Config,
) -> dict[str, Any]:
    figure, axis = plt.subplots(figsize=config.reporting.bar_size)
    axis.axis("off")
    axis.text(
        0.5,
        0.62,
        title,
        ha="center",
        va="center",
        fontsize=config.reporting.title_fontsize,
        fontweight="bold",
    )
    axis.text(
        0.5,
        0.43,
        "not_applicable",
        ha="center",
        va="center",
        fontsize=config.reporting.axis_fontsize,
        color=config.reporting.primary_red,
    )
    axis.text(
        0.5,
        0.30,
        reason,
        ha="center",
        va="center",
        fontsize=config.reporting.legend_fontsize,
        wrap=True,
    )
    _save_figure(figure, path, config=config)
    return {
        "file": path.name,
        "status": "not_applicable",
        "reason": reason,
    }


def _line_chart(
    frame: pd.DataFrame,
    path: Path,
    *,
    x: str,
    y: str,
    series: str | None,
    title: str,
    xlabel: str,
    ylabel: str,
    config: M5Config,
) -> dict[str, Any]:
    if frame.empty:
        return _not_applicable_chart(
            path,
            title=title,
            reason="底层结果表为空，未生成伪造曲线",
            config=config,
        )
    required = [x, y, *((series,) if series else ())]
    _require_columns(frame, required, title)
    working = frame.loc[:, required].copy()
    working[y] = pd.to_numeric(working[y], errors="coerce")
    working = working.loc[np.isfinite(working[y])].copy()
    if working.empty:
        return _not_applicable_chart(
            path,
            title=title,
            reason="底层结果无有限数值，未生成伪造曲线",
            config=config,
        )
    figure, axis = plt.subplots(figsize=config.reporting.time_series_size)
    if series is None:
        axis.plot(
            working[x],
            working[y],
            color=config.reporting.primary_blue,
            linewidth=1.8,
        )
    else:
        groups = tuple(working[series].drop_duplicates().astype(str))
        colors = plt.cm.tab10(np.linspace(0, 1, max(len(groups), 1)))
        for color, (name, group) in zip(
            colors,
            working.groupby(series, observed=True, sort=False),
            strict=False,
        ):
            axis.plot(group[x], group[y], label=str(name), color=color, linewidth=1.4)
        if len(groups) <= 15:
            axis.legend(fontsize=config.reporting.legend_fontsize)
    axis.set_title(title, fontsize=config.reporting.title_fontsize)
    axis.set_xlabel(xlabel, fontsize=config.reporting.axis_fontsize)
    axis.set_ylabel(ylabel, fontsize=config.reporting.axis_fontsize)
    _save_figure(figure, path, config=config)
    return {"file": path.name, "status": "generated", "row_count": len(working)}


def _bar_chart(
    frame: pd.DataFrame,
    path: Path,
    *,
    category: str,
    value: str,
    title: str,
    xlabel: str,
    ylabel: str,
    config: M5Config,
) -> dict[str, Any]:
    if frame.empty:
        return _not_applicable_chart(
            path,
            title=title,
            reason="底层结果表为空，未生成伪造柱图",
            config=config,
        )
    _require_columns(frame, (category, value), title)
    working = frame.loc[:, [category, value]].copy()
    working[value] = pd.to_numeric(working[value], errors="coerce")
    working = working.loc[np.isfinite(working[value])].copy()
    if working.empty:
        return _not_applicable_chart(
            path,
            title=title,
            reason="底层结果无有限数值，未生成伪造柱图",
            config=config,
        )
    colors = np.where(
        working[value].ge(0),
        config.reporting.primary_blue,
        config.reporting.primary_red,
    )
    figure, axis = plt.subplots(figsize=config.reporting.bar_size)
    axis.bar(working[category].astype(str), working[value], color=colors)
    axis.set_title(title, fontsize=config.reporting.title_fontsize)
    axis.set_xlabel(xlabel, fontsize=config.reporting.axis_fontsize)
    axis.set_ylabel(ylabel, fontsize=config.reporting.axis_fontsize)
    axis.tick_params(axis="x", rotation=45)
    _save_figure(figure, path, config=config)
    return {"file": path.name, "status": "generated", "row_count": len(working)}


def build_figure13_long_table(paired_comparisons: pd.DataFrame) -> pd.DataFrame:
    """只把 R4 两类差值转换为 Figure 13 可绘制长表。"""
    required = (
        "element_id",
        "target_market",
        "feature_set",
        "model",
        "portfolio",
        "same_graph_minus_non_graph",
        "cross_minus_same_graph",
    )
    if paired_comparisons.empty:
        return pd.DataFrame(
            columns=[
                "element_id",
                "target_market",
                "feature_set",
                "model",
                "portfolio",
                "delta_type",
                "sharpe_delta",
                "chart_label",
            ]
        )
    _require_columns(paired_comparisons, required, "Figure 13 R4 对照")
    r4 = paired_comparisons.loc[
        paired_comparisons["element_id"].eq("R4"),
        list(required),
    ].copy()
    if r4.empty:
        return pd.DataFrame(
            columns=[
                *required[:5],
                "delta_type",
                "sharpe_delta",
                "chart_label",
            ]
        )
    long = r4.melt(
        id_vars=list(required[:5]),
        value_vars=[
            "same_graph_minus_non_graph",
            "cross_minus_same_graph",
        ],
        var_name="delta_type",
        value_name="sharpe_delta",
    )
    long["sharpe_delta"] = pd.to_numeric(
        long["sharpe_delta"], errors="coerce"
    )
    long["chart_label"] = long[
        ["target_market", "model", "portfolio", "delta_type"]
    ].astype(str).agg("/".join, axis=1)
    return long


def _heatmap_chart(
    matrix: pd.DataFrame,
    path: Path,
    *,
    title: str,
    config: M5Config,
) -> dict[str, Any]:
    if matrix.empty:
        return _not_applicable_chart(
            path,
            title=title,
            reason="底层矩阵为空，未生成伪造热图",
            config=config,
        )
    numeric = matrix.apply(pd.to_numeric, errors="coerce")
    if not np.isfinite(numeric.to_numpy(dtype=float, na_value=np.nan)).any():
        return _not_applicable_chart(
            path,
            title=title,
            reason="底层矩阵无有限值，未生成伪造热图",
            config=config,
        )
    figure, axis = plt.subplots(figsize=config.reporting.time_series_size)
    image = axis.imshow(
        numeric.to_numpy(dtype=float, na_value=np.nan),
        aspect="auto",
        cmap=config.reporting.matrix_cmap,
    )
    axis.set_title(title, fontsize=config.reporting.title_fontsize)
    axis.set_xlabel("源节点", fontsize=config.reporting.axis_fontsize)
    axis.set_ylabel("目标节点", fontsize=config.reporting.axis_fontsize)
    figure.colorbar(image, ax=axis, label="t统计量")
    _save_figure(figure, path, config=config)
    return {
        "file": path.name,
        "status": "generated",
        "matrix_shape": list(numeric.shape),
    }


def _build_monthly_ic(predictions: pd.DataFrame) -> pd.DataFrame:
    required = ("target_date", "model", "prediction", "actual_return")
    if predictions.empty or not set(required).issubset(predictions.columns):
        return pd.DataFrame(columns=["month", "model", "ic"])
    working = predictions.loc[:, list(required)].copy()
    working["target_date"] = pd.to_datetime(
        working["target_date"], errors="coerce"
    )
    working["month"] = working["target_date"].dt.to_period("M").dt.to_timestamp()
    working[["prediction", "actual_return"]] = working[
        ["prediction", "actual_return"]
    ].apply(pd.to_numeric, errors="coerce")
    return (
        working.groupby(["month", "model"], observed=True)
        .apply(
            lambda group: pd.Series(
                {
                    "ic": group["prediction"].corr(group["actual_return"]),
                    "observation_count": int(
                        group[["prediction", "actual_return"]]
                        .dropna()
                        .shape[0]
                    ),
                }
            ),
            include_groups=False,
        )
        .reset_index()
    )


def _build_prediction_scatter_sample(
    predictions: pd.DataFrame,
    *,
    seed: int,
    limit: int = 20_000,
) -> pd.DataFrame:
    required = ("prediction", "actual_return")
    if predictions.empty or not set(required).issubset(predictions.columns):
        return pd.DataFrame(columns=list(required))
    working = predictions.loc[:, list(required)].apply(
        pd.to_numeric, errors="coerce"
    )
    working = working.replace([np.inf, -np.inf], np.nan).dropna()
    if len(working) > limit:
        working = working.sample(n=limit, random_state=seed)
    return working


def _plot_prediction_scatter(
    predictions: pd.DataFrame,
    path: Path,
    *,
    config: M5Config,
) -> dict[str, Any]:
    sample = _build_prediction_scatter_sample(
        predictions,
        seed=config.sensitivity.seed,
    )
    if sample.empty:
        return _not_applicable_chart(
            path,
            title="预测值与实际收益散点",
            reason="无有效预测—实际收益配对，未生成伪造散点",
            config=config,
        )
    figure, axis = plt.subplots(figsize=config.reporting.bar_size)
    axis.scatter(
        sample["prediction"],
        sample["actual_return"],
        s=7,
        alpha=0.25,
        color=config.reporting.primary_blue,
    )
    axis.axhline(0, color=config.reporting.primary_red, linewidth=0.8)
    axis.axvline(0, color=config.reporting.primary_red, linewidth=0.8)
    axis.set_title("预测值与实际收益散点", fontsize=config.reporting.title_fontsize)
    axis.set_xlabel("预测收益", fontsize=config.reporting.axis_fontsize)
    axis.set_ylabel("实际收益", fontsize=config.reporting.axis_fontsize)
    _save_figure(figure, path, config=config)
    return {"file": path.name, "status": "generated", "row_count": len(sample)}


def create_result_charts(
    tables: Mapping[str, pd.DataFrame],
    output_dir: Path,
    *,
    config: M5Config = DEFAULT_M5_CONFIG,
) -> pd.DataFrame:
    """生成 ml 模板图与 FIG3–FIG19；缺表时保留明确 N/A 证据。"""
    _setup_plot_theme()
    output_dir.mkdir(parents=True, exist_ok=True)
    statuses: list[dict[str, Any]] = []

    feature_values = tables.get("feature_values", pd.DataFrame())
    if {
        "feature",
        "absolute_value",
    }.issubset(feature_values.columns) and not feature_values.empty:
        importance = (
            feature_values.groupby("feature", observed=True)["absolute_value"]
            .mean()
            .nlargest(20)
            .sort_values()
            .reset_index()
        )
    else:
        importance = pd.DataFrame(columns=["feature", "absolute_value"])
    statuses.append(
        _bar_chart(
            importance,
            output_dir / "feature_importance.png",
            category="feature",
            value="absolute_value",
            title="特征重要性（绝对值均值前20）",
            xlabel="特征",
            ylabel="绝对重要性",
            config=config,
        )
    )

    predictions = tables.get("predictions", pd.DataFrame())
    monthly_ic = tables.get("monthly_ic", _build_monthly_ic(predictions))
    statuses.append(
        _line_chart(
            monthly_ic,
            output_dir / "ic_series.png",
            x="month",
            y="ic",
            series="model",
            title="预测值月度 IC 时序",
            xlabel="月份",
            ylabel="IC",
            config=config,
        )
    )
    daily_pnl = tables.get("daily_pnl", pd.DataFrame())
    if "cumulative_pnl" in daily_pnl.columns:
        pnl_value = "cumulative_pnl"
    elif "cumulative_pnl_before_cost" in daily_pnl.columns:
        pnl_value = "cumulative_pnl_before_cost"
    else:
        pnl_value = "daily_pnl_before_cost"
    pnl_series = (
        "series_label"
        if "series_label" in daily_pnl.columns
        else ("model" if "model" in daily_pnl.columns else None)
    )
    statuses.append(
        _line_chart(
            daily_pnl,
            output_dir / "group_cumulative_returns.png",
            x="target_date",
            y=pnl_value,
            series=pnl_series,
            title="预测组合累计 PnL",
            xlabel="日期",
            ylabel="累计PnL",
            config=config,
        )
    )
    statuses.append(
        _not_applicable_chart(
            output_dir / "train_val_loss.png",
            title="训练/验证损失曲线",
            reason=(
                "本报告使用 OLS、正则回归、核回归和树模型的滚动闭式/拟合接口，"
                "无统一逐 epoch 训练/验证损失，按模板显式标记不适用"
            ),
            config=config,
        )
    )
    statuses.append(
        _plot_prediction_scatter(
            predictions,
            output_dir / "prediction_scatter.png",
            config=config,
        )
    )

    statuses.append(
        _heatmap_chart(
            tables.get("fig03_graph_matrix", pd.DataFrame()),
            output_dir / "fig03_biadjacency_heatmap.png",
            title=(
                "FIG3 2021-10-21 全量 B1 代理图 "
                "(degraded/substitute, AS27)"
            ),
            config=config,
        )
    )
    statuses.append(
        _not_applicable_chart(
            output_dir / "fig04_sector_block_heatmap.png",
            title="Figure 4 行业块绝对 t 统计量中位数",
            reason="AS4：缺少美股行业字段，F6 双市场行业块为 data_missing",
            config=config,
        )
    )
    indegree = tables.get("fig05_indegree", pd.DataFrame())
    statuses.append(
        _line_chart(
            indegree,
            output_dir / "fig05_indegree_percentiles.png",
            x="anchor_date",
            y="in_degree",
            series="series_label",
            title="Figure 5 目标入度分位数",
            xlabel="图更新日期",
            ylabel="目标入度",
            config=config,
        )
    )

    sharpe = tables.get("sharpe_metrics", pd.DataFrame())
    chart_specs = (
        (
            "fig08_us_to_cn_sharpe.png",
            "Figure 8 美国收益预测中国 OPCL 的 Sharpe",
            "fig08",
        ),
        (
            "fig09_cn_to_us_sharpe.png",
            "Figure 9 中国收益预测美国 OPCL 的 Sharpe",
            "fig09",
        ),
        (
            "fig11_same_market_sharpe.png",
            "Figure 11 中国同市场图预测 Sharpe",
            "fig11",
        ),
        (
            "fig12_three_specification_sharpe.png",
            "Figure 12 三规格 Sharpe 比较",
            "fig12",
        ),
        (
            "fig13_sharpe_delta.png",
            "Figure 13 相对非图基线 Sharpe 差",
            "fig13",
        ),
    )
    for file_name, title, key in chart_specs:
        table = tables.get(key, sharpe if key in {"fig08", "fig09"} else pd.DataFrame())
        if not table.empty and "chart_label" not in table.columns:
            table = table.copy()
            candidate_columns = [
                column
                for column in ("model", "portfolio", "specification")
                if column in table.columns
            ]
            table["chart_label"] = table[candidate_columns].astype(str).agg(
                "/".join, axis=1
            )
        value_column = (
            "sharpe_delta"
            if "sharpe_delta" in table.columns
            else "annualized_sharpe"
        )
        statuses.append(
            _bar_chart(
                table,
                output_dir / file_name,
                category="chart_label",
                value=value_column,
                title=title,
                xlabel="模型/组合/规格",
                ylabel="年化Sharpe",
                config=config,
            )
        )

    fig10 = tables.get("fig10_daily_pnl", pd.DataFrame())
    fig10_value = (
        "cumulative_pnl_before_cost"
        if "cumulative_pnl_before_cost" in fig10.columns
        else "cumulative_pnl"
    )
    statuses.append(
        _line_chart(
            fig10,
            output_dir / "fig10_cumulative_daily_pnl.png",
            x="target_date",
            y=fig10_value,
            series="series_label",
            title="Figure 10 十模型六组合累计日度 PnL",
            xlabel="日期",
            ylabel="累计PnL",
            config=config,
        )
    )
    statuses.append(
        _line_chart(
            tables.get("fig14_edge_lag", pd.DataFrame()),
            output_dir / "fig14_edge_randomization_and_lag.png",
            x="sensitivity_value",
            y="median_sharpe",
            series="series_label",
            title="Figure 14 图边随机化与信息滞后敏感性",
            xlabel="边替换比例或信息滞后",
            ylabel="十模型Sharpe中位数",
            config=config,
        )
    )
    for file_name, title, key in (
        (
            "fig15_opcl_three_specification.png",
            "Figure B.1 OPCL 特征三规格 Sharpe",
            "sa4_opcl",
        ),
        (
            "fig16_opcl_sharpe_delta.png",
            "Figure B.2 OPCL 特征相对非图 Sharpe 差",
            "sa4_opcl_delta",
        ),
        (
            "fig17_us_shock_conditioned.png",
            "Figure C.1 美国冲击条件 Sharpe",
            "sa5_shock",
        ),
        (
            "fig18_sector_sharpe.png",
            "Figure D.1 A股中信一级行业 Sharpe",
            "sa6_industry",
        ),
        (
            "fig19_joint_predictors.png",
            "Figure E.1 美中联合预测变量 Sharpe",
            "sa7_joint",
        ),
    ):
        table = tables.get(key, pd.DataFrame())
        if not table.empty and "chart_label" not in table.columns:
            table = table.copy()
            candidates = [
                column
                for column in (
                    "model",
                    "portfolio",
                    "shock_quantile",
                    "target_industry",
                    "feature_combination",
                    "specification",
                )
                if column in table.columns
            ]
            table["chart_label"] = table[candidates].astype(str).agg(
                "/".join, axis=1
            )
        value_column = (
            "sharpe_delta"
            if "sharpe_delta" in table.columns
            else "annualized_sharpe"
        )
        statuses.append(
            _bar_chart(
                table,
                output_dir / file_name,
                category="chart_label",
                value=value_column,
                title=title,
                xlabel="模型/组合/规格",
                ylabel="年化Sharpe",
                config=config,
            )
        )
    return pd.DataFrame(statuses)


def _excel_safe_frame(frame: pd.DataFrame) -> pd.DataFrame:
    result = frame.copy()
    for column in result.columns:
        if isinstance(result[column].dtype, pd.DatetimeTZDtype):
            result[column] = result[column].astype(str)
        elif result[column].dtype == "object":
            result[column] = result[column].map(
                lambda value: (
                    json.dumps(_jsonable(value), ensure_ascii=False, sort_keys=True)
                    if isinstance(value, (dict, list, tuple, set))
                    else value
                )
            )
    return result


def _sheet_name(value: str, used: set[str]) -> str:
    sanitized = "".join(
        "_" if character in r"[]:*?/\\" else character for character in value
    )[:31] or "Sheet"
    candidate = sanitized
    counter = 1
    while candidate in used:
        suffix = f"_{counter}"
        candidate = f"{sanitized[:31-len(suffix)]}{suffix}"
        counter += 1
    used.add(candidate)
    return candidate


def _format_workbook(
    path: Path,
    *,
    chart_by_sheet: Mapping[str, Path] | None = None,
) -> None:
    workbook = load_workbook(path)
    for worksheet in workbook.worksheets:
        worksheet.freeze_panes = "A2"
        for cell in worksheet[1]:
            cell.font = Font(bold=True)
        for index, column_cells in enumerate(worksheet.columns, start=1):
            maximum = max(
                (
                    len(str(cell.value))
                    for cell in column_cells[: min(2000, len(column_cells))]
                    if cell.value is not None
                ),
                default=0,
            )
            worksheet.column_dimensions[get_column_letter(index)].width = min(
                max(maximum + 2, 10),
                48,
            )
        if chart_by_sheet and worksheet.title in chart_by_sheet:
            chart_path = chart_by_sheet[worksheet.title]
            if chart_path.is_file():
                image = ExcelImage(chart_path)
                image.width = min(image.width, 900)
                image.height = min(image.height, 450)
                worksheet.add_image(image, "J2")
    workbook.save(path)


def write_table_workbook(
    path: Path,
    tables: Mapping[str, pd.DataFrame],
    *,
    chart_by_table: Mapping[str, Path] | None = None,
    config: M5Config = DEFAULT_M5_CONFIG,
) -> None:
    """多 sheet 写表，超 Excel 单 sheet 上限时确定性拆分且不静默截断。"""
    used: set[str] = set()
    chart_by_sheet: dict[str, Path] = {}
    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        if not tables:
            pd.DataFrame(
                [{"status": "not_applicable", "reason": "本工作簿无对应结果表"}]
            ).to_excel(writer, sheet_name="说明", index=False)
        for table_name, frame in tables.items():
            safe = _excel_safe_frame(frame)
            chunk_size = config.reporting.excel_max_rows_per_sheet
            chunks = max(1, int(np.ceil(len(safe) / chunk_size)))
            for index in range(chunks):
                suffix = f"_{index + 1}" if chunks > 1 else ""
                sheet = _sheet_name(f"{table_name}{suffix}", used)
                start = index * chunk_size
                stop = min((index + 1) * chunk_size, len(safe))
                safe.iloc[start:stop].to_excel(
                    writer,
                    sheet_name=sheet,
                    index=False,
                )
                if (
                    index == 0
                    and chart_by_table
                    and table_name in chart_by_table
                ):
                    chart_by_sheet[sheet] = chart_by_table[table_name]
    _format_workbook(path, chart_by_sheet=chart_by_sheet)


def _file_sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _build_manifest(files: Sequence[Path], root: Path) -> pd.DataFrame:
    return pd.DataFrame(
        [
            {
                "relative_path": str(path.relative_to(root)),
                "bytes": path.stat().st_size,
                "sha256": _file_sha256(path),
            }
            for path in sorted(set(files))
            if path.is_file()
        ]
    )


def write_formal_outputs(
    tables: Mapping[str, pd.DataFrame],
    *,
    metrics: Mapping[str, Any],
    comparison: Mapping[str, Any],
    run_log_lines: Sequence[str],
    output_dir: Path | None = None,
    config: M5Config = DEFAULT_M5_CONFIG,
) -> ReportingResult:
    """一次写齐正式 results 输出合同；调用方负责仅在正式模式调用。"""
    destination = output_dir or config.reporting.output_dir
    destination.mkdir(parents=True, exist_ok=True)
    chart_status = create_result_charts(tables, destination, config=config)
    chart_files = tuple(destination / name for name in chart_status["file"])

    workbook_specs = {
        "backtest_summary.xlsx": {
            "核心对照": tables.get("core_comparison", pd.DataFrame()),
            "Sharpe汇总": tables.get("sharpe_metrics", pd.DataFrame()),
            "R2逐格对照": tables.get("r2_comparison", pd.DataFrame()),
            "敏感性方向证据": tables.get(
                "sensitivity_directional", pd.DataFrame()
            ),
        },
        "ml_core_results.xlsx": {
            "预测质量": tables.get("prediction_quality", pd.DataFrame()),
            "月度IC": tables.get("monthly_ic", _build_monthly_ic(
                tables.get("predictions", pd.DataFrame())
            )),
            "组合日度PnL": tables.get("daily_pnl", pd.DataFrame()),
            "模型特征值": tables.get("feature_values", pd.DataFrame()),
            "模型配置": tables.get("model_configuration", pd.DataFrame()),
            "FIG3全量B1代理_degraded": tables.get(
                "fig03_graph_matrix_table", pd.DataFrame()
            ),
            "FIG3代理口径_degraded": tables.get(
                "fig03_graph_anchor_audit", pd.DataFrame()
            ),
            "FIG5入度分位底表": tables.get(
                "fig05_indegree", pd.DataFrame()
            ),
        },
        "sensitivity_analysis.xlsx": {
            "SA1边替换": tables.get("sa1_edge", pd.DataFrame()),
            "SA1边守恒审计": tables.get("sa1_edge_audit", pd.DataFrame()),
            "SA2信息滞后": tables.get("sa2_lag", pd.DataFrame()),
            "SA2时点审计": tables.get("sa2_lag_audit", pd.DataFrame()),
            "FIG14边替换与滞后": tables.get(
                "fig14_edge_lag", pd.DataFrame()
            ),
            "SA3超参数逐配置": tables.get("sa3_runs", pd.DataFrame()),
            "SA3逐配置模型审计": tables.get(
                "sa3_model_audit", pd.DataFrame()
            ),
            "SA3共同样本门禁": tables.get(
                "sa3_grid_audit", pd.DataFrame()
            ),
            "TBL1超参数汇总": tables.get("r2_comparison", pd.DataFrame()),
            "SA4_OPCL": tables.get("sa4_opcl", pd.DataFrame()),
            "SA4差值": tables.get("sa4_opcl_delta", pd.DataFrame()),
            "SA5冲击条件": tables.get("sa5_shock", pd.DataFrame()),
            "SA6行业": tables.get("sa6_industry", pd.DataFrame()),
            "SA7联合特征": tables.get("sa7_joint", pd.DataFrame()),
            "SA7基线配对": tables.get("sa7_comparison", pd.DataFrame()),
        },
        "evidence_tables.xlsx": {
            name[:31]: frame
            for name, frame in tables.items()
            if name.endswith("_audit")
            or name
            in {
                "graph_audit",
                "training_audit",
                "model_audit",
                "model_completeness_audit",
                "comparability_audit",
                "execution_timing_audit",
            }
        },
    }
    workbook_files: list[Path] = []
    workbook_chart_maps: dict[str, dict[str, Path]] = {
        "backtest_summary.xlsx": {
            "核心对照": destination / "fig12_three_specification_sharpe.png",
            "Sharpe汇总": destination / "fig08_us_to_cn_sharpe.png",
        },
        "ml_core_results.xlsx": {
            "月度IC": destination / "ic_series.png",
            "组合日度PnL": destination / "group_cumulative_returns.png",
            "模型特征值": destination / "feature_importance.png",
            "FIG3全量B1代理_degraded": (
                destination / "fig03_biadjacency_heatmap.png"
            ),
            "FIG5入度分位底表": (
                destination / "fig05_indegree_percentiles.png"
            ),
        },
        "sensitivity_analysis.xlsx": {
            "FIG14边替换与滞后": (
                destination / "fig14_edge_randomization_and_lag.png"
            ),
            "SA4_OPCL": destination / "fig15_opcl_three_specification.png",
            "SA4差值": destination / "fig16_opcl_sharpe_delta.png",
            "SA5冲击条件": destination / "fig17_us_shock_conditioned.png",
            "SA6行业": destination / "fig18_sector_sharpe.png",
            "SA7联合特征": destination / "fig19_joint_predictors.png",
        },
    }
    for file_name, workbook_tables in workbook_specs.items():
        path = destination / file_name
        write_table_workbook(
            path,
            workbook_tables,
            chart_by_table=workbook_chart_maps.get(file_name),
            config=config,
        )
        workbook_files.append(path)

    metrics_path = destination / "metrics.json"
    comparison_path = destination / "comparison.json"
    run_log_path = destination / "run_log.md"
    enriched_metrics = {
        **dict(metrics),
        "output_contract": {
            "train_val_loss": {
                "status": "not_applicable",
                "reason": (
                    "传统滚动回归与树模型无统一逐 epoch 训练/验证损失"
                ),
                "evidence_file": "train_val_loss.png",
            },
            "chart_status": chart_status.to_dict(orient="records"),
        },
    }
    _write_json(metrics_path, enriched_metrics)
    _write_json(comparison_path, comparison)
    run_log_path.write_text(
        "# r007 正式运行日志\n\n"
        + "\n".join(f"- {line}" for line in run_log_lines)
        + "\n",
        encoding="utf-8",
    )
    files = (
        *chart_files,
        *workbook_files,
        metrics_path,
        comparison_path,
        run_log_path,
    )
    manifest = _build_manifest(files, destination)
    manifest_path = destination / "artifact_manifest.json"
    _write_json(
        manifest_path,
        {
            "report_name": "r007_cross_market_return_forecasting",
            "files": manifest.to_dict(orient="records"),
        },
    )
    final_files = (*files, manifest_path)
    return ReportingResult(
        files=tuple(final_files),
        chart_status=chart_status,
        manifest=_build_manifest(final_files, destination),
    )
