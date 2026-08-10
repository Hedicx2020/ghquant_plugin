"""r007 无参正式流水线与 ``--smoke`` 极小样本入口。

正式模式按 m1→m2/m3→m4→m5→reporting 顺序执行，并把可续跑检查点写入
workspace；只有 reporting 阶段写 ``output/.../results``。冒烟模式完全
使用合成小样本，产物仅写 ``workspace/.../smoke``。
"""

from __future__ import annotations

import argparse
import gzip
import hashlib
import json
import os
import pickle
import shutil
import warnings
from dataclasses import asdict, dataclass, is_dataclass, replace
from datetime import datetime
from pathlib import Path
from typing import Any, Mapping, Sequence

import numpy as np
import pandas as pd
import pyarrow as pa
import pyarrow.parquet as pq

from .backtest_pipeline import (
    build_same_market_calendar_mapping,
    iter_non_graph_autoregressive_predictions,
    iter_same_market_graph_predictions,
)
from .config import (
    DEFAULT_M5_CONFIG,
    DirectionConfig,
    M5Config,
)
from .data_pipeline import (
    CN_PRICE_REQUIRED_COLUMNS,
    CN_TRADE_REQUIRED_COLUMNS,
    CN_UNIT_CONTRACT,
    US_SOURCE_REQUIRED_COLUMNS,
    US_UNIT_CONTRACT,
    audit_dual_market_return_panels,
    audit_loaded_market_panel,
    audit_m1_source_schemas,
    audit_return_distributions,
    audit_security_calendar_coverage,
    audit_us_merged_panel_contract,
    build_bidirectional_calendar_mappings,
    build_dual_market_return_panels,
    build_dual_market_universes,
    calculate_market_log_returns,
    load_dual_market_sources,
)
from .graph_pipeline import (
    GraphSnapshot,
    audit_snapshot_causality,
    build_update_blocks,
    iter_bidirectional_rolling_graphs,
    iter_rolling_graph_snapshots,
    select_training_mapping,
)
from .model_pipeline import (
    SnapshotPredictionResult,
    predict_graph_snapshot,
    summarize_prediction_quality,
    validate_eligibility_audit,
)
from .portfolio_pipeline import (
    build_core_comparison_tables,
    prepare_portfolio_predictions,
    run_m4_portfolios,
)
from .reporting import build_figure13_long_table, write_formal_outputs
from .sensitivity_pipeline import (
    attach_nested_us_shock_quantiles,
    attach_target_industry,
    build_directional_sensitivity_comparison,
    build_hyperparameter_factory,
    build_hyperparameter_grid,
    build_information_lag_mapping,
    build_r2_comparison,
    build_sa7_us_only_comparison,
    load_citic_level_one_industry,
    predict_joint_graph_snapshots,
    predict_snapshot_single_model,
    randomize_graph_snapshot,
    run_cn_three_specification_portfolios,
    run_sensitivity_portfolios,
    summarize_hyperparameter_sharpe,
    summarize_model_median_sharpe,
)
from .strategy import (
    build_base_model_factories,
    build_model_configuration_evidence,
)


@dataclass
class M1Tables:
    us_panel: pd.DataFrame
    cn_panel: pd.DataFrame
    calendar_mapping: pd.DataFrame
    us_calendar: pd.DataFrame
    cn_calendar: pd.DataFrame
    us_market_returns: pd.DataFrame
    universe_audit: pd.DataFrame
    source_schema_audit: pd.DataFrame
    loaded_panel_audit: pd.DataFrame
    return_panel_audit: pd.DataFrame
    security_coverage_audit: pd.DataFrame
    return_distribution_audit: pd.DataFrame


@dataclass
class M3Tables:
    predictions: pd.DataFrame
    training_audit: pd.DataFrame
    model_audit: pd.DataFrame
    feature_values: pd.DataFrame
    graph_audit: pd.DataFrame
    indegree_quantiles: pd.DataFrame
    eligibility_audit: pd.DataFrame


@dataclass
class SpecificationPredictionTables:
    """单一规格的组合输入及独立动态可预测机会审计。"""

    predictions: pd.DataFrame
    eligibility_audit: pd.DataFrame


@dataclass
class M4Tables:
    prepared_predictions: pd.DataFrame
    sharpe_metrics: pd.DataFrame
    daily_pnl: pd.DataFrame
    portfolio_members: pd.DataFrame
    nested_audit: pd.DataFrame
    model_completeness_audit: pd.DataFrame
    comparability_audit: pd.DataFrame
    execution_timing_audit: pd.DataFrame
    core_observations: pd.DataFrame
    core_paired_comparisons: pd.DataFrame
    core_summary: pd.DataFrame
    formal_coverage_audit: pd.DataFrame


@dataclass(frozen=True)
class CheckpointValidation:
    """检查点复用判定及明确失效原因。"""

    reusable: bool
    reason: str
    expected_context_hash: str
    observed_context_hash: str | None


def _checkpoint_jsonable(value: Any) -> Any:
    if is_dataclass(value) and not isinstance(value, type):
        return _checkpoint_jsonable(asdict(value))
    if isinstance(value, dict):
        return {
            str(key): _checkpoint_jsonable(item)
            for key, item in sorted(value.items(), key=lambda item: str(item[0]))
        }
    if isinstance(value, (list, tuple, set)):
        return [_checkpoint_jsonable(item) for item in value]
    if isinstance(value, Path):
        return str(value.resolve())
    if isinstance(value, (pd.Timestamp, datetime)):
        return value.isoformat()
    if isinstance(value, np.generic):
        return value.item()
    return value


def _json_sha256(payload: Any) -> str:
    encoded = json.dumps(
        _checkpoint_jsonable(payload),
        ensure_ascii=False,
        sort_keys=True,
        separators=(",", ":"),
        default=str,
    ).encode("utf-8")
    return hashlib.sha256(encoded).hexdigest()


def _file_sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _frame_schema_payload(frame: pd.DataFrame) -> list[dict[str, str]]:
    return [
        {"column": str(column), "dtype": str(frame[column].dtype)}
        for column in frame.columns
    ]


_CHECKPOINT_TYPED_VALUE_METADATA_KEY = (
    b"r007.checkpoint.typed_json_columns"
)
_CHECKPOINT_TYPED_VALUE_PROTOCOL = "r007.checkpoint.typed-json"
_CHECKPOINT_TYPED_VALUE_VERSION = 1
_CHECKPOINT_COMPLEX_VALUE_TAGS = frozenset(
    {"timestamp", "datetime", "list", "dict", "pd_na", "nat"}
)


def _checkpoint_typed_value_node(value: Any) -> dict[str, Any]:
    """把审计单元格编码成无歧义、可递归且可稳定排序的类型节点。"""
    if value is None:
        return {"t": "none"}
    if value is pd.NA:
        return {"t": "pd_na"}
    if value is pd.NaT:
        return {"t": "nat"}
    if isinstance(value, np.generic):
        return _checkpoint_typed_value_node(value.item())
    if isinstance(value, bool):
        return {"t": "bool", "v": value}
    if isinstance(value, int):
        return {"t": "int", "v": str(value)}
    if isinstance(value, float):
        if np.isnan(value):
            return {"t": "nan"}
        return {"t": "float", "v": value.hex()}
    if isinstance(value, str):
        return {"t": "str", "v": value}
    if isinstance(value, pd.Timestamp):
        if pd.isna(value):
            return {"t": "nat"}
        return {"t": "timestamp", "v": value.isoformat()}
    if isinstance(value, datetime):
        return {"t": "datetime", "v": value.isoformat()}
    if isinstance(value, list):
        return {
            "t": "list",
            "v": [_checkpoint_typed_value_node(item) for item in value],
        }
    if isinstance(value, dict):
        pairs = [
            [
                _checkpoint_typed_value_node(key),
                _checkpoint_typed_value_node(item),
            ]
            for key, item in value.items()
        ]
        pairs.sort(
            key=lambda pair: json.dumps(
                pair[0],
                ensure_ascii=False,
                sort_keys=True,
                separators=(",", ":"),
            )
        )
        return {"t": "dict", "v": pairs}
    raise TypeError(
        "检查点 typed JSON 不支持单元格类型: "
        f"{type(value).__module__}.{type(value).__qualname__}"
    )


def _checkpoint_typed_value_from_node(node: Any) -> Any:
    """严格解码类型节点；未知 tag、字段或损坏结构立即失败。"""
    if not isinstance(node, dict) or not isinstance(node.get("t"), str):
        raise ValueError("检查点 typed JSON 单元格不是合法类型节点")
    tag = node["t"]
    no_value_tags = {"none", "pd_na", "nat", "nan"}
    value_tags = {
        "bool",
        "int",
        "float",
        "str",
        "timestamp",
        "datetime",
        "list",
        "dict",
    }
    if tag in no_value_tags:
        if set(node) != {"t"}:
            raise ValueError(f"检查点 typed JSON {tag} 节点字段非法")
        return {
            "none": None,
            "pd_na": pd.NA,
            "nat": pd.NaT,
            "nan": np.nan,
        }[tag]
    if tag not in value_tags:
        raise ValueError(f"检查点 typed JSON 存在未知类型 tag: {tag}")
    if set(node) != {"t", "v"}:
        raise ValueError(f"检查点 typed JSON {tag} 节点字段非法")
    value = node["v"]
    if tag == "bool":
        if not isinstance(value, bool):
            raise ValueError("检查点 typed JSON bool 值非法")
        return value
    if tag == "int":
        if not isinstance(value, str):
            raise ValueError("检查点 typed JSON int 值非法")
        try:
            return int(value)
        except ValueError as exc:
            raise ValueError("检查点 typed JSON int 值非法") from exc
    if tag == "float":
        if not isinstance(value, str):
            raise ValueError("检查点 typed JSON float 值非法")
        try:
            return float.fromhex(value)
        except ValueError as exc:
            raise ValueError("检查点 typed JSON float 值非法") from exc
    if tag == "str":
        if not isinstance(value, str):
            raise ValueError("检查点 typed JSON str 值非法")
        return value
    if tag == "timestamp":
        if not isinstance(value, str):
            raise ValueError("检查点 typed JSON timestamp 值非法")
        try:
            return pd.Timestamp(value)
        except (TypeError, ValueError) as exc:
            raise ValueError("检查点 typed JSON timestamp 值非法") from exc
    if tag == "datetime":
        if not isinstance(value, str):
            raise ValueError("检查点 typed JSON datetime 值非法")
        try:
            return datetime.fromisoformat(value)
        except ValueError as exc:
            raise ValueError("检查点 typed JSON datetime 值非法") from exc
    if tag == "list":
        if not isinstance(value, list):
            raise ValueError("检查点 typed JSON list 值非法")
        return [_checkpoint_typed_value_from_node(item) for item in value]
    if not isinstance(value, list):
        raise ValueError("检查点 typed JSON dict 值非法")
    result: dict[Any, Any] = {}
    for pair in value:
        if not isinstance(pair, list) or len(pair) != 2:
            raise ValueError("检查点 typed JSON dict 键值对非法")
        key = _checkpoint_typed_value_from_node(pair[0])
        item = _checkpoint_typed_value_from_node(pair[1])
        try:
            if key in result:
                raise ValueError("检查点 typed JSON dict 存在重复键")
            result[key] = item
        except TypeError as exc:
            raise ValueError("检查点 typed JSON dict 键不可哈希") from exc
    return result


def _encode_checkpoint_typed_value(value: Any) -> str:
    return json.dumps(
        _checkpoint_typed_value_node(value),
        ensure_ascii=False,
        sort_keys=True,
        separators=(",", ":"),
    )


def _decode_checkpoint_typed_value(payload: Any) -> Any:
    if not isinstance(payload, str):
        raise ValueError("检查点 typed JSON 编码单元格必须是字符串")
    try:
        node = json.loads(payload)
    except json.JSONDecodeError as exc:
        raise ValueError("检查点 typed JSON 单元格 JSON 损坏") from exc
    decoded = _checkpoint_typed_value_from_node(node)
    canonical = _encode_checkpoint_typed_value(decoded)
    if canonical != payload:
        raise ValueError("检查点 typed JSON 单元格不是 canonical 编码")
    return decoded


def _checkpoint_columns_requiring_typed_encoding(
    frame: pd.DataFrame,
) -> list[str]:
    """仅选择存在异构或嵌套值的 object 列，并预先拒绝未知类型。"""
    columns: list[str] = []
    for column in frame.columns:
        series = frame[column]
        if not pd.api.types.is_object_dtype(series.dtype):
            continue
        tags: set[str] = set()
        for row_index, value in series.items():
            try:
                tags.add(str(_checkpoint_typed_value_node(value)["t"]))
            except TypeError as exc:
                raise TypeError(
                    f"检查点列 {column!r} 行 {row_index!r} 含未知类型"
                ) from exc
        if len(tags) > 1 or bool(tags & _CHECKPOINT_COMPLEX_VALUE_TAGS):
            if not isinstance(column, str):
                raise TypeError("typed JSON 检查点列名必须是字符串")
            columns.append(column)
    return columns


def _typed_checkpoint_metadata_payload(
    columns: Sequence[str],
) -> bytes:
    return json.dumps(
        {
            "protocol": _CHECKPOINT_TYPED_VALUE_PROTOCOL,
            "version": _CHECKPOINT_TYPED_VALUE_VERSION,
            "columns": [
                {"name": column, "decoded_dtype": "object"}
                for column in columns
            ],
        },
        ensure_ascii=False,
        sort_keys=True,
        separators=(",", ":"),
    ).encode("utf-8")


def _typed_checkpoint_columns_from_schema(
    schema: pa.Schema,
) -> list[str]:
    metadata = schema.metadata or {}
    raw = metadata.get(_CHECKPOINT_TYPED_VALUE_METADATA_KEY)
    if raw is None:
        return []
    try:
        payload = json.loads(raw.decode("utf-8"))
    except (UnicodeDecodeError, json.JSONDecodeError) as exc:
        raise ValueError("检查点 typed JSON schema metadata 损坏") from exc
    if not isinstance(payload, dict):
        raise ValueError("检查点 typed JSON schema metadata 必须是对象")
    if (
        payload.get("protocol") != _CHECKPOINT_TYPED_VALUE_PROTOCOL
        or type(payload.get("version")) is not int
        or payload["version"] != _CHECKPOINT_TYPED_VALUE_VERSION
        or set(payload) != {"protocol", "version", "columns"}
        or not isinstance(payload.get("columns"), list)
    ):
        raise ValueError("检查点 typed JSON schema metadata 协议或版本非法")
    columns: list[str] = []
    for entry in payload["columns"]:
        if (
            not isinstance(entry, dict)
            or set(entry) != {"name", "decoded_dtype"}
            or not isinstance(entry.get("name"), str)
            or entry.get("decoded_dtype") != "object"
        ):
            raise ValueError("检查点 typed JSON schema metadata 列定义非法")
        column = entry["name"]
        if column in columns:
            raise ValueError("检查点 typed JSON schema metadata 含重复列")
        if column not in schema.names:
            raise ValueError(
                f"检查点 typed JSON metadata 指向不存在列: {column}"
            )
        encoded_type = schema.field(column).type
        if not (
            pa.types.is_string(encoded_type)
            or pa.types.is_large_string(encoded_type)
        ):
            raise ValueError(
                f"检查点 typed JSON 编码列不是 Arrow string: {column}"
            )
        columns.append(column)
    if not columns:
        raise ValueError("检查点 typed JSON schema metadata 列清单为空")
    return columns


def _read_checkpoint_frame_path(path: Path) -> pd.DataFrame:
    """读取并自动解码 typed JSON；无协议 metadata 的旧 parquet 原样兼容。"""
    schema = pq.ParquetFile(path).schema_arrow
    encoded_columns = _typed_checkpoint_columns_from_schema(schema)
    frame = pd.read_parquet(path)
    for column in encoded_columns:
        decoded: list[Any] = []
        for row_index, payload in frame[column].items():
            try:
                decoded.append(_decode_checkpoint_typed_value(payload))
            except ValueError as exc:
                raise ValueError(
                    f"检查点列 {column!r} 行 {row_index!r} typed JSON 解码失败"
                ) from exc
        frame[column] = pd.Series(decoded, index=frame.index, dtype=object)
    return frame


def _write_checkpoint_frame_path(
    path: Path,
    frame: pd.DataFrame,
) -> list[str]:
    """按需编码混合 object 列并把协议直接嵌入 parquet schema。"""
    encoded_columns = _checkpoint_columns_requiring_typed_encoding(frame)
    if not encoded_columns:
        frame.to_parquet(path, index=False)
        return []
    encoded = frame.copy()
    for column in encoded_columns:
        values: list[str] = []
        for row_index, value in frame[column].items():
            try:
                values.append(_encode_checkpoint_typed_value(value))
            except TypeError as exc:
                raise TypeError(
                    f"检查点列 {column!r} 行 {row_index!r} 编码失败"
                ) from exc
        encoded[column] = pd.Series(
            values,
            index=encoded.index,
            dtype="string",
        )
    table = pa.Table.from_pandas(encoded, preserve_index=False)
    metadata = dict(table.schema.metadata or {})
    metadata[_CHECKPOINT_TYPED_VALUE_METADATA_KEY] = (
        _typed_checkpoint_metadata_payload(encoded_columns)
    )
    pq.write_table(table.replace_schema_metadata(metadata), path)
    return encoded_columns


def _lossless_object_dtype_roundtrip_columns(
    original: pd.DataFrame,
    reloaded: pd.DataFrame,
    *,
    encoded_columns: Sequence[str],
) -> list[str] | None:
    """识别 Parquet 对同构 object 标量列的无损原生 dtype 推断。

    pandas 的 ``groupby.apply(...).unstack()`` 可能因一行内同时含数值与字符串，
    把展开后的每一列都保留为 object。Parquet 会按单列真实值把这些列重开为
    int64/float64/string。这里只放行单一 bool/int/float/str 类型且 canonical
    类型和值逐单元格完全一致的转换；混合、缺失、嵌套及特殊值仍由 typed JSON
    协议处理或按 schema 不一致失败。
    """
    if (
        list(original.columns) != list(reloaded.columns)
        or original.shape != reloaded.shape
    ):
        return None
    encoded = set(encoded_columns)
    normalized: list[str] = []
    for column in original.columns:
        original_dtype = original[column].dtype
        reloaded_dtype = reloaded[column].dtype
        if str(original_dtype) == str(reloaded_dtype):
            continue
        if (
            column in encoded
            or not pd.api.types.is_object_dtype(original_dtype)
        ):
            return None
        tags = {
            str(_checkpoint_typed_value_node(value)["t"])
            for value in original[column]
        }
        if len(tags) != 1:
            return None
        tag = next(iter(tags))
        dtype_compatible = {
            "bool": pd.api.types.is_bool_dtype(reloaded_dtype),
            "int": pd.api.types.is_integer_dtype(reloaded_dtype),
            "float": pd.api.types.is_float_dtype(reloaded_dtype),
            "str": pd.api.types.is_string_dtype(reloaded_dtype),
        }.get(tag, False)
        if not dtype_compatible:
            return None
        original_values = [
            _encode_checkpoint_typed_value(value)
            for value in original[column]
        ]
        reloaded_values = [
            _encode_checkpoint_typed_value(value)
            for value in reloaded[column]
        ]
        if original_values != reloaded_values:
            return None
        normalized.append(str(column))
    return normalized


def _path_identity_manifest(path: Path) -> dict[str, Any]:
    """以流式完整正文 SHA256 与 schema 绑定正式原始输入身份。"""
    resolved = path.resolve()
    if not resolved.is_file():
        return {
            "path": str(resolved),
            "exists": False,
            "identity_sha256": _json_sha256(
                {"path": str(resolved), "exists": False}
            ),
        }
    stat = resolved.stat()
    schema_hash: str | None = None
    if resolved.suffix == ".parquet":
        import pyarrow.parquet as pq

        schema_hash = hashlib.sha256(
            str(pq.ParquetFile(resolved).schema_arrow).encode("utf-8")
        ).hexdigest()
    content_hash = _file_sha256(resolved)
    identity = {
        "path": str(resolved),
        "exists": True,
        "bytes": stat.st_size,
        "mtime_ns": stat.st_mtime_ns,
        "schema_sha256": schema_hash,
        "content_sha256": content_hash,
    }
    return {
        **identity,
        "identity_sha256": _json_sha256(identity),
    }


class CheckpointStore:
    """绑定代码/config/输入/上游哈希的原子检查点，不触碰 state.json。"""

    def __init__(self, root: Path) -> None:
        self.root = root
        self.frames_dir = root / "frames"
        self.snapshots_dir = root / "snapshots"
        self.meta_dir = root / "meta"
        self.last_validation: dict[str, CheckpointValidation] = {}
        for path in (self.frames_dir, self.snapshots_dir, self.meta_dir):
            path.mkdir(parents=True, exist_ok=True)

    def frame_path(self, name: str) -> Path:
        return self.frames_dir / f"{name}.parquet"

    def has_frame(self, name: str) -> bool:
        return self.frame_path(name).is_file()

    def write_frame(self, name: str, frame: pd.DataFrame) -> Path:
        path = self.frame_path(name)
        temporary = path.with_name(f"{path.name}.part")
        try:
            encoded_columns = _write_checkpoint_frame_path(temporary, frame)
            reloaded = _read_checkpoint_frame_path(temporary)
        except Exception:
            temporary.unlink(missing_ok=True)
            raise
        expected_schema = _frame_schema_payload(frame)
        observed_schema = _frame_schema_payload(reloaded)
        normalized_object_columns: list[str] = []
        if observed_schema != expected_schema:
            normalized = _lossless_object_dtype_roundtrip_columns(
                frame,
                reloaded,
                encoded_columns=encoded_columns,
            )
            if normalized is None:
                temporary.unlink(missing_ok=True)
                raise ValueError(
                    f"检查点写后 schema 校验失败: {name}; "
                    f"expected={expected_schema}; "
                    f"observed={observed_schema}"
                )
            normalized_object_columns = normalized
        if (
            observed_schema != expected_schema
            and not normalized_object_columns
        ):
            temporary.unlink(missing_ok=True)
            raise ValueError(
                f"检查点写后 schema 校验失败: {name}; "
                f"expected={expected_schema}; "
                f"observed={observed_schema}"
            )
        for column in encoded_columns:
            original_values = frame[column].map(
                _encode_checkpoint_typed_value
            )
            observed_values = reloaded[column].map(
                _encode_checkpoint_typed_value
            )
            if not original_values.equals(observed_values):
                temporary.unlink(missing_ok=True)
                raise ValueError(
                    f"检查点写后 typed JSON 值校验失败: {name}.{column}"
                )
        os.replace(temporary, path)
        return path

    def read_frame(self, name: str) -> pd.DataFrame:
        path = self.frame_path(name)
        if not path.is_file():
            raise FileNotFoundError(f"检查点表不存在: {path}")
        return _read_checkpoint_frame_path(path)

    def stage_marker(self, stage: str) -> Path:
        return self.meta_dir / f"{stage}.json"

    def marker_hash(self, stage: str) -> str | None:
        path = self.stage_marker(stage)
        return _file_sha256(path) if path.is_file() else None

    def code_manifest(self) -> dict[str, Any]:
        project_root = Path(__file__).resolve().parents[2]
        paths = sorted(
            [
                *(
                    project_root / "src" / "r007_cross_market_return_forecasting"
                ).glob("*.py"),
                project_root / "common" / "ml_pipeline.py",
                project_root / "common" / "ml_portfolio.py",
                project_root / "common" / "ml_reporting.py",
                project_root / "common" / "data_loader.py",
                project_root / "common" / "backtest.py",
                project_root / "common" / "utils.py",
                project_root / "pyproject.toml",
                project_root / "uv.lock",
            ]
        )
        entries = [
            {
                "path": str(path.relative_to(project_root)),
                "sha256": _file_sha256(path),
            }
            for path in paths
            if path.is_file()
        ]
        return {"files": entries, "aggregate_sha256": _json_sha256(entries)}

    def build_context(
        self,
        stage: str,
        *,
        config: Any,
        input_manifest: Mapping[str, Any],
        upstream_stages: tuple[str, ...] = (),
        code_hash_override: str | None = None,
    ) -> dict[str, Any]:
        code_manifest = self.code_manifest()
        if code_hash_override is not None:
            code_manifest = {
                "files": [],
                "aggregate_sha256": code_hash_override,
            }
        context = {
            "stage": stage,
            "code_manifest": code_manifest,
            "config_sha256": _json_sha256(config),
            "input_manifest": _checkpoint_jsonable(dict(input_manifest)),
            "input_manifest_sha256": _json_sha256(input_manifest),
            "upstream_checkpoint_hashes": {
                upstream: self.marker_hash(upstream)
                for upstream in upstream_stages
            },
        }
        return {
            **context,
            "context_sha256": _json_sha256(context),
        }

    def _output_manifest(self, paths: Sequence[Path]) -> list[dict[str, Any]]:
        entries: list[dict[str, Any]] = []
        for path in sorted(set(path.resolve() for path in paths)):
            if not path.is_file():
                raise FileNotFoundError(f"检查点输出不存在: {path}")
            schema_sha256 = None
            if path.suffix == ".parquet":
                schema_sha256 = _json_sha256(
                    _frame_schema_payload(pd.read_parquet(path))
                )
            entries.append(
                {
                    "path": str(path),
                    "bytes": path.stat().st_size,
                    "sha256": _file_sha256(path),
                    "schema_sha256": schema_sha256,
                }
            )
        return entries

    def validate_stage(
        self,
        stage: str,
        expected_context: Mapping[str, Any],
    ) -> CheckpointValidation:
        path = self.stage_marker(stage)
        if not path.is_file():
            result = CheckpointValidation(
                False,
                "marker_missing",
                str(expected_context["context_sha256"]),
                None,
            )
            self.last_validation[stage] = result
            return result
        try:
            marker = json.loads(path.read_text(encoding="utf-8"))
        except (json.JSONDecodeError, OSError) as exc:
            result = CheckpointValidation(
                False,
                f"marker_unreadable:{type(exc).__name__}",
                str(expected_context["context_sha256"]),
                None,
            )
            self.last_validation[stage] = result
            return result
        observed_hash = marker.get("context", {}).get("context_sha256")
        if marker.get("status") != "complete":
            reason = "marker_not_complete"
        elif observed_hash != expected_context["context_sha256"]:
            reason = "context_hash_mismatch"
        else:
            reason = "reusable"
            for entry in marker.get("output_manifest", []):
                output = Path(entry["path"])
                if not output.is_file():
                    reason = "output_missing"
                    break
                if _file_sha256(output) != entry.get("sha256"):
                    reason = "output_hash_mismatch"
                    break
                if output.suffix == ".parquet":
                    actual_schema = _json_sha256(
                        _frame_schema_payload(pd.read_parquet(output))
                    )
                    if actual_schema != entry.get("schema_sha256"):
                        reason = "output_schema_mismatch"
                        break
        result = CheckpointValidation(
            reason == "reusable",
            reason,
            str(expected_context["context_sha256"]),
            str(observed_hash) if observed_hash is not None else None,
        )
        self.last_validation[stage] = result
        return result

    def validate_stage_fix04_compatible(
        self,
        stage: str,
        expected_context: Mapping[str, Any],
    ) -> CheckpointValidation:
        """只按 fix04/fix05 明示源码哈希放行既有 m1/m2_m3 完整检查点。

        兼容清单要求旧 marker 的 config、输入、上游和所有输出正文不变；
        未列入清单的源码也必须逐文件与旧 marker 完全一致。清单中的当前
        源码哈希必须精确匹配，后续任一代码编辑都会重新使兼容路径失效。
        """
        strict = self.validate_stage(stage, expected_context)
        if (
            strict.reusable
            or strict.reason != "context_hash_mismatch"
            or stage not in {"m1", "m2_m3"}
        ):
            return strict
        try:
            marker = json.loads(
                self.stage_marker(stage).read_text(encoding="utf-8")
            )
        except (json.JSONDecodeError, OSError):
            return strict
        compatibility_paths = (
            Path(__file__).with_name(
                "checkpoint_compatibility_fix05.json"
            ),
            Path(__file__).with_name(
                "checkpoint_compatibility_fix04.json"
            ),
        )
        observed_context = marker.get("context", {})
        observed_manifest = observed_context.get("code_manifest", {})
        legacy_files = {
            str(entry["path"]): str(entry["sha256"])
            for entry in observed_manifest.get("files", [])
        }
        current_files = {
            str(entry["path"]): str(entry["sha256"])
            for entry in expected_context.get("code_manifest", {}).get(
                "files",
                [],
            )
        }
        for compatibility_path in compatibility_paths:
            if not compatibility_path.is_file():
                continue
            try:
                compatibility = json.loads(
                    compatibility_path.read_text(encoding="utf-8")
                )
            except (json.JSONDecodeError, OSError):
                continue
            protocol = str(compatibility.get("protocol"))
            if (
                protocol
                not in {
                    "r007.verify-run-error-fix-04",
                    "r007.verify-run-error-fix-05",
                }
                or stage not in compatibility.get(
                    "compatible_stages",
                    [],
                )
                or marker.get("status") != "complete"
                or observed_manifest.get("aggregate_sha256")
                != compatibility.get("legacy_code_aggregate_sha256")
            ):
                continue
            if any(
                observed_context.get(key) != expected_context.get(key)
                for key in (
                    "stage",
                    "config_sha256",
                    "input_manifest_sha256",
                    "upstream_checkpoint_hashes",
                )
            ):
                continue
            allowed_current_files = {
                str(path): str(digest)
                for path, digest in compatibility.get(
                    "allowed_current_files",
                    {},
                ).items()
            }
            if set(legacy_files) != set(current_files):
                continue
            differing = {
                path
                for path in legacy_files
                if legacy_files[path] != current_files[path]
            }
            if differing != set(allowed_current_files):
                continue
            if any(
                current_files.get(path) != digest
                for path, digest in allowed_current_files.items()
            ):
                continue
            fix_number = "05" if protocol.endswith("05") else "04"
            reason = (
                f"reusable_fix{fix_number}_execution_compatible"
            )
            for entry in marker.get("output_manifest", []):
                output = Path(entry["path"])
                if not output.is_file():
                    reason = "output_missing"
                    break
                if _file_sha256(output) != entry.get("sha256"):
                    reason = "output_hash_mismatch"
                    break
                if output.suffix == ".parquet":
                    actual_schema = _json_sha256(
                        _frame_schema_payload(pd.read_parquet(output))
                    )
                    if actual_schema != entry.get("schema_sha256"):
                        reason = "output_schema_mismatch"
                        break
            result = CheckpointValidation(
                reason
                == f"reusable_fix{fix_number}_execution_compatible",
                reason,
                str(expected_context["context_sha256"]),
                str(observed_context.get("context_sha256")),
            )
            self.last_validation[stage] = result
            return result
        return strict

    def stage_complete(
        self,
        stage: str,
        expected_context: Mapping[str, Any],
    ) -> bool:
        return self.validate_stage(stage, expected_context).reusable

    def mark_complete(
        self,
        stage: str,
        payload: dict[str, Any],
        *,
        context: Mapping[str, Any],
        output_paths: Sequence[Path],
    ) -> None:
        path = self.stage_marker(stage)
        temporary = path.with_name(f"{path.name}.part")
        temporary.write_text(
            json.dumps(
                {
                    "status": "complete",
                    "context": _checkpoint_jsonable(dict(context)),
                    "output_manifest": self._output_manifest(output_paths),
                    **payload,
                },
                ensure_ascii=False,
                indent=2,
                default=str,
                sort_keys=True,
            )
            + "\n",
            encoding="utf-8",
        )
        json.loads(temporary.read_text(encoding="utf-8"))
        os.replace(temporary, path)

    def write_external_frame_atomic(
        self,
        path: Path,
        frame: pd.DataFrame,
        *,
        context: Mapping[str, Any],
    ) -> Path:
        """SA3 单配置 parquet 与 sidecar 均采用 .part+校验+os.replace。"""
        path.parent.mkdir(parents=True, exist_ok=True)
        temporary = path.with_name(f"{path.name}.part")
        frame.to_parquet(temporary, index=False)
        reloaded = pd.read_parquet(temporary)
        if _frame_schema_payload(reloaded) != _frame_schema_payload(frame):
            temporary.unlink(missing_ok=True)
            raise ValueError(f"SA3 单配置 schema 写后校验失败: {path.name}")
        output_hash = _file_sha256(temporary)
        os.replace(temporary, path)
        sidecar = path.with_suffix(".json")
        sidecar_temporary = sidecar.with_name(f"{sidecar.name}.part")
        sidecar_payload = {
            "status": "complete",
            "context": _checkpoint_jsonable(dict(context)),
            "output": {
                "path": str(path.resolve()),
                "sha256": output_hash,
                "schema_sha256": _json_sha256(_frame_schema_payload(reloaded)),
                "row_count": len(reloaded),
            },
        }
        sidecar_temporary.write_text(
            json.dumps(
                sidecar_payload,
                ensure_ascii=False,
                indent=2,
                sort_keys=True,
            )
            + "\n",
            encoding="utf-8",
        )
        json.loads(sidecar_temporary.read_text(encoding="utf-8"))
        os.replace(sidecar_temporary, sidecar)
        return path

    def validate_external_frame(
        self,
        path: Path,
        *,
        expected_context: Mapping[str, Any],
    ) -> CheckpointValidation:
        sidecar = path.with_suffix(".json")
        expected_hash = str(expected_context["context_sha256"])
        if not path.is_file() or not sidecar.is_file():
            return CheckpointValidation(
                False,
                "single_configuration_checkpoint_missing",
                expected_hash,
                None,
            )
        marker = json.loads(sidecar.read_text(encoding="utf-8"))
        observed = marker.get("context", {}).get("context_sha256")
        reason = "reusable"
        if observed != expected_hash:
            reason = "context_hash_mismatch"
        elif _file_sha256(path) != marker.get("output", {}).get("sha256"):
            reason = "output_hash_mismatch"
        else:
            schema_hash = _json_sha256(
                _frame_schema_payload(pd.read_parquet(path))
            )
            if schema_hash != marker.get("output", {}).get("schema_sha256"):
                reason = "output_schema_mismatch"
        return CheckpointValidation(
            reason == "reusable",
            reason,
            expected_hash,
            str(observed) if observed is not None else None,
        )

    def snapshot_path(self, snapshot: GraphSnapshot) -> Path:
        directory = self.snapshots_dir / snapshot.direction
        directory.mkdir(parents=True, exist_ok=True)
        return directory / f"{snapshot.anchor_date:%Y%m%d}.pkl.gz"

    def write_snapshot(self, snapshot: GraphSnapshot) -> Path:
        path = self.snapshot_path(snapshot)
        temporary = path.with_name(f"{path.name}.part")
        with gzip.open(temporary, "wb", compresslevel=3) as handle:
            pickle.dump(snapshot, handle, protocol=pickle.HIGHEST_PROTOCOL)
        with gzip.open(temporary, "rb") as handle:
            reloaded = pickle.load(handle)
        if not isinstance(reloaded, GraphSnapshot):
            temporary.unlink(missing_ok=True)
            raise TypeError("GraphSnapshot 检查点写后类型校验失败")
        os.replace(temporary, path)
        return path

    def iter_snapshots(self, direction: str) -> list[GraphSnapshot]:
        paths = sorted((self.snapshots_dir / direction).glob("*.pkl.gz"))
        snapshots: list[GraphSnapshot] = []
        for path in paths:
            with gzip.open(path, "rb") as handle:
                snapshot = pickle.load(handle)
            if not isinstance(snapshot, GraphSnapshot):
                raise TypeError(f"检查点不是 GraphSnapshot: {path}")
            snapshots.append(snapshot)
        return snapshots


def _concat_result_frames(
    results: list[SnapshotPredictionResult],
    attribute: str,
) -> pd.DataFrame:
    frames = [
        getattr(result, attribute)
        for result in results
        if not getattr(result, attribute).empty
    ]
    return pd.concat(frames, ignore_index=True) if frames else pd.DataFrame()


def _marker_input_manifest(
    store: CheckpointStore,
    stages: Sequence[str],
) -> dict[str, Any]:
    manifest: dict[str, Any] = {}
    for stage in stages:
        path = store.stage_marker(stage)
        if not path.is_file():
            manifest[stage] = {"marker_missing": True}
            continue
        marker = json.loads(path.read_text(encoding="utf-8"))
        manifest[stage] = {
            "marker_sha256": _file_sha256(path),
            "output_manifest": marker.get("output_manifest", []),
        }
    return manifest


def _m1_raw_input_manifest(config: M5Config) -> dict[str, Any]:
    pipeline = config.m4.m3.pipeline
    paths = {
        "us_cache_manifest": pipeline.us_yahoo.manifest_path,
        "us_constituents": pipeline.us_yahoo.constituents_path,
        "us_merged_panel": pipeline.us_yahoo.merged_panel_path,
        "cn_stock_price": pipeline.data.path(
            pipeline.data.cn_stock_price_forward
        ),
        "cn_stock_trade": pipeline.data.path(pipeline.data.cn_stock_trade),
        "cn_index_trade": pipeline.data.path(pipeline.data.cn_index_trade),
        "cn_calendar": pipeline.data.path(pipeline.data.cn_trade_calendar),
        "us_index_metadata": pipeline.data.path(
            pipeline.data.us_index_metadata
        ),
        "overseas_index_price": pipeline.data.path(
            pipeline.data.overseas_index_price
        ),
    }
    return {
        name: _path_identity_manifest(path)
        for name, path in paths.items()
    }


def _m5_input_manifest(
    store: CheckpointStore,
    config: M5Config,
) -> dict[str, Any]:
    """绑定上游检查点及 SA6 行业文件正文/schema，防止支线陈旧复用。"""
    pipeline = config.m4.m3.pipeline
    return {
        "upstream": _marker_input_manifest(
            store,
            ("m1", "m2_m3", "m4"),
        ),
        "sa6_industry_input": _path_identity_manifest(
            pipeline.data.path(pipeline.data.cn_stock_industry)
        ),
    }


def _stage_m1(
    store: CheckpointStore,
    *,
    config: M5Config,
) -> M1Tables:
    names = (
        "m1_us_panel",
        "m1_cn_panel",
        "m1_calendar_mapping",
        "m1_us_calendar",
        "m1_cn_calendar",
        "m1_us_market_returns",
        "m1_universe_audit",
        "m1_source_schema_audit",
        "m1_loaded_panel_audit",
        "m1_return_panel_audit",
        "m1_security_coverage_audit",
        "m1_return_distribution_audit",
    )
    context = store.build_context(
        "m1",
        config=config.m4.m3.pipeline,
        input_manifest=_m1_raw_input_manifest(config),
    )
    if (
        store.validate_stage_fix04_compatible("m1", context).reusable
        and all(store.has_frame(name) for name in names)
    ):
        return M1Tables(*(store.read_frame(name) for name in names))

    us_universe, cn_universe = build_dual_market_universes(config.m4.m3.pipeline)
    sources = load_dual_market_sources(
        us_universe["stock_code"].astype(str).tolist(),
        cn_universe["stock_code"].astype(str).tolist(),
        config=config.m4.m3.pipeline,
    )
    panels = build_dual_market_return_panels(
        sources,
        config=config.m4.m3.pipeline,
    )
    us_contract_audit = audit_us_merged_panel_contract(
        sources.us_quotes,
        us_universe["stock_code"].astype(str).tolist(),
        config=config.m4.m3.pipeline,
    )
    calendar_mapping = build_bidirectional_calendar_mappings(
        sources.us_calendar,
        sources.cn_calendar,
        config=config.m4.m3.pipeline,
    )
    us_market_returns = calculate_market_log_returns(sources.us_index)
    universe_audit = pd.concat(
        [
            us_universe.assign(universe_source="AS24_AS26_fixed_yahoo_pool"),
            cn_universe.assign(universe_source="AS11_average_market_cap"),
        ],
        ignore_index=True,
        sort=False,
    )
    source_schema_audit = audit_m1_source_schemas(config.m4.m3.pipeline)
    loaded_panel_audit = pd.concat(
        [
            audit_loaded_market_panel(
                sources.us_quotes,
                dataset_name="US_loaded_source_panel",
                key_columns=("stock_code", "date"),
                required_columns=US_SOURCE_REQUIRED_COLUMNS,
                unit_contract=US_UNIT_CONTRACT,
                expected_calendar=sources.us_calendar,
            ),
            audit_loaded_market_panel(
                sources.cn_quotes,
                dataset_name="CN_loaded_source_panel",
                key_columns=("stock_code", "date"),
                required_columns=tuple(
                    dict.fromkeys(
                        (*CN_PRICE_REQUIRED_COLUMNS, *CN_TRADE_REQUIRED_COLUMNS)
                    )
                ),
                unit_contract=CN_UNIT_CONTRACT,
                expected_calendar=sources.cn_calendar,
            ),
        ],
        ignore_index=True,
    )
    return_panel_audit = audit_dual_market_return_panels(
        panels,
        us_calendar=sources.us_calendar,
        cn_calendar=sources.cn_calendar,
    )
    security_coverage_audit = pd.concat(
        [
            us_contract_audit.assign(audit_section="us_manifest_panel_contract"),
            audit_security_calendar_coverage(
                panels.us,
                us_universe,
                sources.us_calendar,
                market="US",
            ).assign(audit_section="security_calendar_coverage"),
            audit_security_calendar_coverage(
                panels.cn,
                cn_universe,
                sources.cn_calendar,
                market="CN",
            ).assign(audit_section="security_calendar_coverage"),
        ],
        ignore_index=True,
        sort=False,
    )
    return_distribution_audit = audit_return_distributions(panels)
    failed_m1_checks = pd.concat(
        [source_schema_audit, loaded_panel_audit, return_panel_audit],
        ignore_index=True,
        sort=False,
    )
    failed_m1_checks = failed_m1_checks.loc[
        failed_m1_checks["status"].astype(str).eq("fail")
    ]
    if not failed_m1_checks.empty:
        raise ValueError(
            "M1 数据字段、键、日历或 OHLC 门禁失败: "
            f"{failed_m1_checks.head(20).to_dict(orient='records')}"
        )
    tables = M1Tables(
        panels.us,
        panels.cn,
        calendar_mapping,
        sources.us_calendar,
        sources.cn_calendar,
        us_market_returns,
        universe_audit,
        source_schema_audit,
        loaded_panel_audit,
        return_panel_audit,
        security_coverage_audit,
        return_distribution_audit,
    )
    output_paths = [
        store.write_frame(name, frame)
        for name, frame in zip(names, tables.__dict__.values(), strict=True)
    ]
    store.mark_complete(
        "m1",
        {
            "us_stock_count": panels.us["stock_code"].nunique(),
            "cn_stock_count": panels.cn["stock_code"].nunique(),
        },
        context=context,
        output_paths=output_paths,
    )
    return tables


def _stage_m2_m3(
    m1: M1Tables,
    store: CheckpointStore,
    *,
    load_reporting_frames: bool = True,
    config: M5Config,
) -> M3Tables:
    names = (
        "m3_predictions",
        "m3_training_audit",
        "m3_model_audit",
        "m3_feature_values",
        "m2_graph_audit",
        "m2_indegree_quantiles",
        "m3_eligibility_audit",
    )
    context = store.build_context(
        "m2_m3",
        config=config.m4.m3,
        input_manifest=_marker_input_manifest(store, ("m1",)),
        upstream_stages=("m1",),
    )
    if (
        store.validate_stage_fix04_compatible("m2_m3", context).reusable
        and all(store.has_frame(name) for name in names)
        and any(
            (store.snapshots_dir / "US_TO_CN").glob("*.pkl.gz")
        )
        and any(
            (store.snapshots_dir / "CN_TO_US").glob("*.pkl.gz")
        )
    ):
        if load_reporting_frames:
            return M3Tables(*(store.read_frame(name) for name in names))
        return M3Tables(
            predictions=store.read_frame("m3_predictions"),
            training_audit=pd.DataFrame(),
            model_audit=pd.DataFrame(),
            feature_values=pd.DataFrame(),
            graph_audit=pd.DataFrame(),
            indegree_quantiles=store.read_frame("m2_indegree_quantiles"),
            eligibility_audit=store.read_frame("m3_eligibility_audit"),
        )
    existing_snapshots = sorted(
        store.snapshots_dir.glob("*/*.pkl.gz")
    )
    if existing_snapshots:
        validation = store.last_validation.get("m2_m3")
        raise RuntimeError(
            "M2/M3 已存在正式图快照但完整 marker 不可复用；"
            "为避免静默删除 299 个正式快照，本次拒绝自动重建。"
            f" validation={validation}"
        )
    for direction in ("US_TO_CN", "CN_TO_US"):
        for stale_snapshot in (
            store.snapshots_dir / direction
        ).glob("*.pkl.gz"):
            stale_snapshot.unlink()

    factories = build_base_model_factories(config.m4.m3)
    prediction_results: list[SnapshotPredictionResult] = []
    graph_audits: list[pd.DataFrame] = []
    indegree_tables: list[pd.DataFrame] = []
    snapshots = iter_bidirectional_rolling_graphs(
        m1.us_panel,
        m1.cn_panel,
        m1.calendar_mapping,
        prediction_start=config.m4.m3.prediction.prediction_start,
        prediction_end=config.m4.m3.prediction.prediction_end,
        source_value_column=config.m4.m3.prediction.source_value_column,
        target_value_column=config.m4.m3.prediction.target_value_column,
        config=config.m4.m3.pipeline,
    )
    panels = {"US": m1.us_panel, "CN": m1.cn_panel}
    for snapshot in snapshots:
        store.write_snapshot(snapshot)
        graph_audits.append(snapshot.graph_audit)
        indegree_tables.append(snapshot.indegree_quantiles)
        prediction_results.append(
            predict_graph_snapshot(
                snapshot,
                panels[snapshot.source_market],
                panels[snapshot.target_market],
                model_factories=factories,
                config=config.m4.m3,
            )
        )
    tables = M3Tables(
        predictions=_concat_result_frames(prediction_results, "predictions"),
        training_audit=_concat_result_frames(
            prediction_results, "training_audit"
        ),
        model_audit=_concat_result_frames(prediction_results, "model_audit"),
        feature_values=_concat_result_frames(
            prediction_results, "feature_values"
        ),
        graph_audit=(
            pd.concat(graph_audits, ignore_index=True)
            if graph_audits
            else pd.DataFrame()
        ),
        indegree_quantiles=(
            pd.concat(indegree_tables, ignore_index=True)
            if indegree_tables
            else pd.DataFrame()
        ),
        eligibility_audit=_concat_result_frames(
            prediction_results, "eligibility_audit"
        ),
    )
    output_paths = [
        store.write_frame(name, frame)
        for name, frame in zip(names, tables.__dict__.values(), strict=True)
    ]
    output_paths.extend(
        sorted(store.snapshots_dir.glob("*/*.pkl.gz"))
    )
    store.mark_complete(
        "m2_m3",
        {
            "prediction_rows": len(tables.predictions),
            "snapshot_count": len(graph_audits),
        },
        context=context,
        output_paths=output_paths,
    )
    return tables


def _prepare_cross_predictions(
    predictions: pd.DataFrame,
    eligibility_audit: pd.DataFrame,
    m1: M1Tables,
    *,
    feature_set: str,
    config: M5Config,
) -> SpecificationPredictionTables:
    parts: list[pd.DataFrame] = []
    for direction, target_market, target_panel in (
        ("US_TO_CN", "CN", m1.cn_panel),
        ("CN_TO_US", "US", m1.us_panel),
    ):
        frame = predictions.loc[predictions["direction"].eq(direction)].copy()
        if frame.empty:
            continue
        parts.append(
            prepare_portfolio_predictions(
                frame,
                target_panel,
                target_market=target_market,
                specification=config.m4.baseline.cross_market_specification,
                feature_set=feature_set,
                config=config.m4,
            )
        )
    eligibility = eligibility_audit.copy()
    eligibility["specification"] = (
        config.m4.baseline.cross_market_specification
    )
    eligibility["feature_set"] = feature_set
    eligibility["target_market"] = eligibility["direction"].map(
        {"US_TO_CN": "CN", "CN_TO_US": "US"}
    )
    if eligibility["target_market"].isna().any():
        raise ValueError("跨市场动态资格审计含未知方向")
    return SpecificationPredictionTables(
        predictions=(
            pd.concat(parts, ignore_index=True) if parts else pd.DataFrame()
        ),
        eligibility_audit=eligibility,
    )


def _m4_component_context(
    parent_context: Mapping[str, Any],
    *,
    component: str,
    key: str,
) -> dict[str, Any]:
    payload = {
        "stage": "m4_component",
        "parent_context_sha256": str(parent_context["context_sha256"]),
        "component": component,
        "key": key,
    }
    return {**payload, "context_sha256": _json_sha256(payload)}


def _fix05_legacy_m4_component_context(
    parent_context: Mapping[str, Any],
    *,
    component: str,
    key: str,
) -> dict[str, Any] | None:
    """对 fix05 明示的 attempt5 M4 组件构造只读兼容上下文。

    仅 checkpoint 写后 dtype 校验源码发生变化时，既有 M4 组件的业务计算
    正文仍有效。兼容清单精确绑定当前完整代码 aggregate、config、输入、上游
    marker 与新旧 parent context；任一身份变化均拒绝复用。
    """
    compatibility_path = Path(__file__).with_name(
        "checkpoint_compatibility_fix05.json"
    )
    if not compatibility_path.is_file():
        return None
    try:
        compatibility = json.loads(
            compatibility_path.read_text(encoding="utf-8")
        )
    except (json.JSONDecodeError, OSError):
        return None
    component_compatibility = compatibility.get(
        "m4_component_compatibility",
        {},
    )
    if (
        compatibility.get("protocol")
        != "r007.verify-run-error-fix-05"
        or component
        not in component_compatibility.get("components", [])
        or parent_context.get("stage") != "m4"
        or parent_context.get("context_sha256")
        != component_compatibility.get(
            "current_parent_context_sha256"
        )
        or parent_context.get("config_sha256")
        != component_compatibility.get("config_sha256")
        or parent_context.get("input_manifest_sha256")
        != component_compatibility.get("input_manifest_sha256")
        or parent_context.get("upstream_checkpoint_hashes")
        != component_compatibility.get("upstream_checkpoint_hashes")
        or parent_context.get("code_manifest", {}).get(
            "aggregate_sha256"
        )
        != component_compatibility.get(
            "current_code_aggregate_sha256"
        )
    ):
        return None
    legacy_parent_hash = component_compatibility.get(
        "legacy_parent_context_sha256"
    )
    if not isinstance(legacy_parent_hash, str):
        return None
    return _m4_component_context(
        {"context_sha256": legacy_parent_hash},
        component=component,
        key=key,
    )


def _m4_component_paths(
    store: CheckpointStore,
    *,
    component: str,
    key: str,
) -> tuple[Path, Path]:
    allowed = set(
        "abcdefghijklmnopqrstuvwxyz"
        "ABCDEFGHIJKLMNOPQRSTUVWXYZ"
        "0123456789_-"
    )
    if not component or not key or not set(component + key) <= allowed:
        raise ValueError("M4 组件名称与键只能包含字母、数字、下划线和连字符")
    directory = store.root / "components" / "m4" / component / key
    return (
        directory / "predictions.parquet",
        directory / "eligibility_audit.parquet",
    )


def _read_m4_component(
    store: CheckpointStore,
    parent_context: Mapping[str, Any],
    *,
    component: str,
    key: str,
) -> SpecificationPredictionTables | None:
    current_context = _m4_component_context(
        parent_context,
        component=component,
        key=key,
    )
    contexts = [current_context]
    legacy_context = _fix05_legacy_m4_component_context(
        parent_context,
        component=component,
        key=key,
    )
    if (
        legacy_context is not None
        and legacy_context["context_sha256"]
        != current_context["context_sha256"]
    ):
        contexts.append(legacy_context)
    prediction_path, eligibility_path = _m4_component_paths(
        store,
        component=component,
        key=key,
    )
    for context in contexts:
        if (
            store.validate_external_frame(
                prediction_path,
                expected_context=context,
            ).reusable
            and store.validate_external_frame(
                eligibility_path,
                expected_context=context,
            ).reusable
        ):
            return SpecificationPredictionTables(
                predictions=pd.read_parquet(prediction_path),
                eligibility_audit=pd.read_parquet(eligibility_path),
            )
    return None


def _write_m4_component(
    store: CheckpointStore,
    parent_context: Mapping[str, Any],
    tables: SpecificationPredictionTables,
    *,
    component: str,
    key: str,
) -> SpecificationPredictionTables:
    context = _m4_component_context(
        parent_context,
        component=component,
        key=key,
    )
    prediction_path, eligibility_path = _m4_component_paths(
        store,
        component=component,
        key=key,
    )
    store.write_external_frame_atomic(
        prediction_path,
        tables.predictions,
        context=context,
    )
    store.write_external_frame_atomic(
        eligibility_path,
        tables.eligibility_audit,
        context=context,
    )
    reloaded = _read_m4_component(
        store,
        parent_context,
        component=component,
        key=key,
    )
    if reloaded is None:
        raise RuntimeError(f"M4 组件写后校验失败: {component}/{key}")
    return reloaded


def _m4_expected_anchor_dates(
    panel: pd.DataFrame,
    *,
    market: str,
    config: M5Config,
) -> tuple[pd.Timestamp, ...]:
    mapping = build_same_market_calendar_mapping(
        panel,
        market=market,
        config=config.m4,
    )
    blocks = build_update_blocks(
        mapping,
        prediction_start=config.m4.m3.prediction.prediction_start,
        prediction_end=config.m4.m3.prediction.prediction_end,
        update_days=config.m4.m3.pipeline.returns.graph_update_days,
    )
    return tuple(block.anchor_date for block in blocks)


def _m4_result_anchor(
    eligibility_audit: pd.DataFrame,
    *,
    component: str,
) -> pd.Timestamp:
    if "snapshot_date" not in eligibility_audit.columns:
        raise KeyError(f"{component} 资格审计缺少 snapshot_date")
    anchors = pd.DatetimeIndex(
        pd.to_datetime(
            eligibility_audit["snapshot_date"],
            errors="raise",
        ).dt.normalize().unique()
    )
    if len(anchors) != 1:
        raise ValueError(f"{component} 单块资格审计锚点数不是 1: {len(anchors)}")
    return pd.Timestamp(anchors[0]).normalize()


def _m4_block_key(market: str, anchor: pd.Timestamp) -> str:
    return f"{market}_{pd.Timestamp(anchor):%Y%m%d}"


def _collect_same_market_predictions(
    m1: M1Tables,
    store: CheckpointStore,
    parent_context: Mapping[str, Any],
    *,
    source_value_column: str,
    config: M5Config,
) -> SpecificationPredictionTables:
    aggregate_component = "same_market_graph_aggregate"
    existing_aggregate = _read_m4_component(
        store,
        parent_context,
        component=aggregate_component,
        key="all",
    )
    if existing_aggregate is not None:
        return existing_aggregate
    parts: list[pd.DataFrame] = []
    eligibility_parts: list[pd.DataFrame] = []
    for market, panel in (("US", m1.us_panel), ("CN", m1.cn_panel)):
        expected_anchors = _m4_expected_anchor_dates(
            panel,
            market=market,
            config=config,
        )
        completed_anchors = {
            anchor
            for anchor in expected_anchors
            if _read_m4_component(
                store,
                parent_context,
                component="same_market_graph_block",
                key=_m4_block_key(market, anchor),
            )
            is not None
        }
        for result in iter_same_market_graph_predictions(
            panel,
            market=market,
            source_value_column=source_value_column,
            skip_anchor_dates=completed_anchors,
            config=config.m4,
        ):
            anchor = _m4_result_anchor(
                result.eligibility_audit,
                component="same_market_graph_block",
            )
            if anchor not in expected_anchors:
                raise ValueError(
                    f"{market} 同市场图产生意外锚点: {anchor.date()}"
                )
            _write_m4_component(
                store,
                parent_context,
                SpecificationPredictionTables(
                    predictions=result.predictions,
                    eligibility_audit=result.eligibility_audit,
                ),
                component="same_market_graph_block",
                key=_m4_block_key(market, anchor),
            )
        market_blocks: list[SpecificationPredictionTables] = []
        for anchor in expected_anchors:
            block = _read_m4_component(
                store,
                parent_context,
                component="same_market_graph_block",
                key=_m4_block_key(market, anchor),
            )
            if block is None:
                raise RuntimeError(
                    f"{market} 同市场图块缺失或损坏: {anchor.date()}"
                )
            market_blocks.append(block)
        predictions = pd.concat(
            [block.predictions for block in market_blocks],
            ignore_index=True,
        )
        eligibility = pd.concat(
            [block.eligibility_audit for block in market_blocks],
            ignore_index=True,
        )
        if not eligibility.empty:
            eligibility["specification"] = (
                config.m4.baseline.same_market_graph_specification
            )
            eligibility["feature_set"] = source_value_column
            eligibility["target_market"] = market
            eligibility_parts.append(eligibility)
        if predictions.empty:
            continue
        parts.append(
            prepare_portfolio_predictions(
                predictions,
                panel,
                target_market=market,
                specification=(
                    config.m4.baseline.same_market_graph_specification
                ),
                feature_set=source_value_column,
                config=config.m4,
            )
        )
    tables = SpecificationPredictionTables(
        predictions=(
            pd.concat(parts, ignore_index=True) if parts else pd.DataFrame()
        ),
        eligibility_audit=(
            pd.concat(eligibility_parts, ignore_index=True)
            if eligibility_parts
            else pd.DataFrame()
        ),
    )
    return _write_m4_component(
        store,
        parent_context,
        tables,
        component=aggregate_component,
        key="all",
    )


def _collect_non_graph_predictions(
    m1: M1Tables,
    store: CheckpointStore,
    parent_context: Mapping[str, Any],
    *,
    source_value_column: str,
    config: M5Config,
) -> SpecificationPredictionTables:
    aggregate_component = "non_graph_aggregate"
    existing_aggregate = _read_m4_component(
        store,
        parent_context,
        component=aggregate_component,
        key="all",
    )
    if existing_aggregate is not None:
        return existing_aggregate
    parts: list[pd.DataFrame] = []
    eligibility_parts: list[pd.DataFrame] = []
    for market, panel in (("US", m1.us_panel), ("CN", m1.cn_panel)):
        expected_anchors = _m4_expected_anchor_dates(
            panel,
            market=market,
            config=config,
        )
        completed_anchors = {
            anchor
            for anchor in expected_anchors
            if _read_m4_component(
                store,
                parent_context,
                component="non_graph_block",
                key=_m4_block_key(market, anchor),
            )
            is not None
        }
        for result in iter_non_graph_autoregressive_predictions(
            panel,
            market=market,
            source_value_column=source_value_column,
            skip_anchor_dates=completed_anchors,
            config=config.m4,
        ):
            anchor = _m4_result_anchor(
                result.eligibility_audit,
                component="non_graph_block",
            )
            if anchor not in expected_anchors:
                raise ValueError(
                    f"{market} 非图基线产生意外锚点: {anchor.date()}"
                )
            _write_m4_component(
                store,
                parent_context,
                SpecificationPredictionTables(
                    predictions=result.predictions,
                    eligibility_audit=result.eligibility_audit,
                ),
                component="non_graph_block",
                key=_m4_block_key(market, anchor),
            )
        market_blocks: list[SpecificationPredictionTables] = []
        for anchor in expected_anchors:
            block = _read_m4_component(
                store,
                parent_context,
                component="non_graph_block",
                key=_m4_block_key(market, anchor),
            )
            if block is None:
                raise RuntimeError(
                    f"{market} 非图基线块缺失或损坏: {anchor.date()}"
                )
            market_blocks.append(block)
        predictions = pd.concat(
            [block.predictions for block in market_blocks],
            ignore_index=True,
        )
        eligibility = pd.concat(
            [block.eligibility_audit for block in market_blocks],
            ignore_index=True,
        )
        if not eligibility.empty:
            eligibility["specification"] = (
                config.m4.baseline.non_graph_specification
            )
            eligibility["feature_set"] = source_value_column
            eligibility["target_market"] = market
            eligibility_parts.append(eligibility)
        if predictions.empty:
            continue
        parts.append(
            prepare_portfolio_predictions(
                predictions,
                panel,
                target_market=market,
                specification=config.m4.baseline.non_graph_specification,
                feature_set=source_value_column,
                config=config.m4,
            )
        )
    tables = SpecificationPredictionTables(
        predictions=(
            pd.concat(parts, ignore_index=True) if parts else pd.DataFrame()
        ),
        eligibility_audit=(
            pd.concat(eligibility_parts, ignore_index=True)
            if eligibility_parts
            else pd.DataFrame()
        ),
    )
    return _write_m4_component(
        store,
        parent_context,
        tables,
        component=aggregate_component,
        key="all",
    )


def _stage_m4(
    m1: M1Tables,
    m3: M3Tables,
    store: CheckpointStore,
    *,
    config: M5Config,
) -> M4Tables:
    names = (
        "m4_prepared_predictions",
        "m4_sharpe_metrics",
        "m4_daily_pnl",
        "m4_portfolio_members",
        "m4_nested_audit",
        "m4_model_completeness_audit",
        "m4_comparability_audit",
        "m4_execution_timing_audit",
        "m4_core_observations",
        "m4_core_paired_comparisons",
        "m4_core_summary",
        "m4_formal_coverage_audit",
    )
    context = store.build_context(
        "m4",
        config=config.m4,
        input_manifest=_marker_input_manifest(store, ("m1", "m2_m3")),
        upstream_stages=("m1", "m2_m3"),
    )
    if (
        store.stage_complete("m4", context)
        and all(store.has_frame(name) for name in names)
    ):
        return M4Tables(*(store.read_frame(name) for name in names))
    feature_set = config.m4.m3.prediction.source_value_column
    cross = _read_m4_component(
        store,
        context,
        component="cross_market_graph_aggregate",
        key="all",
    )
    if cross is None:
        cross = _write_m4_component(
            store,
            context,
            _prepare_cross_predictions(
                m3.predictions,
                m3.eligibility_audit,
                m1,
                feature_set=feature_set,
                config=config,
            ),
            component="cross_market_graph_aggregate",
            key="all",
        )
    # 跨市场规格已原子落盘后释放 10M+ 行 m3 面板；正式报告前会从
    # m2_m3 完整 marker 绑定的只读 frame 重新载入，不改变任何结果值。
    m3.predictions = pd.DataFrame()
    m3.eligibility_audit = pd.DataFrame()
    same_graph = _collect_same_market_predictions(
        m1,
        store,
        context,
        source_value_column=feature_set,
        config=config,
    )
    non_graph = _collect_non_graph_predictions(
        m1,
        store,
        context,
        source_value_column=feature_set,
        config=config,
    )
    result = run_m4_portfolios(
        {
            config.m4.baseline.cross_market_specification: cross.predictions,
            config.m4.baseline.same_market_graph_specification: (
                same_graph.predictions
            ),
            config.m4.baseline.non_graph_specification: (
                non_graph.predictions
            ),
        },
        eligibility_frames={
            config.m4.baseline.cross_market_specification: (
                cross.eligibility_audit
            ),
            config.m4.baseline.same_market_graph_specification: (
                same_graph.eligibility_audit
            ),
            config.m4.baseline.non_graph_specification: (
                non_graph.eligibility_audit
            ),
        },
        config=config.m4,
    )
    comparison = build_core_comparison_tables(
        result.portfolios.sharpe_metrics,
        config=config.m4,
    )
    tables = M4Tables(
        prepared_predictions=result.predictions,
        sharpe_metrics=result.portfolios.sharpe_metrics,
        daily_pnl=result.portfolios.daily_pnl,
        portfolio_members=result.portfolios.portfolio_members,
        nested_audit=result.portfolios.nested_audit,
        model_completeness_audit=result.model_completeness_audit,
        comparability_audit=result.comparability_audit,
        execution_timing_audit=result.execution_timing_audit,
        core_observations=comparison.observations,
        core_paired_comparisons=comparison.paired_comparisons,
        core_summary=comparison.summary,
        formal_coverage_audit=result.formal_coverage_audit,
    )
    output_paths = [
        store.write_frame(name, frame)
        for name, frame in zip(names, tables.__dict__.values(), strict=True)
    ]
    store.mark_complete(
        "m4",
        {
            "prepared_prediction_rows": len(tables.prepared_predictions),
            "sharpe_rows": len(tables.sharpe_metrics),
        },
        context=context,
        output_paths=output_paths,
    )
    return tables


def _us_to_cn_prepared(m4: M4Tables) -> pd.DataFrame:
    return m4.prepared_predictions.loc[
        m4.prepared_predictions["direction"].eq("US_TO_CN")
        & m4.prepared_predictions["target_market"].eq("CN")
        & m4.prepared_predictions["specification"].eq("cross_market_graph")
    ].copy()


def _predict_snapshots(
    snapshots: list[GraphSnapshot],
    m1: M1Tables,
    *,
    config: M5Config,
) -> pd.DataFrame:
    factories = build_base_model_factories(config.m4.m3)
    results = [
        predict_graph_snapshot(
            snapshot,
            m1.us_panel if snapshot.source_market == "US" else m1.cn_panel,
            m1.us_panel if snapshot.target_market == "US" else m1.cn_panel,
            model_factories=factories,
            config=config.m4.m3,
        )
        for snapshot in snapshots
    ]
    return _concat_result_frames(results, "predictions")


def _run_sa1(
    m1: M1Tables,
    m4: M4Tables,
    store: CheckpointStore,
    *,
    config: M5Config,
) -> tuple[pd.DataFrame, pd.DataFrame, pd.DataFrame]:
    base = _us_to_cn_prepared(m4)
    summary_parts: list[pd.DataFrame] = []
    audit_parts: list[pd.DataFrame] = []
    median_audit_parts: list[pd.DataFrame] = []
    snapshots = store.iter_snapshots("US_TO_CN")
    for fraction in config.sensitivity.edge_replacement_fractions:
        if fraction == 0:
            prepared = base.copy()
        else:
            randomized_snapshots: list[GraphSnapshot] = []
            for snapshot in snapshots:
                randomized = randomize_graph_snapshot(
                    snapshot,
                    replacement_fraction=fraction,
                    seed=config.sensitivity.seed,
                    config=config,
                )
                randomized_snapshots.append(randomized.snapshot)
                audit_parts.append(randomized.audit)
            predictions = _predict_snapshots(
                randomized_snapshots,
                m1,
                config=config,
            )
            prepared = prepare_portfolio_predictions(
                predictions,
                m1.cn_panel,
                target_market="CN",
                specification="cross_market_graph",
                feature_set=config.m4.m3.prediction.source_value_column,
                config=config.m4,
            )
        prepared["replacement_fraction"] = fraction
        result = run_sensitivity_portfolios(
            prepared,
            group_columns=(
                "replacement_fraction",
                "feature_set",
                "direction",
                "target_market",
                "model",
            ),
            config=config,
        )
        median = summarize_model_median_sharpe(
            result.sharpe_metrics,
            sensitivity_columns=("replacement_fraction",),
            config=config,
        )
        summary_parts.append(median.summary)
        median_audit_parts.append(median.audit)
    summary = pd.concat(summary_parts, ignore_index=True)
    audit = pd.concat(audit_parts, ignore_index=True) if audit_parts else pd.DataFrame()
    median_audit = pd.concat(median_audit_parts, ignore_index=True)
    store.write_frame("m5_sa1_summary", summary)
    store.write_frame("m5_sa1_audit", audit)
    store.write_frame("m5_sa1_median_audit", median_audit)
    return summary, audit, median_audit


def _run_sa2(
    m1: M1Tables,
    m4: M4Tables,
    store: CheckpointStore,
    *,
    config: M5Config,
) -> tuple[pd.DataFrame, pd.DataFrame, pd.DataFrame]:
    base = _us_to_cn_prepared(m4)
    summary_parts: list[pd.DataFrame] = []
    audit_parts: list[pd.DataFrame] = []
    median_audit_parts: list[pd.DataFrame] = []
    for lag in config.sensitivity.information_lags:
        if lag == 1:
            prepared = base.copy()
        else:
            mapping = build_information_lag_mapping(
                m1.cn_calendar,
                m1.us_calendar,
                lag=lag,
                config=config,
            )
            direction = replace(
                config.m4.m3.pipeline.calendar.direction("US_TO_CN"),
                lag_parameter=lag,
            )
            local_calendar = replace(
                config.m4.m3.pipeline.calendar,
                directions=(direction,),
            )
            local_m1 = replace(
                config.m4.m3.pipeline,
                calendar=local_calendar,
            )
            local_m3 = replace(config.m4.m3, pipeline=local_m1)
            snapshots = list(
                iter_rolling_graph_snapshots(
                    m1.us_panel,
                    m1.cn_panel,
                    mapping,
                    direction="US_TO_CN",
                    prediction_start=local_m3.prediction.prediction_start,
                    prediction_end=local_m3.prediction.prediction_end,
                    source_value_column=local_m3.prediction.source_value_column,
                    target_value_column=local_m3.prediction.target_value_column,
                    config=local_m1,
                )
            )
            local_config = replace(config, m4=replace(config.m4, m3=local_m3))
            predictions = _predict_snapshots(snapshots, m1, config=local_config)
            prepared = prepare_portfolio_predictions(
                predictions,
                m1.cn_panel,
                target_market="CN",
                specification="cross_market_graph",
                feature_set=local_m3.prediction.source_value_column,
                config=local_config.m4,
            )
            audit_parts.append(
                mapping.assign(
                    sensitivity_id="SA2",
                    lag=lag,
                )
            )
        prepared["information_lag"] = lag
        result = run_sensitivity_portfolios(
            prepared,
            group_columns=(
                "information_lag",
                "feature_set",
                "direction",
                "target_market",
                "model",
            ),
            config=config,
        )
        median = summarize_model_median_sharpe(
            result.sharpe_metrics,
            sensitivity_columns=("information_lag",),
            config=config,
        )
        summary_parts.append(median.summary)
        median_audit_parts.append(median.audit)
    summary = pd.concat(summary_parts, ignore_index=True)
    audit = pd.concat(audit_parts, ignore_index=True) if audit_parts else pd.DataFrame()
    median_audit = pd.concat(median_audit_parts, ignore_index=True)
    store.write_frame("m5_sa2_summary", summary)
    store.write_frame("m5_sa2_audit", audit)
    store.write_frame("m5_sa2_median_audit", median_audit)
    return summary, audit, median_audit


def _sa3_sample_key_sha256(frame: pd.DataFrame) -> str:
    keys = (
        frame.loc[:, ["target_date", "target_code"]]
        .assign(
            target_date=lambda value: pd.to_datetime(
                value["target_date"], errors="raise"
            ).dt.strftime("%Y-%m-%d"),
            target_code=lambda value: value["target_code"].astype(str),
        )
        .drop_duplicates()
        .sort_values(["target_date", "target_code"])
    )
    payload = "\n".join(
        f"{row.target_date}\x1f{row.target_code}"
        for row in keys.itertuples(index=False)
    ).encode("utf-8")
    return hashlib.sha256(payload).hexdigest()


def _run_sa3(
    m1: M1Tables,
    store: CheckpointStore,
    *,
    config: M5Config,
) -> tuple[pd.DataFrame, pd.DataFrame, pd.DataFrame, pd.DataFrame]:
    snapshots = store.iter_snapshots("US_TO_CN")
    grid = build_hyperparameter_grid(config=config)
    prepared_parts: list[pd.DataFrame] = []
    model_audit_parts: list[pd.DataFrame] = []
    sa3_dir = store.frames_dir / "m5_sa3_configurations"
    sa3_dir.mkdir(parents=True, exist_ok=True)
    for row in grid.to_dict(orient="records"):
        configuration_id = str(row["configuration_id"])
        prediction_checkpoint = (
            sa3_dir / f"{configuration_id}_prepared_predictions.parquet"
        )
        audit_checkpoint = sa3_dir / f"{configuration_id}_model_audit.parquet"
        configuration_context = store.build_context(
            f"m5_sa3_{configuration_id}",
            config={
                "m5": config,
                "configuration": row,
            },
            input_manifest=_marker_input_manifest(
                store,
                ("m1", "m2_m3", "m4"),
            ),
            upstream_stages=("m1", "m2_m3", "m4"),
        )
        predictions_reusable = store.validate_external_frame(
            prediction_checkpoint,
            expected_context=configuration_context,
        ).reusable
        audit_reusable = store.validate_external_frame(
            audit_checkpoint,
            expected_context=configuration_context,
        ).reusable
        if predictions_reusable and audit_reusable:
            prepared_parts.append(pd.read_parquet(prediction_checkpoint))
            model_audit_parts.append(pd.read_parquet(audit_checkpoint))
            continue
        parameters = json.loads(str(row["parameters_json"]))
        factory = build_hyperparameter_factory(
            str(row["model"]),
            parameters,
            m3_config=config.m4.m3,
        )
        prediction_parts: list[pd.DataFrame] = []
        configuration_model_audits: list[pd.DataFrame] = []
        for snapshot in snapshots:
            prediction = predict_snapshot_single_model(
                snapshot,
                m1.us_panel,
                m1.cn_panel,
                model=str(row["model"]),
                model_factory=factory,
                config=config,
            )
            if not prediction.predictions.empty:
                prediction_parts.append(prediction.predictions)
            configuration_model_audits.append(prediction.model_audit)
        predictions = (
            pd.concat(prediction_parts, ignore_index=True)
            if prediction_parts
            else pd.DataFrame()
        )
        model_audit = (
            pd.concat(configuration_model_audits, ignore_index=True)
            if configuration_model_audits
            else pd.DataFrame()
        )
        model_audit.insert(0, "configuration_id", configuration_id)
        model_audit.insert(1, "parameters_json", row["parameters_json"])
        prepared = prepare_portfolio_predictions(
            predictions,
            m1.cn_panel,
            target_market="CN",
            specification="cross_market_graph",
            feature_set=config.m4.m3.prediction.source_value_column,
            config=config.m4,
        )
        prepared["configuration_id"] = configuration_id
        prepared["parameters_json"] = row["parameters_json"]
        store.write_external_frame_atomic(
            prediction_checkpoint,
            prepared,
            context=configuration_context,
        )
        store.write_external_frame_atomic(
            audit_checkpoint,
            model_audit,
            context=configuration_context,
        )
        prepared_parts.append(prepared)
        model_audit_parts.append(model_audit)

    prepared_all = pd.concat(prepared_parts, ignore_index=True)
    model_audit_all = pd.concat(
        model_audit_parts,
        ignore_index=True,
        sort=False,
    )
    eligibility = store.read_frame("m3_eligibility_audit")
    required_eligibility = {
        "direction",
        "prediction_date",
        "snapshot_date",
        "target_code",
        "graph_has_in_edges",
        "feature_count",
        "complete_case_nobs",
        "min_required_nobs",
        "predictor_complete",
        "actual_return_available",
        "historical_liquidity_available",
        "eligibility_status",
        "exclusion_reason",
    }
    missing_eligibility = sorted(
        required_eligibility - set(eligibility.columns)
    )
    if missing_eligibility:
        raise KeyError(f"SA3 动态资格审计缺少字段: {missing_eligibility}")
    eligibility = eligibility.loc[
        eligibility["direction"].astype(str).eq("US_TO_CN")
    ].copy()
    eligibility["prediction_date"] = pd.to_datetime(
        eligibility["prediction_date"], errors="raise"
    ).dt.normalize()
    if eligibility.duplicated(["prediction_date", "target_code"]).any():
        raise ValueError("SA3 动态资格审计存在重复日期/股票机会键")
    eligibility = validate_eligibility_audit(
        eligibility,
        context="SA3 动态资格消费门禁",
    )
    exclusion_reason = eligibility["exclusion_reason"].fillna("").astype(str)
    predictable_mask = eligibility["expected_eligibility_status"].astype(
        str
    ).eq("predictable")
    expected_keys = eligibility.loc[
        predictable_mask,
        ["prediction_date", "target_code"],
    ].rename(columns={"prediction_date": "target_date"})
    expected_keys["target_code"] = expected_keys["target_code"].astype(str)
    expected_key_hash = _sa3_sample_key_sha256(expected_keys)
    eligibility_lookup = eligibility.loc[
        :,
        [
            "prediction_date",
            "target_code",
            "eligibility_status",
            "exclusion_reason",
        ],
    ].rename(columns={"prediction_date": "target_date"})
    prepared_all["target_date"] = pd.to_datetime(
        prepared_all["target_date"], errors="raise"
    ).dt.normalize()
    prepared_all["target_code"] = prepared_all["target_code"].astype(str)
    prepared_all = prepared_all.merge(
        eligibility_lookup,
        on=["target_date", "target_code"],
        how="left",
        validate="many_to_one",
        indicator="_eligibility_key_merge",
    )
    unregistered_outputs = ~prepared_all[
        "_eligibility_key_merge"
    ].astype(str).eq("both")
    if unregistered_outputs.any():
        raise ValueError("SA3 配置输出存在资格机会全集之外的未登记键")
    raw_output_counts = prepared_all.groupby(
        "configuration_id", observed=True
    ).size()
    removed_legal_counts = (
        prepared_all.loc[
            prepared_all["eligibility_status"].astype(str).eq(
                "legally_excluded"
            )
        ]
        .groupby("configuration_id", observed=True)
        .size()
    )
    prepared_all = prepared_all.loc[
        prepared_all["eligibility_status"].astype(str).eq("predictable")
    ].copy()
    daily_eligibility = (
        eligibility.assign(
            _no_in_edges=exclusion_reason.eq("no_in_edges"),
            _complete_case_insufficient=exclusion_reason.eq(
                "complete_case_insufficient"
            ),
            _required_observation_missing=exclusion_reason.eq(
                "required_observation_missing"
            ),
            _predictable=predictable_mask,
        )
        .groupby("prediction_date", observed=True)
        .agg(
            opportunity_count=("target_code", "size"),
            no_in_edges_count=("_no_in_edges", "sum"),
            complete_case_insufficient_count=(
                "_complete_case_insufficient",
                "sum",
            ),
            required_observation_missing_count=(
                "_required_observation_missing",
                "sum",
            ),
            predictable_count=("_predictable", "sum"),
        )
        .reset_index()
    )
    audited_universe_count = eligibility["target_code"].astype(str).nunique()
    if (
        ~daily_eligibility["opportunity_count"].eq(audited_universe_count)
    ).any() or daily_eligibility["predictable_count"].eq(0).any():
        raise ValueError("SA3 机会全集不完整或存在整日动态可预测键为零")
    for count_column in (
        "no_in_edges_count",
        "complete_case_insufficient_count",
        "required_observation_missing_count",
        "predictable_count",
    ):
        daily_eligibility[count_column.replace("_count", "_rate")] = (
            daily_eligibility[count_column]
            / daily_eligibility["opportunity_count"]
        )
    daily_eligibility["audit_level"] = "daily_dynamic_eligibility"
    daily_eligibility["model"] = pd.NA
    daily_eligibility["configuration_id"] = pd.NA
    daily_eligibility["status"] = "complete"
    configuration_audit_rows: list[dict[str, Any]] = []
    for row in grid.to_dict(orient="records"):
        configuration_id = str(row["configuration_id"])
        scope = prepared_all.loc[
            prepared_all["configuration_id"].astype(str).eq(configuration_id)
        ].copy()
        scope["target_date"] = pd.to_datetime(
            scope["target_date"], errors="raise"
        ).dt.normalize()
        if scope.duplicated(["target_date", "target_code"]).any():
            raise ValueError(f"SA3 {configuration_id} 存在重复样本键")
        observed_keys = scope.loc[
            :, ["target_date", "target_code"]
        ].drop_duplicates()
        audit_scope = model_audit_all.loc[
            model_audit_all["configuration_id"].astype(str).eq(
                configuration_id
            )
        ]
        skipped_rows = int(
            (~audit_scope["status"].astype(str).eq("fitted")).sum()
        )
        audit_rows = len(audit_scope)
        key_comparison = expected_keys.merge(
            observed_keys,
            on=["target_date", "target_code"],
            how="outer",
            indicator=True,
            validate="one_to_one",
        )
        keys_complete = key_comparison["_merge"].astype(str).eq("both").all()
        configuration_audit_rows.append(
            {
                "audit_level": "configuration",
                "model": str(row["model"]),
                "configuration_id": configuration_id,
                "opportunity_count": len(eligibility),
                "no_in_edges_count": int(
                    exclusion_reason.eq("no_in_edges").sum()
                ),
                "complete_case_insufficient_count": int(
                    exclusion_reason.eq("complete_case_insufficient").sum()
                ),
                "required_observation_missing_count": int(
                    exclusion_reason.eq("required_observation_missing").sum()
                ),
                "predictable_count": len(expected_keys),
                "prediction_row_count": len(scope),
                "raw_prediction_row_count": int(
                    raw_output_counts.get(configuration_id, 0)
                ),
                "legal_exclusion_output_row_count": int(
                    removed_legal_counts.get(configuration_id, 0)
                ),
                "prediction_key_count": int(
                    scope[["target_date", "target_code"]]
                    .drop_duplicates()
                    .shape[0]
                ),
                "prediction_key_sha256": _sa3_sample_key_sha256(scope),
                "expected_prediction_key_sha256": expected_key_hash,
                "missing_predictable_key_count": int(
                    key_comparison["_merge"].astype(str).eq("left_only").sum()
                ),
                "unexpected_prediction_key_count": int(
                    key_comparison["_merge"].astype(str).eq("right_only").sum()
                ),
                "model_audit_row_count": audit_rows,
                "skipped_model_audit_row_count": skipped_rows,
                "skipped_model_audit_ratio": (
                    skipped_rows / audit_rows if audit_rows else np.nan
                ),
                "dynamic_key_set_complete": keys_complete,
                "sample_keys_identical_across_configurations": pd.NA,
                "status": "complete" if keys_complete else "incomplete",
            }
        )
    configuration_audit = pd.DataFrame(configuration_audit_rows)
    common_hashes = set(configuration_audit["prediction_key_sha256"].astype(str))
    keys_identical = common_hashes == {expected_key_hash}
    configuration_audit[
        "sample_keys_identical_across_configurations"
    ] = keys_identical
    if not keys_identical:
        configuration_audit["status"] = "incomplete"
    failed_configurations = configuration_audit.loc[
        ~configuration_audit["status"].eq("complete")
    ]
    if not failed_configurations.empty:
        raise ValueError(
            "SA3 逐配置共同样本/覆盖门禁失败: "
            f"{failed_configurations.head(20).to_dict(orient='records')}"
        )

    run_parts: list[pd.DataFrame] = []
    for row in grid.to_dict(orient="records"):
        configuration_id = str(row["configuration_id"])
        prepared = prepared_all.loc[
            prepared_all["configuration_id"].astype(str).eq(configuration_id)
        ].copy()
        portfolio = run_sensitivity_portfolios(
            prepared,
            group_columns=(
                "configuration_id",
                "feature_set",
                "direction",
                "target_market",
                "model",
            ),
            config=config,
        )
        metrics = portfolio.sharpe_metrics.copy()
        metrics["parameters_json"] = row["parameters_json"]
        run_parts.append(metrics)
    runs = pd.concat(run_parts, ignore_index=True)
    summary, model_grid_audit = summarize_hyperparameter_sharpe(
        runs,
        config=config,
    )
    grid_audit = pd.concat(
        [
            daily_eligibility,
            configuration_audit,
            model_grid_audit.assign(audit_level="model_grid"),
        ],
        ignore_index=True,
        sort=False,
    )
    r2 = build_r2_comparison(summary, config=config)
    store.write_frame("m5_sa3_runs", runs)
    store.write_frame("m5_sa3_grid_audit", grid_audit)
    store.write_frame("m5_sa3_model_audit", model_audit_all)
    store.write_frame("m5_r2_comparison", r2)
    return runs, grid_audit, model_audit_all, r2


def _run_sa4(
    m1: M1Tables,
    store: CheckpointStore,
    *,
    config: M5Config,
) -> tuple[pd.DataFrame, pd.DataFrame]:
    local_m3 = replace(
        config.m4.m3,
        prediction=replace(
            config.m4.m3.prediction,
            source_value_column="opcl_excess",
        ),
    )
    local_config = replace(config, m4=replace(config.m4, m3=local_m3))
    snapshots = list(
        iter_rolling_graph_snapshots(
            m1.us_panel,
            m1.cn_panel,
            m1.calendar_mapping,
            direction="US_TO_CN",
            prediction_start=local_m3.prediction.prediction_start,
            prediction_end=local_m3.prediction.prediction_end,
            source_value_column="opcl_excess",
            target_value_column=local_m3.prediction.target_value_column,
            config=local_m3.pipeline,
        )
    )
    cross_predictions = _predict_snapshots(snapshots, m1, config=local_config)
    cross = prepare_portfolio_predictions(
        cross_predictions,
        m1.cn_panel,
        target_market="CN",
        specification=local_config.m4.baseline.cross_market_specification,
        feature_set="opcl_excess",
        config=local_config.m4,
    )
    same_all = _collect_same_market_predictions(
        m1,
        source_value_column="opcl_excess",
        config=local_config,
    )
    same_predictions = (
        same_all.predictions
        if isinstance(same_all, SpecificationPredictionTables)
        else same_all
    )
    same = same_predictions.loc[
        same_predictions["target_market"].eq("CN")
    ].copy()
    non_graph_all = _collect_non_graph_predictions(
        m1,
        source_value_column="opcl_excess",
        config=local_config,
    )
    non_graph_predictions = (
        non_graph_all.predictions
        if isinstance(non_graph_all, SpecificationPredictionTables)
        else non_graph_all
    )
    non_graph = non_graph_predictions.loc[
        non_graph_predictions["target_market"].eq("CN")
    ].copy()
    result = run_cn_three_specification_portfolios(
        {
            local_config.m4.baseline.cross_market_specification: cross,
            local_config.m4.baseline.same_market_graph_specification: same,
            local_config.m4.baseline.non_graph_specification: non_graph,
        },
        config=local_config,
    )
    sharpe = result.portfolios.sharpe_metrics.copy()
    raw_checkpoint_sharpe = sharpe["annualized_sharpe"].copy()
    sharpe["annualized_sharpe"] = pd.to_numeric(
        raw_checkpoint_sharpe,
        errors="coerce",
    ).astype("float64")
    sharpe["annualized_sharpe_conversion_failed"] = (
        raw_checkpoint_sharpe.notna()
        & sharpe["annualized_sharpe"].isna()
    )
    sharpe["annualized_sharpe_conversion_failure_count"] = int(
        sharpe["annualized_sharpe_conversion_failed"].sum()
    )
    for metric_column in ("pnl_mean", "pnl_std", "total_pnl"):
        sharpe[metric_column] = pd.to_numeric(
            sharpe[metric_column],
            errors="raise",
        ).astype("float64")
    sharpe["observation_count"] = pd.to_numeric(
        sharpe["observation_count"],
        errors="raise",
    ).astype("int64")
    sharpe["sharpe_status"] = sharpe["sharpe_status"].astype(str)
    keys = ["feature_set", "model", "portfolio"]
    baseline = sharpe.loc[
        sharpe["specification"].eq(
            local_config.m4.baseline.non_graph_specification
        ),
        [
            *keys,
            "annualized_sharpe",
            "sharpe_status",
            "annualized_sharpe_conversion_failed",
        ],
    ].rename(
        columns={
            "annualized_sharpe": "non_graph_sharpe",
            "sharpe_status": "non_graph_sharpe_status",
            "annualized_sharpe_conversion_failed": (
                "non_graph_sharpe_conversion_failed"
            ),
        }
    )
    delta = sharpe.merge(baseline, on=keys, how="left", validate="many_to_one")
    raw_sharpe = delta["annualized_sharpe"].copy()
    raw_non_graph_sharpe = delta["non_graph_sharpe"].copy()
    delta["annualized_sharpe"] = pd.to_numeric(
        raw_sharpe,
        errors="coerce",
    ).astype("float64")
    delta["non_graph_sharpe"] = pd.to_numeric(
        raw_non_graph_sharpe,
        errors="coerce",
    ).astype("float64")
    current_conversion_failures = int(
        delta["annualized_sharpe_conversion_failed"].sum()
    )
    baseline_conversion_failures = int(
        delta["non_graph_sharpe_conversion_failed"].sum()
    )
    delta["annualized_sharpe_conversion_failure_count"] = (
        current_conversion_failures
    )
    delta["non_graph_sharpe_conversion_failure_count"] = (
        baseline_conversion_failures
    )
    delta["finite_sharpe_pair"] = (
        np.isfinite(delta["annualized_sharpe"])
        & np.isfinite(delta["non_graph_sharpe"])
    )
    delta["delta_type"] = (
        delta["specification"].astype(str) + "_minus_non_graph"
    )
    delta["paired_complete"] = (
        delta["finite_sharpe_pair"]
        & delta["sharpe_status"].eq("ok")
        & delta["non_graph_sharpe_status"].eq("ok")
    )
    delta["delta_status"] = np.select(
        [
            (
                delta["annualized_sharpe_conversion_failed"]
                & delta["non_graph_sharpe_conversion_failed"]
            ),
            delta["annualized_sharpe_conversion_failed"],
            delta["non_graph_sharpe_conversion_failed"],
            ~delta["finite_sharpe_pair"],
            ~delta["sharpe_status"].eq("ok"),
            ~delta["non_graph_sharpe_status"].eq("ok"),
        ],
        [
            "both_sharpe_conversions_failed",
            "annualized_sharpe_conversion_failed",
            "non_graph_sharpe_conversion_failed",
            "nonfinite_sharpe_pair",
            "annualized_sharpe_status_not_ok",
            "non_graph_sharpe_status_not_ok",
        ],
        default="ok",
    )
    delta["sharpe_delta"] = np.where(
        delta["paired_complete"],
        delta["annualized_sharpe"] - delta["non_graph_sharpe"],
        np.nan,
    ).astype("float64")
    finite_pair_count = int(delta["finite_sharpe_pair"].sum())
    complete_pair_count = int(delta["paired_complete"].sum())
    delta["finite_sharpe_pair_count"] = finite_pair_count
    delta["complete_sharpe_pair_count"] = complete_pair_count
    comparability_audit = result.comparability_audit.copy()
    comparability_audit["annualized_sharpe_conversion_failure_count"] = (
        current_conversion_failures
    )
    comparability_audit["non_graph_sharpe_conversion_failure_count"] = (
        baseline_conversion_failures
    )
    comparability_audit["finite_sharpe_pair_count"] = finite_pair_count
    comparability_audit["complete_sharpe_pair_count"] = complete_pair_count
    store.write_frame("m5_sa4_sharpe", sharpe)
    store.write_frame("m5_sa4_delta", delta)
    store.write_frame(
        "m5_sa4_model_completeness_audit",
        result.model_completeness_audit,
    )
    store.write_frame("m5_sa4_comparability_audit", comparability_audit)
    store.write_frame(
        "m5_sa4_execution_timing_audit",
        result.execution_timing_audit,
    )
    return sharpe, delta


def _run_sa5(
    m1: M1Tables,
    m4: M4Tables,
    store: CheckpointStore,
    *,
    config: M5Config,
) -> tuple[pd.DataFrame, pd.DataFrame]:
    conditioned, audit = attach_nested_us_shock_quantiles(
        _us_to_cn_prepared(m4),
        m1.us_market_returns,
        config=config,
    )
    result = run_sensitivity_portfolios(
        conditioned,
        group_columns=(
            "shock_quantile",
            "feature_set",
            "direction",
            "target_market",
            "model",
        ),
        portfolio_fractions=(("all_stocks", 1.0),),
        config=config,
    )
    store.write_frame("m5_sa5_sharpe", result.sharpe_metrics)
    store.write_frame("m5_sa5_audit", audit)
    return result.sharpe_metrics, audit


def _run_sa6(
    m4: M4Tables,
    store: CheckpointStore,
    *,
    config: M5Config,
) -> tuple[pd.DataFrame, pd.DataFrame, pd.DataFrame]:
    industry_mapping = load_citic_level_one_industry(config=config)
    industry, audit = attach_target_industry(
        _us_to_cn_prepared(m4),
        industry_mapping.mapping,
    )
    result = run_sensitivity_portfolios(
        industry,
        group_columns=(
            "target_industry",
            "feature_set",
            "direction",
            "target_market",
            "model",
        ),
        portfolio_fractions=(("all_industry_stocks", 1.0),),
        config=config,
    )
    store.write_frame("m5_sa6_sharpe", result.sharpe_metrics)
    store.write_frame("m5_sa6_audit", audit)
    store.write_frame(
        "m5_sa6_mapping_summary_audit",
        industry_mapping.summary_audit,
    )
    store.write_frame(
        "m5_sa6_duplicate_rows_audit",
        industry_mapping.duplicate_rows,
    )
    mapping_audit = pd.concat(
        [
            industry_mapping.summary_audit.assign(audit_section="summary"),
            audit.assign(audit_section="prediction_coverage"),
        ],
        ignore_index=True,
        sort=False,
    )
    store.write_frame("m5_sa6_combined_audit", mapping_audit)
    return result.sharpe_metrics, mapping_audit, industry_mapping.duplicate_rows


def _joint_snapshot_pairs(
    m1: M1Tables,
    *,
    us_feature_column: str,
    cn_feature_column: str,
    config: M5Config,
) -> list[tuple[GraphSnapshot, GraphSnapshot]]:
    us_snapshots = list(
        iter_rolling_graph_snapshots(
            m1.us_panel,
            m1.cn_panel,
            m1.calendar_mapping,
            direction="US_TO_CN",
            prediction_start=config.m4.m3.prediction.prediction_start,
            prediction_end=config.m4.m3.prediction.prediction_end,
            source_value_column=us_feature_column,
            target_value_column=config.m4.m3.prediction.target_value_column,
            config=config.m4.m3.pipeline,
        )
    )
    same_mapping = build_same_market_calendar_mapping(
        m1.cn_panel,
        market="CN",
        config=config.m4,
    )
    same_direction = DirectionConfig(
        name=config.m4.baseline.cn_same_market_direction,
        source_market="CN",
        target_market="CN",
        lag_parameter=config.m4.baseline.same_market_lag,
    )
    same_pipeline = replace(
        config.m4.m3.pipeline,
        calendar=replace(
            config.m4.m3.pipeline.calendar,
            directions=(same_direction,),
        ),
    )
    cn_snapshots = list(
        iter_rolling_graph_snapshots(
            m1.cn_panel,
            m1.cn_panel,
            same_mapping,
            direction=same_direction.name,
            prediction_start=config.m4.m3.prediction.prediction_start,
            prediction_end=config.m4.m3.prediction.prediction_end,
            source_value_column=cn_feature_column,
            target_value_column=config.m4.m3.prediction.target_value_column,
            config=same_pipeline,
        )
    )
    if len(us_snapshots) != len(cn_snapshots):
        raise ValueError("SA7 US→CN 与 CN→CN 快照数量不一致")
    return list(zip(us_snapshots, cn_snapshots, strict=True))


def _run_sa7(
    m1: M1Tables,
    m4: M4Tables,
    store: CheckpointStore,
    *,
    config: M5Config,
) -> tuple[
    pd.DataFrame,
    pd.DataFrame,
    pd.DataFrame,
    pd.DataFrame,
    pd.DataFrame,
]:
    sharpe_parts: list[pd.DataFrame] = []
    baseline_sharpe_parts: list[pd.DataFrame] = []
    audit_parts: list[pd.DataFrame] = []
    factories = build_base_model_factories(config.m4.m3)
    expected_models = frozenset(config.m4.m3.prediction.all_model_names)
    baseline_source = _us_to_cn_prepared(m4)

    def complete_keys(frame: pd.DataFrame) -> pd.DataFrame:
        keys = ["target_date", "target_code"]
        numeric = frame[["prediction", "actual_return", "historical_liquidity"]].apply(
            pd.to_numeric,
            errors="coerce",
        )
        working = frame.assign(
            _finite=np.isfinite(
                numeric.to_numpy(dtype=float, na_value=np.nan)
            ).all(axis=1)
        )
        grouped = (
            working.groupby(keys, observed=True)
            .agg(
                model_set=("model", lambda values: frozenset(values.astype(str))),
                finite_count=("_finite", "sum"),
                row_count=("model", "size"),
            )
            .reset_index()
        )
        return grouped.loc[
            grouped["model_set"].map(lambda values: values == expected_models)
            & grouped["finite_count"].eq(len(expected_models))
            & grouped["row_count"].eq(len(expected_models)),
            keys,
        ]

    for combination, us_feature, cn_feature in (
        config.sensitivity.feature_combinations
    ):
        prediction_parts: list[pd.DataFrame] = []
        for us_snapshot, cn_snapshot in _joint_snapshot_pairs(
            m1,
            us_feature_column=us_feature,
            cn_feature_column=cn_feature,
            config=config,
        ):
            result = predict_joint_graph_snapshots(
                us_snapshot,
                cn_snapshot,
                m1.us_panel,
                m1.cn_panel,
                us_feature_column=us_feature,
                cn_feature_column=cn_feature,
                model_factories=factories,
                config=config,
            )
            if not result.predictions.empty:
                prediction_parts.append(result.predictions)
            audit_parts.append(
                result.training_audit.assign(feature_combination=combination)
            )
        predictions = (
            pd.concat(prediction_parts, ignore_index=True)
            if prediction_parts
            else pd.DataFrame()
        )
        prepared = prepare_portfolio_predictions(
            predictions,
            m1.cn_panel,
            target_market="CN",
            specification="cross_market_graph",
            feature_set=combination,
            config=config.m4,
        )
        prepared["feature_combination"] = combination
        joint_keys = complete_keys(prepared)
        baseline_keys = complete_keys(baseline_source)
        common_keys = joint_keys.merge(
            baseline_keys,
            on=["target_date", "target_code"],
            how="inner",
            validate="one_to_one",
        )
        joint_aligned = prepared.merge(
            common_keys,
            on=["target_date", "target_code"],
            how="inner",
            validate="many_to_one",
        )
        baseline_aligned = baseline_source.merge(
            common_keys,
            on=["target_date", "target_code"],
            how="inner",
            validate="many_to_one",
        )
        baseline_aligned["feature_combination"] = combination
        if joint_aligned.empty or baseline_aligned.empty:
            raise ValueError(f"SA7 {combination} 与 US-only 基线无共同完整观测")
        label_check = joint_aligned.loc[
            :, ["target_date", "target_code", "model", "actual_return"]
        ].merge(
            baseline_aligned.loc[
                :, ["target_date", "target_code", "model", "actual_return"]
            ],
            on=["target_date", "target_code", "model"],
            how="inner",
            validate="one_to_one",
            suffixes=("_joint", "_baseline"),
        )
        label_mismatch = (
            label_check["actual_return_joint"]
            .sub(label_check["actual_return_baseline"])
            .abs()
            .gt(config.m4.portfolio.actual_return_match_tolerance)
        )
        if label_mismatch.any():
            raise ValueError(f"SA7 {combination} 联合与基线标签不一致")
        joint_portfolio = run_sensitivity_portfolios(
            joint_aligned,
            group_columns=(
                "feature_combination",
                "direction",
                "target_market",
                "model",
            ),
            config=config,
        )
        baseline_portfolio = run_sensitivity_portfolios(
            baseline_aligned,
            group_columns=(
                "feature_combination",
                "direction",
                "target_market",
                "model",
            ),
            config=config,
        )
        sharpe_parts.append(joint_portfolio.sharpe_metrics)
        baseline_sharpe_parts.append(baseline_portfolio.sharpe_metrics)
        audit_parts.append(
            pd.DataFrame(
                [
                    {
                        "feature_combination": combination,
                        "joint_complete_observation_count": len(joint_keys),
                        "baseline_complete_observation_count": len(baseline_keys),
                        "common_observation_count": len(common_keys),
                        "joint_prediction_rows": len(joint_aligned),
                        "baseline_prediction_rows": len(baseline_aligned),
                        "labels_identical": not label_mismatch.any(),
                        "portfolio_contract": (
                            "same common target_date/target_code, label, "
                            "liquidity, capacity and qr fractions"
                        ),
                    }
                ]
            )
        )
    sharpe = pd.concat(sharpe_parts, ignore_index=True)
    baseline_sharpe = pd.concat(baseline_sharpe_parts, ignore_index=True)
    audit = pd.concat(audit_parts, ignore_index=True, sort=False)
    comparison = build_sa7_us_only_comparison(
        sharpe,
        baseline_sharpe,
        config=config,
    )
    store.write_frame("m5_sa7_sharpe", sharpe)
    store.write_frame("m5_sa7_audit", audit)
    store.write_frame("m5_sa7_us_only_sharpe", baseline_sharpe)
    store.write_frame("m5_sa7_comparison", comparison.paired)
    store.write_frame("m5_sa7_pairing_audit", comparison.audit)
    store.write_frame("m5_sa7_comparison_summary", comparison.summary)
    return (
        sharpe,
        audit,
        comparison.paired,
        comparison.audit,
        comparison.summary,
    )


def _stage_m5(
    m1: M1Tables,
    m4: M4Tables,
    store: CheckpointStore,
    *,
    config: M5Config,
) -> dict[str, pd.DataFrame]:
    frame_map = {
        "sa1_edge": "m5_sa1_summary",
        "sa1_edge_audit": "m5_sa1_audit",
        "sa1_median_audit": "m5_sa1_median_audit",
        "sa2_lag": "m5_sa2_summary",
        "sa2_lag_audit": "m5_sa2_audit",
        "sa2_median_audit": "m5_sa2_median_audit",
        "sa3_runs": "m5_sa3_runs",
        "sa3_grid_audit": "m5_sa3_grid_audit",
        "sa3_model_audit": "m5_sa3_model_audit",
        "r2_comparison": "m5_r2_comparison",
        "sa4_opcl": "m5_sa4_sharpe",
        "sa4_opcl_delta": "m5_sa4_delta",
        "sa4_model_completeness_audit": (
            "m5_sa4_model_completeness_audit"
        ),
        "sa4_comparability_audit": "m5_sa4_comparability_audit",
        "sa4_execution_timing_audit": "m5_sa4_execution_timing_audit",
        "sa5_shock": "m5_sa5_sharpe",
        "sa5_shock_audit": "m5_sa5_audit",
        "sa6_industry": "m5_sa6_sharpe",
        "sa6_industry_audit": "m5_sa6_combined_audit",
        "sa6_duplicate_rows_audit": "m5_sa6_duplicate_rows_audit",
        "sa7_joint": "m5_sa7_sharpe",
        "sa7_joint_audit": "m5_sa7_audit",
        "sa7_us_only": "m5_sa7_us_only_sharpe",
        "sa7_comparison": "m5_sa7_comparison",
        "sa7_pairing_audit": "m5_sa7_pairing_audit",
        "sa7_comparison_summary": "m5_sa7_comparison_summary",
        "sensitivity_directional": "m5_sensitivity_directional",
    }
    context = store.build_context(
        "m5",
        config=config,
        input_manifest=_m5_input_manifest(store, config),
        upstream_stages=("m1", "m2_m3", "m4"),
    )
    if (
        store.stage_complete("m5", context)
        and all(store.has_frame(name) for name in frame_map.values())
    ):
        return {
            output_name: store.read_frame(checkpoint_name)
            for output_name, checkpoint_name in frame_map.items()
        }

    sa1, sa1_audit, sa1_median_audit = _run_sa1(
        m1, m4, store, config=config
    )
    sa2, sa2_audit, sa2_median_audit = _run_sa2(
        m1, m4, store, config=config
    )
    sa3_runs, sa3_audit, sa3_model_audit, r2 = _run_sa3(
        m1,
        store,
        config=config,
    )
    sa4, sa4_delta = _run_sa4(m1, store, config=config)
    sa5, sa5_audit = _run_sa5(m1, m4, store, config=config)
    sa6, sa6_audit, sa6_duplicates = _run_sa6(
        m4, store, config=config
    )
    (
        sa7,
        sa7_audit,
        sa7_comparison,
        sa7_pairing_audit,
        sa7_summary,
    ) = _run_sa7(m1, m4, store, config=config)
    directional = build_directional_sensitivity_comparison(
        {
            "SA1": sa1,
            "SA2": sa2,
            "SA3": sa3_runs,
            "SA4": sa4,
            "SA5": sa5,
            "SA6": sa6,
            "SA7": sa7_comparison,
        }
    )
    store.write_frame("m5_sensitivity_directional", directional)
    results = {
        "sa1_edge": sa1,
        "sa1_edge_audit": sa1_audit,
        "sa1_median_audit": sa1_median_audit,
        "sa2_lag": sa2,
        "sa2_lag_audit": sa2_audit,
        "sa2_median_audit": sa2_median_audit,
        "sa3_runs": sa3_runs,
        "sa3_grid_audit": sa3_audit,
        "sa3_model_audit": sa3_model_audit,
        "r2_comparison": r2,
        "sa4_opcl": sa4,
        "sa4_opcl_delta": sa4_delta,
        "sa4_model_completeness_audit": store.read_frame(
            "m5_sa4_model_completeness_audit"
        ),
        "sa4_comparability_audit": store.read_frame(
            "m5_sa4_comparability_audit"
        ),
        "sa4_execution_timing_audit": store.read_frame(
            "m5_sa4_execution_timing_audit"
        ),
        "sa5_shock": sa5,
        "sa5_shock_audit": sa5_audit,
        "sa6_industry": sa6,
        "sa6_industry_audit": sa6_audit,
        "sa6_duplicate_rows_audit": sa6_duplicates,
        "sa7_joint": sa7,
        "sa7_joint_audit": sa7_audit,
        "sa7_us_only": store.read_frame("m5_sa7_us_only_sharpe"),
        "sa7_comparison": sa7_comparison,
        "sa7_pairing_audit": sa7_pairing_audit,
        "sa7_comparison_summary": sa7_summary,
        "sensitivity_directional": directional,
    }
    output_paths = [
        store.frame_path(name)
        for name in frame_map.values()
    ]
    store.mark_complete(
        "m5",
        {
            "r2_cells": len(r2),
            "directional_elements": len(directional),
        },
        context=context,
        output_paths=output_paths,
    )
    return results


def _validate_fig03_training_mapping(
    training_mapping: pd.DataFrame,
    prediction_mapping: pd.DataFrame,
    *,
    requested: pd.Timestamp,
    direction: DirectionConfig,
    config: M5Config,
) -> dict[str, Any]:
    """在 FIG3 构图前逐行验证 250 条历史因果配对。"""
    required = {
        "target_date",
        "source_date",
        "causal_available",
        "target_open_utc",
        "source_close_utc",
    }
    missing = sorted(required - set(training_mapping.columns))
    if missing:
        raise KeyError(f"FIG3 历史配对缺少因果字段: {missing}")
    if len(prediction_mapping) != 1:
        raise ValueError("FIG3 精确图日必须存在唯一目标开盘映射")
    missing_prediction = sorted(
        {"target_date", "target_open_utc"} - set(prediction_mapping.columns)
    )
    if missing_prediction:
        raise KeyError(f"FIG3 图日映射缺少字段: {missing_prediction}")

    window_days = config.m4.m3.pipeline.returns.training_window_days
    causal_audit = audit_snapshot_causality(
        direction=direction,
        anchor_date=requested,
        training_mapping=training_mapping,
        prediction_mapping=prediction_mapping,
        window_days=window_days,
        update_days=config.m4.m3.pipeline.returns.graph_update_days,
    ).iloc[0]
    target_dates = pd.to_datetime(
        training_mapping["target_date"], errors="raise"
    ).dt.normalize()
    source_dates = pd.to_datetime(
        training_mapping["source_date"], errors="coerce"
    ).dt.normalize()
    causal_available = training_mapping["causal_available"].astype(
        "boolean"
    )
    target_open_utc = pd.to_datetime(
        training_mapping["target_open_utc"],
        errors="coerce",
        utc=True,
    )
    source_close_utc = pd.to_datetime(
        training_mapping["source_close_utc"],
        errors="coerce",
        utc=True,
    )
    graph_open_utc = pd.to_datetime(
        prediction_mapping.iloc[0]["target_open_utc"],
        errors="raise",
        utc=True,
    )
    target_dates_before_graph = target_dates.lt(requested)
    source_dates_before_graph = source_dates.lt(requested)
    source_before_target_open = source_close_utc.lt(target_open_utc)
    source_before_graph_open = source_close_utc.lt(graph_open_utc)
    legal_pair = (
        causal_available.fillna(False)
        & source_dates.notna()
        & target_open_utc.notna()
        & source_close_utc.notna()
        & target_dates_before_graph
        & source_dates_before_graph
        & source_before_target_open
        & source_before_graph_open
    )
    causal_all = bool(causal_available.fillna(False).all())
    source_date_complete = bool(source_dates.notna().all())
    source_timestamp_complete = bool(source_close_utc.notna().all())
    target_timestamp_complete = bool(target_open_utc.notna().all())
    target_dates_strict = bool(target_dates_before_graph.all())
    source_dates_strict = bool(source_dates_before_graph.all())
    source_before_target_all = bool(source_before_target_open.all())
    source_before_graph_all = bool(source_before_graph_open.all())
    legal_pair_count = int(legal_pair.sum())
    audit_count_consistent = (
        int(causal_audit["training_causal_available_rows"]) == window_days
        and int(causal_audit["training_source_date_missing_rows"]) == 0
    )
    if not (
        len(training_mapping) == window_days
        and legal_pair_count == window_days
        and causal_all
        and source_date_complete
        and source_timestamp_complete
        and target_timestamp_complete
        and target_dates_strict
        and source_dates_strict
        and source_before_target_all
        and source_before_graph_all
        and audit_count_consistent
    ):
        failed_rows = training_mapping.loc[~legal_pair].head(20)
        raise AssertionError(
            "FIG3 不具备250条逐行合法历史因果配对: "
            f"legal_pair_count={legal_pair_count}, "
            f"failed={failed_rows.to_dict(orient='records')}"
        )
    return {
        "training_pair_count": legal_pair_count,
        "training_causal_available_count": int(
            causal_available.fillna(False).sum()
        ),
        "training_causal_available_all": causal_all,
        "training_source_date_missing_count": int(source_dates.isna().sum()),
        "training_source_date_complete": source_date_complete,
        "training_source_timestamp_complete": source_timestamp_complete,
        "training_target_timestamp_complete": target_timestamp_complete,
        "training_target_dates_strictly_before_graph_date": (
            target_dates_strict
        ),
        "training_dates_strictly_before_graph_date": target_dates_strict,
        "training_source_dates_strictly_before_graph_date": (
            source_dates_strict
        ),
        "source_close_strictly_before_target_open_all": (
            source_before_target_all
        ),
        "source_close_strictly_before_graph_open_all": (
            source_before_graph_all
        ),
        "training_target_date_start": target_dates.min(),
        "training_target_date_end": target_dates.max(),
        "training_source_date_start": source_dates.min(),
        "training_source_date_end": source_dates.max(),
        "training_target_open_utc_start": target_open_utc.min(),
        "training_target_open_utc_end": target_open_utc.max(),
        "training_source_close_utc_start": source_close_utc.min(),
        "training_source_close_utc_end": source_close_utc.max(),
        "effective_graph_open_utc": graph_open_utc,
    }


def _build_fig03_graph_evidence(
    m1: M1Tables,
    *,
    config: M5Config,
) -> tuple[pd.DataFrame, pd.DataFrame]:
    """在论文精确日一次性构图，输出全量 B1 代理而非行业限额原图。"""
    requested = pd.Timestamp(
        config.reporting.fig03_requested_anchor_date
    ).normalize()
    direction = config.m4.m3.pipeline.calendar.direction("US_TO_CN")
    direction_mapping = m1.calendar_mapping.loc[
        m1.calendar_mapping["direction"].astype(str).eq(direction.name)
    ].copy()
    prediction_mapping = direction_mapping.loc[
        pd.to_datetime(
            direction_mapping["target_date"], errors="raise"
        ).dt.normalize().eq(requested)
    ].copy()
    prevalidated_training_mapping = select_training_mapping(
        direction_mapping,
        anchor_date=requested,
        window_days=config.m4.m3.pipeline.returns.training_window_days,
    )
    causal_contract = _validate_fig03_training_mapping(
        prevalidated_training_mapping,
        prediction_mapping,
        requested=requested,
        direction=direction,
        config=config,
    )
    snapshots = list(
        iter_rolling_graph_snapshots(
            m1.us_panel,
            m1.cn_panel,
            m1.calendar_mapping,
            direction="US_TO_CN",
            prediction_start=requested,
            prediction_end=requested,
            source_value_column=config.m4.m3.prediction.source_value_column,
            target_value_column=config.m4.m3.prediction.target_value_column,
            config=config.m4.m3.pipeline,
        )
    )
    if len(snapshots) != 1:
        raise ValueError(
            "FIG3 精确日无法形成唯一合法图快照，不允许回退邻近日期"
        )
    snapshot = snapshots[0]
    if snapshot.anchor_date != requested:
        raise AssertionError("FIG3 一次性构图未严格使用论文日期")
    training_mapping = snapshot.training_mapping.copy()
    pd.testing.assert_frame_equal(
        training_mapping.loc[
            :,
            [
                "target_date",
                "source_date",
                "causal_available",
                "target_open_utc",
                "source_close_utc",
            ],
        ].reset_index(drop=True),
        prevalidated_training_mapping.loc[
            :,
            [
                "target_date",
                "source_date",
                "causal_available",
                "target_open_utc",
                "source_close_utc",
            ],
        ].reset_index(drop=True),
        check_dtype=False,
        check_exact=True,
    )
    source_by_target = snapshot.adjacency.to_dense_frame()
    matrix = source_by_target.T
    matrix.index.name = "cn_target_code"
    matrix.columns.name = "us_source_code"
    row_nodes = [str(value) for value in matrix.index]
    column_nodes = [str(value) for value in matrix.columns]

    def ordered_list_hash(values: list[str]) -> str:
        return hashlib.sha256("\x1f".join(values).encode("utf-8")).hexdigest()

    matrix_payload = (
        "\x1e".join(row_nodes).encode("utf-8")
        + b"\x00"
        + "\x1e".join(column_nodes).encode("utf-8")
        + b"\x00"
        + np.ascontiguousarray(
            matrix.to_numpy(dtype=np.float64).astype("<f8", copy=False)
        ).tobytes()
    )
    audit = pd.DataFrame(
        [
            {
                "figure_id": "FIG3",
                "paper_requested_date": requested,
                "effective_graph_date": snapshot.anchor_date,
                "used_nearest_update_anchor": False,
                "paper_universe": config.reporting.fig03_paper_universe,
                "implemented_universe": (
                    config.reporting.fig03_implemented_universe
                ),
                "industry_limit_applied": False,
                "missing_dependency": (
                    config.reporting.fig03_missing_dependency
                ),
                "reproduction_status": (
                    config.reporting.fig03_reproduction_status
                ),
                "assumption_id": config.reporting.fig03_assumption_id,
                "direction": snapshot.direction,
                "source_market": snapshot.source_market,
                "target_market": snapshot.target_market,
                "matrix_internal_orientation": "source_US_x_target_CN",
                "output_transposed": True,
                "matrix_row_market": "CN",
                "matrix_column_market": "US",
                "matrix_value_contract": "signed t_beta; absent edge=0",
                "row_count": matrix.shape[0],
                "column_count": matrix.shape[1],
                "paper_row_count_cn": (
                    config.reporting.fig03_paper_cn_row_count
                ),
                "paper_column_count_us": (
                    config.reporting.fig03_paper_us_column_count
                ),
                "row_node_ordered_list_json": json.dumps(
                    row_nodes, ensure_ascii=False
                ),
                "column_node_ordered_list_json": json.dumps(
                    column_nodes, ensure_ascii=False
                ),
                "row_node_ordered_list_sha256": ordered_list_hash(row_nodes),
                "column_node_ordered_list_sha256": ordered_list_hash(
                    column_nodes
                ),
                **causal_contract,
                "matrix_data_sha256": hashlib.sha256(
                    matrix_payload
                ).hexdigest(),
                "nonzero_edge_count": snapshot.adjacency.nnz,
            }
        ]
    )
    return matrix, audit


def _build_fig05_long_table(indegree: pd.DataFrame) -> pd.DataFrame:
    """把两个方向的 q25/q50/q75 展开为六条可审计时序。"""
    required = (
        "direction",
        "anchor_date",
        "in_degree_q25",
        "in_degree_q50",
        "in_degree_q75",
    )
    missing = sorted(set(required) - set(indegree.columns))
    if missing:
        raise KeyError(f"FIG5 入度底表缺少字段: {missing}")
    long = indegree.loc[:, list(required)].melt(
        id_vars=["direction", "anchor_date"],
        value_vars=["in_degree_q25", "in_degree_q50", "in_degree_q75"],
        var_name="percentile",
        value_name="in_degree",
    )
    long["percentile"] = long["percentile"].str.replace(
        "in_degree_",
        "",
        regex=False,
    )
    long["series_label"] = (
        long["direction"].astype(str)
        + "/"
        + long["percentile"].astype(str)
    )
    expected_series = {
        f"{direction}/{percentile}"
        for direction in ("US_TO_CN", "CN_TO_US")
        for percentile in ("q25", "q50", "q75")
    }
    observed_series = set(long["series_label"].astype(str))
    if observed_series != expected_series:
        raise ValueError(
            "FIG5 六序列合同不完整: "
            f"actual={sorted(observed_series)}, expected={sorted(expected_series)}"
        )
    return long.sort_values(
        ["direction", "percentile", "anchor_date"]
    ).reset_index(drop=True)


def _build_chart_tables(
    m1: M1Tables,
    m3: M3Tables,
    m4: M4Tables,
    m5: dict[str, pd.DataFrame],
    store: CheckpointStore,
    *,
    config: M5Config,
) -> dict[str, pd.DataFrame]:
    daily_pnl = m4.daily_pnl.copy()
    if not daily_pnl.empty:
        group_columns = [
            column
            for column in (
                "specification",
                "feature_set",
                "direction",
                "target_market",
                "model",
                "portfolio",
            )
            if column in daily_pnl.columns
        ]
        daily_pnl = daily_pnl.sort_values([*group_columns, "target_date"])
        if "daily_pnl_before_cost" in daily_pnl.columns:
            daily_pnl["cumulative_pnl"] = daily_pnl.groupby(
                group_columns,
                observed=True,
            )["daily_pnl_before_cost"].cumsum()
        daily_pnl["series_label"] = daily_pnl[group_columns].astype(str).agg(
            "/".join,
            axis=1,
        )
    sharpe = m4.sharpe_metrics.copy()
    fig08 = sharpe.loc[
        sharpe["direction"].eq("US_TO_CN")
        & sharpe["specification"].eq("cross_market_graph")
    ].copy()
    fig09 = sharpe.loc[
        sharpe["direction"].eq("CN_TO_US")
        & sharpe["specification"].eq("cross_market_graph")
    ].copy()
    fig11 = sharpe.loc[
        sharpe["target_market"].eq("CN")
        & sharpe["specification"].eq("same_market_graph")
    ].copy()
    fig12 = sharpe.loc[sharpe["target_market"].eq("CN")].copy()
    fig13 = build_figure13_long_table(m4.core_paired_comparisons)
    fig10 = daily_pnl.loc[
        daily_pnl["specification"].eq("cross_market_graph")
        & daily_pnl["feature_set"].eq("pvclcl_excess")
        & daily_pnl["direction"].eq("US_TO_CN")
        & daily_pnl["target_market"].eq("CN")
    ].copy()
    if not fig10.empty:
        fig10["series_label"] = fig10[["model", "portfolio"]].astype(str).agg(
            "/".join,
            axis=1,
        )
        if fig10.duplicated(["series_label", "target_date"]).any():
            raise ValueError("Figure 10 model×portfolio×date 存在重复路径")
    edge = m5["sa1_edge"].rename(
        columns={"replacement_fraction": "sensitivity_value"}
    ).assign(series_label=lambda frame: "边替换/" + frame["portfolio"].astype(str))
    lag = m5["sa2_lag"].rename(
        columns={"information_lag": "sensitivity_value"}
    ).assign(series_label=lambda frame: "信息滞后/" + frame["portfolio"].astype(str))
    fig03_matrix, fig03_audit = _build_fig03_graph_evidence(
        m1,
        config=config,
    )
    fig05_long = _build_fig05_long_table(m3.indegree_quantiles)
    return {
        "predictions": m4.prepared_predictions,
        "feature_values": m3.feature_values,
        "prediction_quality": summarize_prediction_quality(m3.predictions),
        "daily_pnl": daily_pnl,
        "sharpe_metrics": sharpe,
        "fig03_graph_matrix": fig03_matrix,
        "fig03_graph_matrix_table": fig03_matrix.reset_index(),
        "fig03_graph_anchor_audit": fig03_audit,
        "fig05_indegree": fig05_long,
        "fig08": fig08,
        "fig09": fig09,
        "fig10_daily_pnl": fig10,
        "fig11": fig11,
        "fig12": fig12,
        "fig13": fig13,
        "fig14_edge_lag": pd.concat([edge, lag], ignore_index=True),
        **m5,
    }


def run_formal_pipeline(
    *,
    config: M5Config = DEFAULT_M5_CONFIG,
) -> Path:
    """运行完整正式流水线；不在 coder 冒烟阶段调用。"""
    warnings.filterwarnings(
        "ignore",
        message=(
            "X does not have valid feature names, but "
            "LGBMRegressor was fitted with feature names"
        ),
        category=UserWarning,
        module=r"sklearn\.utils\.validation",
    )
    store = CheckpointStore(config.reporting.checkpoint_dir)
    started = datetime.now().astimezone()
    model_configuration = build_model_configuration_evidence(config.m4.m3)
    m1 = _stage_m1(store, config=config)
    m3 = _stage_m2_m3(
        m1,
        store,
        load_reporting_frames=False,
        config=config,
    )
    m4 = _stage_m4(m1, m3, store, config=config)
    m5 = _stage_m5(m1, m4, store, config=config)
    m3 = M3Tables(
        *(
            store.read_frame(name)
            for name in (
                "m3_predictions",
                "m3_training_audit",
                "m3_model_audit",
                "m3_feature_values",
                "m2_graph_audit",
                "m2_indegree_quantiles",
                "m3_eligibility_audit",
            )
        )
    )
    tables = _build_chart_tables(m1, m3, m4, m5, store, config=config)
    tables.update(
        {
            "m1_source_schema_audit": m1.source_schema_audit,
            "m1_loaded_panel_audit": m1.loaded_panel_audit,
            "m1_return_panel_audit": m1.return_panel_audit,
            "m1_security_coverage_audit": m1.security_coverage_audit,
            "m1_return_distribution_audit": (
                m1.return_distribution_audit
            ),
            "graph_audit": m3.graph_audit,
            "training_audit": m3.training_audit,
            "model_audit": m3.model_audit,
            "model_completeness_audit": m4.model_completeness_audit,
            "comparability_audit": m4.comparability_audit,
            "execution_timing_audit": m4.execution_timing_audit,
            "formal_coverage_audit": m4.formal_coverage_audit,
            "core_comparison": m4.core_summary,
            "model_configuration": pd.DataFrame(
                [
                    {
                        "evidence_type": "full_get_params_and_versions",
                        "payload_json": json.dumps(
                            model_configuration,
                            ensure_ascii=False,
                            sort_keys=True,
                        ),
                    }
                ]
            ),
        }
    )
    metrics = {
        "report_name": "r007_cross_market_return_forecasting",
        "mode": "formal",
        "seed": config.sensitivity.seed,
        "prediction_rows": len(m3.predictions),
        "formal_coverage_gate_rows": len(m4.formal_coverage_audit),
        "m4_sharpe_rows": len(m4.sharpe_metrics),
        "r2_comparison_cells": len(m5["r2_comparison"]),
        "assumption_scope": [
            "AS4",
            "AS7",
            "AS8",
            "AS15",
            "AS16",
            "AS19",
            "AS24",
            "AS25",
            "AS26",
            "AS27",
        ],
        "fig03_reproduction_status": (
            config.reporting.fig03_reproduction_status
        ),
        "fig03_effective_graph_date": (
            config.reporting.fig03_requested_anchor_date
        ),
        "fig03_used_nearest_update_anchor": False,
        "fig03_missing_dependency": config.reporting.fig03_missing_dependency,
        "fig03_assumption_id": config.reporting.fig03_assumption_id,
    }
    comparison = {
        "R2": m5["r2_comparison"].to_dict(orient="records"),
        "SA_directional": m5["sensitivity_directional"].to_dict(
            orient="records"
        ),
        "R1_R3_R4": m4.core_summary.to_dict(orient="records"),
        "verdict_owner": "verifier_and_check_gates",
    }
    completed = datetime.now().astimezone()
    report = write_formal_outputs(
        tables,
        metrics=metrics,
        comparison=comparison,
        run_log_lines=(
            f"开始时间：{started.isoformat()}",
            f"结束时间：{completed.isoformat()}",
            "模式：formal，无静默缩样本",
            f"确定性随机种子：{config.sensitivity.seed}",
            f"检查点目录：{config.reporting.checkpoint_dir}",
            "阶段：m1→m2/m3→m4→m5→reporting",
            "训练损失图：not_applicable，传统模型无统一 epoch 损失",
            (
                "FIG3：2021-10-21 精确日全量 B1 代理图，"
                "degraded/substitute；D9 缺失；AS27；未使用邻近锚点"
            ),
            "裁决边界：本入口仅产观测与证据，不宣布复现通过",
        ),
        output_dir=config.reporting.output_dir,
        config=config,
    )
    return report.files[-1]


def run_smoke(
    *,
    config: M5Config = DEFAULT_M5_CONFIG,
) -> Path:
    """动态覆盖 CA-A01–CA-A08；只用极小合成样本和只读行业映射。"""
    smoke_dir = config.reporting.smoke_dir
    if smoke_dir.exists():
        shutil.rmtree(smoke_dir)
    smoke_dir.mkdir(parents=True, exist_ok=True)
    regression_rows: list[dict[str, Any]] = []

    def record(check_id: str, condition: bool, evidence: str) -> None:
        if not condition:
            raise AssertionError(f"{check_id} 动态回归失败: {evidence}")
        regression_rows.append(
            {
                "check_id": check_id,
                "status": "passed",
                "dynamic_evidence": evidence,
            }
        )

    edges = pd.DataFrame(
        {
            "source_code": ["U1", "U2", "U3"],
            "target_code": ["C1", "C2", "C3"],
            "t_beta": [2.4, -2.7, 3.1],
        }
    )
    from .sensitivity_pipeline import randomly_replace_graph_edges

    randomized = randomly_replace_graph_edges(
        edges,
        source_codes=("U1", "U2", "U3", "U4", "U5", "U6"),
        replacement_fraction=1.0,
        seed=config.sensitivity.seed,
    )
    grid = build_hyperparameter_grid(config=config)
    dates = pd.date_range("2021-01-04", periods=5, freq="B")
    models = config.m4.m3.prediction.all_model_names
    portfolios = tuple(
        name for name, _ in config.m4.portfolio.portfolio_fractions
    )
    predictions = pd.DataFrame(
        [
            {
                "specification": "cross_market_graph",
                "feature_set": "pvclcl_excess",
                "direction": "US_TO_CN",
                "target_market": "CN",
                "target_date": date,
                "target_code": code,
                "model": model,
                "prediction": (
                    (index + 1)
                    * (1 if code != "C2" else -1)
                    * (1 + model_index / 100)
                    * 0.001
                ),
                "actual_return": (
                    (index + 1)
                    * (1 if code != "C3" else -1)
                    * 0.0008
                ),
                "historical_liquidity": 10_000_000.0 + index * 100_000.0,
                "source_date": date - pd.offsets.BDay(1),
                "signal_date": date - pd.offsets.BDay(1),
                "execution_date": date,
            }
            for index, date in enumerate(dates)
            for code in ("C1", "C2", "C3", "C4")
            for model_index, model in enumerate(models)
        ]
    )
    portfolio = run_sensitivity_portfolios(
        predictions,
        group_columns=(
            "specification",
            "feature_set",
            "direction",
            "target_market",
            "model",
        ),
        config=config,
    )

    specification_frames: dict[str, pd.DataFrame] = {}
    for specification in config.m4.baseline.specifications:
        frame = predictions.copy()
        frame["specification"] = specification
        frame["feature_set"] = "opcl_excess"
        frame["direction"] = config.m4.baseline.expected_direction(
            specification,
            "CN",
        )
        specification_frames[specification] = frame
    real_cn_runner = run_cn_three_specification_portfolios

    def object_dtype_cn_runner(
        frames: Mapping[str, pd.DataFrame],
        *,
        config: M5Config,
    ) -> Any:
        result = real_cn_runner(frames, config=config)
        object_sharpe = result.portfolios.sharpe_metrics.copy()
        object_sharpe["annualized_sharpe"] = object_sharpe[
            "annualized_sharpe"
        ].map(
            lambda value: str(value) if pd.notna(value) else None
        ).astype(object)
        invalid_candidates = object_sharpe.index[
            object_sharpe["specification"].eq(
                config.m4.baseline.non_graph_specification
            )
            & object_sharpe["sharpe_status"].eq("ok")
        ]
        if invalid_candidates.empty:
            raise AssertionError("SA4 dtype 回归缺少可注入的有限非图 Sharpe")
        object_sharpe.loc[
            invalid_candidates[0],
            "annualized_sharpe",
        ] = "invalid-sharpe"
        result.portfolios.sharpe_metrics = object_sharpe
        return result

    smoke_m1 = M1Tables(
        us_panel=pd.DataFrame(),
        cn_panel=pd.DataFrame(),
        calendar_mapping=pd.DataFrame(),
        us_calendar=pd.DataFrame(),
        cn_calendar=pd.DataFrame(),
        us_market_returns=pd.DataFrame(),
        universe_audit=pd.DataFrame(),
        source_schema_audit=pd.DataFrame(),
        loaded_panel_audit=pd.DataFrame(),
        return_panel_audit=pd.DataFrame(),
        security_coverage_audit=pd.DataFrame(),
        return_distribution_audit=pd.DataFrame(),
    )
    sa4_store = CheckpointStore(smoke_dir / "sa4_dtype_regression")
    from unittest.mock import patch

    with (
        patch(
            f"{__name__}.iter_rolling_graph_snapshots",
            return_value=iter(()),
        ),
        patch(f"{__name__}._predict_snapshots", return_value=pd.DataFrame()),
        patch(
            f"{__name__}.prepare_portfolio_predictions",
            return_value=specification_frames[
                config.m4.baseline.cross_market_specification
            ],
        ),
        patch(
            f"{__name__}._collect_same_market_predictions",
            return_value=specification_frames[
                config.m4.baseline.same_market_graph_specification
            ],
        ),
        patch(
            f"{__name__}._collect_non_graph_predictions",
            return_value=specification_frames[
                config.m4.baseline.non_graph_specification
            ],
        ),
        patch(
            f"{__name__}.run_cn_three_specification_portfolios",
            side_effect=object_dtype_cn_runner,
        ),
    ):
        sa4_sharpe, sa4_delta = _run_sa4(
            smoke_m1,
            sa4_store,
            config=config,
        )
    expected_cn_sharpe_rows = (
        len(config.m4.baseline.specifications)
        * len(models)
        * len(portfolios)
    )
    record(
        "CA-A01",
        len(sa4_sharpe) == expected_cn_sharpe_rows
        and set(sa4_sharpe["specification"])
        == set(config.m4.baseline.specifications),
        (
            "真实 _run_sa4 经 CN-only 三规格低层审计后成功组合，"
            f"Sharpe 行数={len(sa4_sharpe)}"
        ),
    )
    baseline_for_expected = sa4_sharpe.loc[
        sa4_sharpe["specification"].eq(
            config.m4.baseline.non_graph_specification
        ),
        ["feature_set", "model", "portfolio", "annualized_sharpe", "sharpe_status"],
    ].rename(
        columns={
            "annualized_sharpe": "non_graph_sharpe",
            "sharpe_status": "non_graph_sharpe_status",
        }
    )
    expected_delta = sa4_sharpe.merge(
        baseline_for_expected,
        on=["feature_set", "model", "portfolio"],
        how="left",
        validate="many_to_one",
    )
    expected_current_numeric = pd.to_numeric(
        expected_delta["annualized_sharpe"],
        errors="coerce",
    )
    expected_baseline_numeric = pd.to_numeric(
        expected_delta["non_graph_sharpe"],
        errors="coerce",
    )
    expected_finite_pairs = (
        np.isfinite(expected_current_numeric)
        & np.isfinite(expected_baseline_numeric)
    )
    expected_complete_pairs = (
        expected_finite_pairs
        & expected_delta["sharpe_status"].eq("ok")
        & expected_delta["non_graph_sharpe_status"].eq("ok")
    )
    persisted_sa4_delta = sa4_store.read_frame("m5_sa4_delta")
    persisted_sa4_audit = sa4_store.read_frame(
        "m5_sa4_comparability_audit"
    )
    record(
        "SA4-DTYPE",
        str(sa4_delta["annualized_sharpe"].dtype) == "float64"
        and str(sa4_delta["non_graph_sharpe"].dtype) == "float64"
        and str(sa4_delta["sharpe_delta"].dtype) == "float64"
        and int(sa4_delta["finite_sharpe_pair"].sum())
        == int(expected_finite_pairs.sum())
        and int(sa4_delta["paired_complete"].sum())
        == int(expected_complete_pairs.sum())
        and int(sa4_delta["sharpe_delta"].notna().sum())
        == int(expected_complete_pairs.sum())
        and int(
            sa4_delta[
                "annualized_sharpe_conversion_failure_count"
            ].iloc[0]
        )
        == 1
        and int(
            sa4_delta[
                "non_graph_sharpe_conversion_failure_count"
            ].iloc[0]
        )
        == len(config.m4.baseline.specifications)
        and "ok"
        not in set(
            sa4_delta.loc[
                ~sa4_delta["paired_complete"],
                "delta_status",
            ]
        )
        and len(persisted_sa4_delta) == expected_cn_sharpe_rows
        and not persisted_sa4_audit.empty,
        (
            "真实 _run_sa4 消费 object Sharpe 且无 TypeError；"
            f"180 行 checkpoint 落盘，有限配对={int(expected_finite_pairs.sum())}，"
            f"完整配对={int(expected_complete_pairs.sum())}；"
            "无法转换值保留 NaN、失败计数和 delta_status"
        ),
    )

    industry = load_citic_level_one_industry(config=config)
    industry_summary = industry.summary_audit.iloc[0]
    record(
        "CA-A02",
        int(industry_summary["raw_row_count"]) == 6505
        and int(industry_summary["unique_stock_count"]) == 6501
        and int(industry_summary["duplicate_key_count"]) == 4
        and int(industry_summary["conflicting_duplicate_key_count"]) == 0
        and not industry.mapping.duplicated("target_code").any(),
        (
            "standard_code=37 实读 6505 行/6501 股票，"
            "4 个完全同值重复键，0 个冲突键，确定性去重后键唯一"
        ),
    )

    r4_pair = pd.DataFrame(
        [
            {
                "element_id": "R4",
                "target_market": "CN",
                "feature_set": "pvclcl_excess",
                "model": "OLS",
                "portfolio": "qr1",
                "same_graph_minus_non_graph": 0.15,
                "cross_minus_same_graph": -0.05,
            }
        ]
    )
    fig13 = build_figure13_long_table(r4_pair)
    record(
        "CA-A03",
        len(fig13) == 2
        and set(fig13["delta_type"])
        == {
            "same_graph_minus_non_graph",
            "cross_minus_same_graph",
        }
        and np.isfinite(fig13["sharpe_delta"]).all(),
        "R4 单行转为两类 Sharpe 差长表，Figure 13 有 2 条有限值",
    )

    incomplete_metrics = pd.DataFrame(
        [
            {
                "replacement_fraction": 0.2,
                "portfolio": "qr1",
                "model": model,
                "annualized_sharpe": float(index + 1),
                "sharpe_status": "ok",
            }
            for index, model in enumerate(models[:-1])
        ]
    )
    incomplete_median = summarize_model_median_sharpe(
        incomplete_metrics,
        sensitivity_columns=("replacement_fraction",),
        config=config,
    )
    incomplete_row = incomplete_median.audit.iloc[0]
    record(
        "CA-A04",
        pd.isna(incomplete_row["median_sharpe"])
        and incomplete_row["median_status"] == "incomplete_model_set"
        and models[-1] in incomplete_row["missing_models"],
        (
            "缺 1 个模型时 median_sharpe=N/A，"
            f"missing_models={incomplete_row['missing_models']}"
        ),
    )

    sa7_rows: list[dict[str, Any]] = []
    us_only_rows: list[dict[str, Any]] = []
    for combination_index, (combination, _, _) in enumerate(
        config.sensitivity.feature_combinations
    ):
        for model_index, model in enumerate(models):
            for portfolio_index, portfolio_name in enumerate(portfolios):
                base_sharpe = (
                    0.25
                    + combination_index / 10
                    + model_index / 100
                    + portfolio_index / 1000
                )
                delta = 0.02 if (model_index + portfolio_index) % 2 == 0 else -0.01
                common = {
                    "feature_combination": combination,
                    "model": model,
                    "portfolio": portfolio_name,
                    "sharpe_status": "ok",
                }
                us_only_rows.append(
                    {**common, "annualized_sharpe": base_sharpe}
                )
                sa7_rows.append(
                    {**common, "annualized_sharpe": base_sharpe + delta}
                )
    sa7_comparison = build_sa7_us_only_comparison(
        pd.DataFrame(sa7_rows),
        pd.DataFrame(us_only_rows),
        config=config,
    )
    expected_sa7_pairs = (
        len(config.sensitivity.feature_combinations)
        * len(models)
        * len(portfolios)
    )
    record(
        "CA-A05",
        len(sa7_comparison.paired) == expected_sa7_pairs
        and sa7_comparison.paired["paired_complete"].all()
        and sa7_comparison.summary["pairing_complete"].all(),
        (
            "四特征组合×十模型×六 qr 的 US-only 同键严格配对完整，"
            f"配对数={len(sa7_comparison.paired)}"
        ),
    )

    checkpoint_store = CheckpointStore(smoke_dir / "checkpoint_regression")
    checkpoint_output = checkpoint_store.write_frame(
        "probe_output",
        pd.DataFrame({"value": [1.0, 2.0]}),
    )
    checkpoint_context = checkpoint_store.build_context(
        "probe",
        config={"parameter": 1},
        input_manifest={"input_identity": "input-a"},
        code_hash_override="code-a",
    )
    checkpoint_store.mark_complete(
        "probe",
        {"purpose": "dynamic_regression"},
        context=checkpoint_context,
        output_paths=(checkpoint_output,),
    )
    unchanged_validation = checkpoint_store.validate_stage(
        "probe",
        checkpoint_context,
    )
    config_changed = checkpoint_store.validate_stage(
        "probe",
        checkpoint_store.build_context(
            "probe",
            config={"parameter": 2},
            input_manifest={"input_identity": "input-a"},
            code_hash_override="code-a",
        ),
    )
    input_changed = checkpoint_store.validate_stage(
        "probe",
        checkpoint_store.build_context(
            "probe",
            config={"parameter": 1},
            input_manifest={"input_identity": "input-b"},
            code_hash_override="code-a",
        ),
    )
    code_changed = checkpoint_store.validate_stage(
        "probe",
        checkpoint_store.build_context(
            "probe",
            config={"parameter": 1},
            input_manifest={"input_identity": "input-a"},
            code_hash_override="code-b",
        ),
    )
    single_path = smoke_dir / "checkpoint_regression" / "single.parquet"
    single_context = checkpoint_store.build_context(
        "sa3_single",
        config={"configuration_id": "smoke"},
        input_manifest={"upstream": checkpoint_store.marker_hash("probe")},
        code_hash_override="code-a",
    )
    checkpoint_store.write_external_frame_atomic(
        single_path,
        pd.DataFrame({"configuration_id": ["smoke"], "value": [1.0]}),
        context=single_context,
    )
    single_validation = checkpoint_store.validate_external_frame(
        single_path,
        expected_context=single_context,
    )
    record(
        "CA-A06",
        unchanged_validation.reusable
        and not config_changed.reusable
        and not input_changed.reusable
        and not code_changed.reusable
        and {
            config_changed.reason,
            input_changed.reason,
            code_changed.reason,
        }
        == {"context_hash_mismatch"}
        and single_validation.reusable
        and not tuple(single_path.parent.glob("*.part")),
        (
            "未变上下文可复用；config/input/code 任一变化均失效；"
            "SA3 单配置原子写入、schema/hash 复验通过且无 .part 残留"
        ),
    )

    daily = portfolio.daily_pnl.copy()
    cumulative_groups = [
        "specification",
        "feature_set",
        "direction",
        "target_market",
        "model",
        "portfolio",
    ]
    daily = daily.sort_values([*cumulative_groups, "target_date"])
    daily["cumulative_pnl"] = daily.groupby(
        cumulative_groups,
        observed=True,
    )["daily_pnl_before_cost"].cumsum()
    daily["series_label"] = daily[cumulative_groups].astype(str).agg(
        "/".join,
        axis=1,
    )
    fig10 = daily.loc[
        daily["specification"].eq("cross_market_graph")
        & daily["feature_set"].eq("pvclcl_excess")
        & daily["direction"].eq("US_TO_CN")
        & daily["target_market"].eq("CN")
    ].copy()
    fig10["series_label"] = fig10[["model", "portfolio"]].astype(str).agg(
        "/".join,
        axis=1,
    )
    expected_fig10_series = len(models) * len(portfolios)
    record(
        "CA-A08",
        fig10["series_label"].nunique() == expected_fig10_series
        and not fig10.duplicated(["series_label", "target_date"]).any()
        and daily["series_label"].nunique() == expected_fig10_series,
        (
            "Figure 10 严格筛至 cross/pvCLCL/US_TO_CN/CN，"
            f"保留 model×portfolio={expected_fig10_series} 条独立路径；"
            "模板累计图保留全部规格维度"
        ),
    )

    fig14 = pd.DataFrame(
        {
            "sensitivity_value": [0.0, 0.2, 1.0, 2.0],
            "median_sharpe": [0.8, 0.7, 0.65, 0.55],
            "series_label": ["边替换/qr1", "边替换/qr1", "信息滞后/qr1", "信息滞后/qr1"],
        }
    )
    directional = build_directional_sensitivity_comparison(
        {"SA1": randomized.audit, "SA7": sa7_comparison.paired}
    )
    smoke_audit = pd.DataFrame(regression_rows)
    report = write_formal_outputs(
        {
            "predictions": predictions,
            "daily_pnl": daily,
            "fig10_daily_pnl": fig10,
            "fig13": fig13,
            "fig14_edge_lag": fig14,
            "sharpe_metrics": portfolio.sharpe_metrics,
            "sa1_edge": randomized.audit,
            "sa1_edge_audit": randomized.audit,
            "sa2_lag": fig14.iloc[2:].assign(
                information_lag=lambda frame: frame["sensitivity_value"]
            ),
            "sa2_lag_audit": pd.DataFrame(
                {
                    "information_lag": [1, 2],
                    "history_only": [True, True],
                }
            ),
            "sa3_runs": grid,
            "sa4_opcl": sa4_sharpe,
            "sa4_opcl_delta": sa4_delta,
            "sa7_joint": pd.DataFrame(sa7_rows),
            "sa7_comparison": sa7_comparison.paired,
            "sensitivity_directional": directional,
            "core_comparison": pd.DataFrame(),
            "r2_comparison": pd.DataFrame(),
            "smoke_regression_audit": smoke_audit,
        },
        metrics={
            "report_name": "r007_cross_market_return_forecasting",
            "mode": "smoke",
            "formal_data_used": False,
            "hyperparameter_configuration_count": len(grid),
            "randomized_edge_count": len(randomized.edges),
        },
        comparison={
            "mode": "smoke",
            "verdict": None,
            "note": "仅验证接口可执行，不构成复现验证结论",
        },
        run_log_lines=(
            "模式：smoke",
            "数据：极小合成样本",
            f"随机种子：{config.sensitivity.seed}",
            "正式 output/results：未写入",
        ),
        output_dir=smoke_dir,
        config=config,
    )
    from openpyxl import load_workbook

    sensitivity_workbook = load_workbook(
        smoke_dir / "sensitivity_analysis.xlsx",
        read_only=False,
        data_only=False,
    )
    fig14_sheet = sensitivity_workbook["FIG14边替换与滞后"]
    sa4_delta_sheet = sensitivity_workbook["SA4差值"]
    fig16_generated = bool(
        report.chart_status.loc[
            report.chart_status["file"].eq(
                "fig16_opcl_sharpe_delta.png"
            ),
            "status",
        ].eq("generated").all()
    )
    record(
        "CA-A07",
        fig14_sheet.max_row > 1
        and len(fig14_sheet._images) == 1
        and sa4_delta_sheet.max_row == len(sa4_delta) + 1
        and len(sa4_delta_sheet._images) == 1
        and fig16_generated,
        (
            "sensitivity_analysis.xlsx 的 FIG14边替换与滞后、SA4差值"
            "均含非空数据且各嵌入 1 张对应图片；FIG16 状态 generated"
        ),
    )
    sensitivity_workbook.close()
    parsed_smoke = _parse_args(["--smoke"])
    parsed_formal = _parse_args([])
    record(
        "CLI",
        parsed_smoke.smoke and not parsed_formal.smoke,
        "`--smoke` 与无参正式入口分派互斥且可解析",
    )
    print(
        pd.DataFrame(regression_rows).to_string(index=False)
    )
    return report.files[-1]


def _parse_args(argv: Sequence[str] | None = None) -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="r007 跨市场收益预测复现流水线"
    )
    parser.add_argument(
        "--smoke",
        action="store_true",
        help="只运行 workspace 临时目录下的极小合成冒烟",
    )
    parser.add_argument(
        "--output-dir",
        type=Path,
        help="正式结果目录覆盖；smoke 模式忽略",
    )
    parser.add_argument(
        "--checkpoint-dir",
        type=Path,
        help="正式检查点目录覆盖；smoke 模式忽略",
    )
    return parser.parse_args(argv)


def main() -> int:
    args = _parse_args()
    config = DEFAULT_M5_CONFIG
    if args.output_dir is not None or args.checkpoint_dir is not None:
        reporting = replace(
            config.reporting,
            output_dir=args.output_dir or config.reporting.output_dir,
            checkpoint_dir=(
                args.checkpoint_dir or config.reporting.checkpoint_dir
            ),
        )
        config = replace(config, reporting=reporting)
    if args.smoke:
        artifact = run_smoke(config=config)
        print(f"冒烟产物清单：{artifact}")
    else:
        artifact = run_formal_pipeline(config=config)
        print(f"正式产物清单：{artifact}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
