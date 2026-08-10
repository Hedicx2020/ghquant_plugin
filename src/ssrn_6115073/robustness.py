"""m7 稳健性与敏感性分析 · ssrn_6115073（市场迁移复现 experimental）。

在 m1→m6 全链（``main.build_and_run``）产出的**同一基础链**上，用其组件在**不同子区间 /
参数网格 / 宏观因子子集**下参数化重跑，产出四张稳健性/敏感性表（对齐 SA1/SA2/SA3/SA4、B3/B4）：

  - **SA1 表3 滚动窗口稳健性（B3）**：A股测试段平移的 3 个 4 年滚动子区间（2019-2022/
    2020-2023/2021-2024，AS21/AS59），各算 AR/AStd/SR/PSR × 10 方法 × 3 市场，看 CPMM 系列
    是否**跨子区间稳健领先**。切基础链已实现收益序列的子窗（窗内重定基），策略 train-once
    不随子窗变，无泄漏。
  - **SA2 表5 样本外稳健性（B4）**：A股测试段尾窗 2023-2024（AS34/AS60，与后续 oos 阶段
    「2024 之后延伸」区分），6 内部设置 × AR/AStd/SR/PSR × 3 市场。
  - **SA3 表7 参数敏感性**：主市场 hs300 上 nc/w/λ 各 ±2 共 5 点（AS14 表头修正对称）：
    nc（CPMM-EW，复用 base fs_panel + 强制层次切割）、w（CPMM-EW，仅重跑 m4 stage2）、
    λ（CPMM-MV，复用 base 选股 + 重建 MV 权重，AS3 网格 {0.25,0.5,1,2,4}）。**直接回应
    前序脆弱性观察**（m3 nc 触顶 AS48、m4 w=7/9/9、m6 MV λ）。
  - **SA4 图15 单因子纳入/剔除敏感性**：主市场 hs300 CPMM-EW，逐宏观因子「仅用该因子」
    与「剔除该因子用其余六个」测 SR 变化（AS24 自算不采信矛盾描述，AS62），回应 R4
    （CPI/IR 哪个因子最关键）。

未来函数：各子区间/参数版本沿用主链纪律——训练期定 w/nc/γ̃（train-once，≤train_end）、
T+1 执行（``portfolio_backtest`` 硬保证）、stage2 历史窗口严格早于信号月；滚动子区间只切
**已实现收益序列**（策略参数不随子窗重估），无跨区间泄漏。

运行：``uv run python -m src.ssrn_6115073.robustness``（冒烟；不宣布 §九达标结论——归 verify）。
文件所有权：本文件复用 common/m2-m6 引擎（import），不改其它 milestone 模块逻辑。
"""

from __future__ import annotations

from pathlib import Path
from typing import Optional

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
import seaborn as sns
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

from common.allocation_backtest import portfolio_backtest
from src.ssrn_6115073.clustering import _hierarchical, relation_to_distance
from src.ssrn_6115073.config import (
    MARKET_NAMES_CN,
    PRIMARY_MARKET,
)
from src.ssrn_6115073.pattern_matching import (
    EP_SPAN_MONTHS,
    MACRO_WINDOW_MONTHS,
    WINDOW_GRID,
    signal_from_rhat,
    stage1_macro_period,
    stage2_asset_forecast,
)
from src.ssrn_6115073.portfolio import (
    ALL_MODELS,
    MV_LAMBDA,
    compute_metrics,
    selected_to_weights,
)
from src.ssrn_6115073.spillover import (
    FS_LONG_ONLY_FLOOR,
    aggregate_fs,
    select_per_cluster,
)

# 配色与样式（CLAUDE.md：蓝#1f77b4=主模型/正收益，红#d62728=对比/剔除，whitegrid，中文）
BLUE = "#1f77b4"
RED = "#d62728"
GREEN = "#2ca02c"
sns.set_style("whitegrid")
plt.rcParams["font.sans-serif"] = ["Songti SC", "Heiti TC", "STHeiti", "PingFang SC", "SimHei"]
plt.rcParams["axes.unicode_minus"] = False
plt.rcParams["figure.dpi"] = 300

RESULTS_DIR = Path("output/ssrn_6115073/results")

# ---------------------------------------------------------------------------
# m7 参数常量（集中于顶部；均可反查 spec/assumptions，杜绝魔法数字）
# ---------------------------------------------------------------------------
# SA1 滚动 4 年子区间（AS21/AS59）：原文 2016-2020/2017-2021/2018-2023（1 年滚动、末窗取满 4 年）
# → A股测试段（2019-2024）平移为等长 4 年窗、按 1 年滚动。窗边界用日期，切已实现收益序列。
ROLLING_SUBPERIODS: tuple[tuple[str, str, str], ...] = (
    ("2019-2022", "2019-01-01", "2022-12-31"),
    ("2020-2023", "2020-01-01", "2023-12-31"),
    ("2021-2024", "2021-01-01", "2024-12-31"),
)
# SA1：spec「ten approaches」——无 Index 列（6 内部消融 + 4 外部配置基准）
SA1_METHODS: tuple[str, ...] = (
    "CPMM-EW", "CPMM-MV", "CPM-EW", "CPM-MV", "PMM", "EWMFM", "HRP", "HERC", "OLMAR", "NCO",
)
# SA1/SA2 四比率（spec 表3/表5：AR/AStd/SR/PSR）
SA_RATIOS: tuple[str, ...] = ("AR", "AStd", "SR", "PSR")

# SA2 样本外尾窗（AS34/AS60）：原文 OOS=主测试之后的尾段（2023-2025）→ A股测试段尾部 2023-2024
# （与后续 oos 阶段「2024 之后延伸」明确区分，见 plan §八 m7 说明）
OOS_TAIL: tuple[str, str] = ("2023-01-01", "2024-12-31")
# SA2：spec「six settings」——仅 6 内部设置（不含外部基准）
SA2_SETTINGS: tuple[str, ...] = ("CPMM-EW", "CPMM-MV", "CPM-EW", "CPM-MV", "PMM", "EWMFM")

# SA3 参数敏感性（±2 共 5 点，AS14 表头 nc+1 修正后对称）
SA3_MARKET: str = PRIMARY_MARKET             # 算力控制：主市场 hs300（对齐原文 S&P500 主分析）
PARAM_OFFSETS: tuple[int, ...] = (-2, -1, 0, 1, 2)
# λ 网格（AS3）：λ−2..λ+2 = {0.25, 0.5, 1, 2, 4}，base λ=1=MV_LAMBDA
LAMBDA_GRID: tuple[float, ...] = (0.25, 0.5, 1.0, 2.0, 4.0)
# 表7 口径：Return=AR（年化算术收益）、Std=AStd（年化波动）、SR、PSR
SA3_RATIO_MAP: dict[str, str] = {"Return": "AR", "Std": "AStd", "SR": "SR", "PSR": "PSR"}

# SA4 单因子纳入/剔除（主市场 hs300 CPMM-EW；AS24 自算 / AS62 只用因子权重 γ̃≡1 隔离信号）
SA4_MARKET: str = PRIMARY_MARKET


# ---------------------------------------------------------------------------
# 通用小工具：子窗指标 / 强制 nc 簇标签 / FS 子集聚合 / 选股→回测
# ---------------------------------------------------------------------------

def _window_metrics(
    port_returns: pd.Series,
    benchmark_returns: pd.Series,
    rf: pd.Series,
    start: str,
    end: str,
) -> dict[str, float]:
    """在 [start,end] 日期子窗上重算 B7 指标（SA1/SA2 用；窗内 nav 重定基 nav=None）。

    切**已实现收益序列**（策略 train-once，参数不随子窗重估）——子区间稳健性=同一策略
    在不同评估窗的表现，无泄漏。空窗返回全 NaN。
    """
    if port_returns is None or port_returns.empty:
        return {k: float("nan") for k in (*SA_RATIOS, "SOR", "MR", "MDD", "VaR", "cum_return", "n_periods")}
    lo, hi = pd.Timestamp(start), pd.Timestamp(end)
    r = port_returns.loc[(port_returns.index >= lo) & (port_returns.index <= hi)].dropna()
    if r.empty:
        return {k: float("nan") for k in (*SA_RATIOS, "SOR", "MR", "MDD", "VaR", "cum_return", "n_periods")}
    return compute_metrics(r, benchmark_returns, rf)      # nav=None → 窗内重定基（MDD/累计口径正确）


def _forced_nc_labels(su_matrix: pd.DataFrame, nc: int) -> dict[str, int]:
    """强制 nc 簇标签：复刻 m3 canonical(hierarchical+SU) 路径但指定簇数 nc。

    D=1−SU（``relation_to_distance``）→ ward 层次切 nc 簇（``_hierarchical``）→ {股票→簇id}。
    nc=base 时与 m3 canonical cluster_labels 完全一致（冒烟校验）。
    """
    D = relation_to_distance(su_matrix, "su")
    labels = _hierarchical(D, int(nc))
    idx = list(su_matrix.index)
    return {str(idx[i]): int(labels[i]) for i in range(len(idx))}


def _select_from_fs(
    fs_panel: dict[pd.Timestamp, pd.Series],
    cluster_labels: dict[str, int],
    *,
    fs_floor: float = FS_LONG_ONLY_FLOOR,
) -> dict[pd.Timestamp, list[str]]:
    """逐测试月对给定 FS 面板 + 簇标签做每簇选一（F10/AS17，复用 m5 ``select_per_cluster``）。"""
    out: dict[pd.Timestamp, list[str]] = {}
    for t, fs in fs_panel.items():
        picks, _detail = select_per_cluster(fs, cluster_labels, fs_floor=fs_floor)
        out[t] = picks
    return out


def _fs_from_signals(
    signals: dict[pd.Timestamp, pd.DataFrame],
    gamma: pd.Series,
    factor_subset: list[str],
) -> dict[pd.Timestamp, pd.Series]:
    """按宏观因子子集重聚合 FS（式37，复用 m5 ``aggregate_fs``）。

    factor_subset=全 7 因子 → 复现 base FS；单因子/去一因子 → SA4；配 gamma 传全 1 可隔离信号。
    """
    out: dict[pd.Timestamp, pd.Series] = {}
    for t, sig in signals.items():
        cols = [f for f in factor_subset if f in sig.columns]
        if not cols:
            out[t] = pd.Series(np.nan, index=sig.index, name="FS")
            continue
        out[t] = aggregate_fs(sig[cols], gamma.reindex(cols))
    return out


def _backtest_selected(
    selected: dict[pd.Timestamp, list[str]],
    md,
    rf: pd.Series,
    *,
    build: str = "ew",
    lam: float = MV_LAMBDA,
) -> dict[str, float]:
    """选股面板 → EW/MV 建仓 → 序贯回测 → B7 指标（复用 m6 ``selected_to_weights`` + 引擎）。

    与 ``portfolio.run_market_portfolio`` 内部路径**同构且同口径**（无摩擦 cost_bps=0，建仓前
    ``member∩tradeable`` 时点过滤——**显式传 ``member_mask=md.member_mask``** 使 SA 表与 R1 主
    结果同为 member∩tradeable 口径，B1/CA-CDX-C-01；MV 协方差 ≤T），故 base 参数下逐值复现
    member 过滤后的 m6 新基准（冒烟校验）。
    """
    wp = selected_to_weights(selected, md.stock_returns, md.tradeable_mask, build,
                             member_mask=md.member_mask, lam=lam)
    if wp.empty:
        return {k: float("nan") for k in (*SA_RATIOS, "SOR", "MR", "MDD", "VaR", "cum_return", "n_periods")}
    bt = portfolio_backtest(wp, md.stock_returns, cost_bps=0.0)
    pr = bt["port_return"] if "port_return" in bt.columns else pd.Series(dtype=float)
    nav = bt["nav"] if "nav" in bt.columns else None
    return compute_metrics(pr, md.index_returns, rf, nav=nav)


def _signals_forced_w(
    md,
    level_panel: pd.DataFrame,
    candidates: list[str],
    factors: list[str],
    month_ends: pd.DatetimeIndex,
    test_months: list[pd.Timestamp],
    w: int,
    *,
    stage1_cache: Optional[dict[pd.Timestamp, dict[str, Optional[int]]]] = None,
) -> dict[pd.Timestamp, pd.DataFrame]:
    """参数化重跑 m4 两阶段匹配、仅注入强制窗口 w（stage1 与 w 无关，跨 w 缓存复用）。

    完全复刻 ``pattern_matching.run_market_pattern_matching`` 的内层循环（时间轴 reindex 到
    month_ends、整数位置索引、stage1 逐因子相似经济周期缓存复用给所有资产、stage2 强制窗口 w
    预测→式27 二值化），只把 w 从 ``search_optimal_window`` 定值改为外部指定——w=base 时逐格
    复现 base signals（冒烟校验）。未来函数由 stage2 内部整数位置约束（历史窗口严格早于信号月）保证。
    """
    R = md.stock_returns.reindex(month_ends)
    L = level_panel.reindex(month_ends)
    facs = [f for f in factors if f in L.columns]
    cands = [c for c in candidates if c in R.columns]
    ret_arrs = {c: R[c].to_numpy(dtype=float) for c in cands}
    level_arrs = {j: L[j].to_numpy(dtype=float) for j in facs}
    pos_of = {ts: i for i, ts in enumerate(month_ends)}

    signals: dict[pd.Timestamp, pd.DataFrame] = {}
    for t in test_months:
        pos_t = pos_of[t]
        if stage1_cache is not None and t in stage1_cache:
            ep_end = stage1_cache[t]
        else:
            ep_end = {j: stage1_macro_period(level_arrs[j], pos_t, MACRO_WINDOW_MONTHS) for j in facs}
            if stage1_cache is not None:
                stage1_cache[t] = ep_end
        rhat = pd.DataFrame(np.nan, index=cands, columns=facs, dtype=float)
        for j in facs:
            h = ep_end[j]
            if h is None:
                continue
            for c in cands:
                rhat.at[c, j] = stage2_asset_forecast(ret_arrs[c], pos_t, w, h, ep_span=EP_SPAN_MONTHS)
        signals[t] = rhat.map(signal_from_rhat)
    return signals


# ---------------------------------------------------------------------------
# SA1 表3：滚动 4 年子区间稳健性（B3；3 子区间 × 4 比率 × 10 方法 × 3 市场）
# ---------------------------------------------------------------------------

def build_sa1(prep, ports) -> tuple[pd.DataFrame, pd.DataFrame]:
    """SA1：各滚动子区间上重算 10 方法 4 比率（长表）+ CPMM-EW 跨区间领先性摘要。

    Returns:
        (sa1_long, leadership)：sa1_long 列=[子区间, 指标, 市场, 各方法]；leadership 列=
        [子区间, 市场, CPMM-EW_SR, CPMM-EW_排名/10, 最优方法, 最优SR]。
    """
    rf = prep.rf
    rows: list[dict] = []
    lead_rows: list[dict] = []
    for label, start, end in ROLLING_SUBPERIODS:
        # 逐市场逐方法算子窗指标
        win_metrics: dict[str, dict[str, dict[str, float]]] = {}   # market → method → metrics
        for market, mp in ports.items():
            bench = prep.markets[market].index_returns
            win_metrics[market] = {}
            for method in SA1_METHODS:
                mr = mp.models.get(method)
                pr = mr.port_returns if mr is not None else pd.Series(dtype=float)
                win_metrics[market][method] = _window_metrics(pr, bench, rf, start, end)
            # 领先性摘要（CPMM-EW 的 SR 在 10 方法中的排名）
            srs = {m: win_metrics[market][m].get("SR", np.nan) for m in SA1_METHODS}
            order = sorted(srs, key=lambda k: np.nan_to_num(srs[k], nan=-1e9), reverse=True)
            rank = order.index("CPMM-EW") + 1
            lead_rows.append({
                "子区间": label, "市场": MARKET_NAMES_CN.get(market, market),
                "CPMM-EW_SR": round(srs["CPMM-EW"], 4),
                "CPMM-EW_排名": f"{rank}/{len(SA1_METHODS)}",
                "最优方法": order[0], "最优SR": round(srs[order[0]], 4),
            })
        # 长表（指标 × 市场 × 方法）
        for ratio in SA_RATIOS:
            for market, mp in ports.items():
                rec: dict = {"子区间": label, "指标": ratio,
                             "市场": MARKET_NAMES_CN.get(market, market)}
                for method in SA1_METHODS:
                    rec[method] = round(win_metrics[market][method].get(ratio, np.nan), 4)
                rows.append(rec)
    sa1_long = pd.DataFrame(rows, columns=["子区间", "指标", "市场", *SA1_METHODS])
    leadership = pd.DataFrame(lead_rows,
                              columns=["子区间", "市场", "CPMM-EW_SR", "CPMM-EW_排名",
                                       "最优方法", "最优SR"])
    return sa1_long, leadership


# ---------------------------------------------------------------------------
# SA2 表5：样本外尾窗稳健性（B4；6 内部设置 × 4 比率 × 3 市场）
# ---------------------------------------------------------------------------

def build_sa2(prep, ports) -> pd.DataFrame:
    """SA2：样本外尾窗（OOS_TAIL）上重算 6 内部设置 4 比率（长表：指标 × 市场 × 6 设置）。"""
    rf = prep.rf
    start, end = OOS_TAIL
    rows: list[dict] = []
    for ratio in SA_RATIOS:
        for market, mp in ports.items():
            bench = prep.markets[market].index_returns
            rec: dict = {"指标": ratio, "市场": MARKET_NAMES_CN.get(market, market)}
            for method in SA2_SETTINGS:
                mr = mp.models.get(method)
                pr = mr.port_returns if mr is not None else pd.Series(dtype=float)
                rec[method] = round(_window_metrics(pr, bench, rf, start, end).get(ratio, np.nan), 4)
            rows.append(rec)
    return pd.DataFrame(rows, columns=["指标", "市场", *SA2_SETTINGS])


# ---------------------------------------------------------------------------
# SA3 表7：参数敏感性 nc/w/λ（主市场 hs300；每参数 ±2 共 5 点）
# ---------------------------------------------------------------------------

def build_sa3(prep, sel, clus, sig, sp) -> pd.DataFrame:
    """SA3：hs300 上 CPMM-EW 对 nc/w 敏感性 + CPMM-MV 对 λ 敏感性（Return/Std/SR/PSR）。

    - nc（CPMM-EW）：复用 base ``sp.fs_panel``（FS 与 nc 无关）+ 强制层次切 nc 簇 → 每簇选一 → EW。
    - w （CPMM-EW）：仅重跑 m4 stage2 定信号（stage1 缓存复用）→ base γ̃ 聚合 FS → base 簇标签
      每簇选一 → EW。
    - λ （CPMM-MV）：复用 base CPMM 选股（``sp.selected``）+ 重建 MV 权重（lam=λ）→ 回测。

    列：[模型, 参数, 列1..列5(含基准标注), 指标]，每参数组 4 行（Return/Std/SR/PSR）。
    """
    m = SA3_MARKET
    md = prep.markets[m]
    rf = prep.rf
    base_nc = int(clus[m].nc)
    base_w = int(sig[m].optimal_w)
    su = clus[m].su_matrix
    fs_panel = sp[m].fs_panel
    base_labels = dict(clus[m].cluster_labels)
    gamma = pd.Series(sp[m].gamma_weight, dtype=float)
    factors = list(sig[m].factors)
    base_selected = {t: list(v) for t, v in sp[m].selected.items()}

    n_cand = len(sel[m].candidates)
    w_lo, w_hi = WINDOW_GRID[0], WINDOW_GRID[-1]

    # nc 网格（clip 到 [2, n_candidates]）、w 网格（clip 到 WINDOW_GRID 边界）
    nc_grid = [max(2, min(n_cand, base_nc + o)) for o in PARAM_OFFSETS]
    w_grid = [max(w_lo, min(w_hi, base_w + o)) for o in PARAM_OFFSETS]

    # --- nc 敏感性（CPMM-EW）---
    nc_metrics: list[dict[str, float]] = []
    for nc in nc_grid:
        labels = _forced_nc_labels(su, nc)
        selected = _select_from_fs(fs_panel, labels)
        nc_metrics.append(_backtest_selected(selected, md, rf, build="ew"))

    # --- w 敏感性（CPMM-EW）---
    stage1_cache: dict = {}
    w_metrics: list[dict[str, float]] = []
    for w in w_grid:
        signals_w = _signals_forced_w(
            md, prep.macro.level_panel, list(sel[m].candidates), factors,
            prep.month_ends, list(sig[m].test_months), w, stage1_cache=stage1_cache)
        fs_w = _fs_from_signals(signals_w, gamma, factors)      # 全 7 因子 base γ̃
        selected = _select_from_fs(fs_w, base_labels)           # base 簇标签（nc 不变）
        w_metrics.append(_backtest_selected(selected, md, rf, build="ew"))

    # --- λ 敏感性（CPMM-MV）---
    lam_metrics: list[dict[str, float]] = []
    for lam in LAMBDA_GRID:
        lam_metrics.append(_backtest_selected(base_selected, md, rf, build="mv", lam=lam))

    # 组表（每参数组 4 指标行；列头标注参数取值 + 基准列）
    def _col_headers(grid: list, base_val, fmt: str) -> list[str]:
        return [f"{fmt.format(v)}{'(基准)' if v == base_val else ''}" for v in grid]

    nc_heads = _col_headers(nc_grid, base_nc, "nc={:d}")
    w_heads = _col_headers(w_grid, base_w, "w={:d}")
    lam_heads = _col_headers(list(LAMBDA_GRID), MV_LAMBDA, "λ={:g}")

    rows: list[dict] = []
    for disp, metric_key in SA3_RATIO_MAP.items():
        rows.append({"模型": "CPMM-EW", "参数": "nc", "指标": disp,
                     **{h: round(mm.get(metric_key, np.nan), 4) for h, mm in zip(nc_heads, nc_metrics)}})
    for disp, metric_key in SA3_RATIO_MAP.items():
        rows.append({"模型": "CPMM-EW", "参数": "w", "指标": disp,
                     **{h: round(mm.get(metric_key, np.nan), 4) for h, mm in zip(w_heads, w_metrics)}})
    for disp, metric_key in SA3_RATIO_MAP.items():
        rows.append({"模型": "CPMM-MV", "参数": "λ", "指标": disp,
                     **{h: round(mm.get(metric_key, np.nan), 4) for h, mm in zip(lam_heads, lam_metrics)}})

    # 三参数各自 5 列头不同 → 统一列名为「列1..列5」并附一行取值说明，便于单表呈现
    generic_cols = [f"列{i+1}" for i in range(len(PARAM_OFFSETS))]
    out_rows: list[dict] = []
    header_map = {"nc": nc_heads, "w": w_heads, "λ": lam_heads}
    for r in rows:
        heads = header_map[r["参数"]]
        rec = {"模型": r["模型"], "参数": r["参数"], "指标": r["指标"]}
        for gc, h in zip(generic_cols, heads):
            rec[gc] = r[h]
        out_rows.append(rec)
    # 取值说明行（参数 → 5 列实际取值）
    legend = {"nc": nc_heads, "w": w_heads, "λ": lam_heads}
    sa3 = pd.DataFrame(out_rows, columns=["模型", "参数", "指标", *generic_cols])
    sa3.attrs["legend"] = legend
    sa3.attrs["base"] = {"nc": base_nc, "w": base_w, "λ": MV_LAMBDA}
    return sa3


# ---------------------------------------------------------------------------
# SA4 图15：单因子纳入/剔除敏感性（主市场 hs300 CPMM-EW）
# ---------------------------------------------------------------------------

def build_sa4(prep, sig, sp, clus) -> pd.DataFrame:
    """SA4：hs300 CPMM-EW 逐宏观因子「仅用该因子」/「剔除该因子」的 SR（及 AR/累计）变化。

    - 仅用 f（γ̃≡1 隔离信号，AS62）：FS_j=signal_{f,j} → 每簇选一 → EW；避免 min-max γ̃=0
      （如 GDP）致单因子退化空仓，忠实测该因子**信号**的选股价值。
    - 剔除 f（base γ̃，其余 6 因子）：FS_j=Σ_{i≠f} γ̃_i·signal_{i,j} → 每簇选一 → EW，即
      CPMM-EW 去掉因子 f。
    - baseline（全 7 因子 base γ̃）= base CPMM-EW（冒烟校验）。

    列：[宏观因子, 仅用_SR, 剔除_SR, baseline_SR, 剔除后ΔSR, 仅用_AR, 剔除_AR]。
    """
    m = SA4_MARKET
    md = prep.markets[m]
    rf = prep.rf
    signals = sig[m].signals
    gamma = pd.Series(sp[m].gamma_weight, dtype=float)
    labels = dict(clus[m].cluster_labels)
    factors = list(sig[m].factors)
    ones = pd.Series(1.0, index=factors)

    # baseline：全 7 因子 base γ̃（应复现 base CPMM-EW）
    fs_all = _fs_from_signals(signals, gamma, factors)
    base_metrics = _backtest_selected(_select_from_fs(fs_all, labels), md, rf, build="ew")
    base_sr = base_metrics.get("SR", np.nan)

    rows: list[dict] = []
    for f in factors:
        # 仅用 f（γ̃≡1）
        fs_only = _fs_from_signals(signals, ones, [f])
        only_m = _backtest_selected(_select_from_fs(fs_only, labels), md, rf, build="ew")
        # 剔除 f（base γ̃，其余 6）
        others = [x for x in factors if x != f]
        fs_excl = _fs_from_signals(signals, gamma, others)
        excl_m = _backtest_selected(_select_from_fs(fs_excl, labels), md, rf, build="ew")
        rows.append({
            "宏观因子": f,
            "仅用_SR": round(only_m.get("SR", np.nan), 4),
            "剔除_SR": round(excl_m.get("SR", np.nan), 4),
            "baseline_SR": round(base_sr, 4),
            "剔除后ΔSR": round(excl_m.get("SR", np.nan) - base_sr, 4),
            "仅用_AR": round(only_m.get("AR", np.nan), 4),
            "剔除_AR": round(excl_m.get("AR", np.nan), 4),
        })
    sa4 = pd.DataFrame(rows, columns=["宏观因子", "仅用_SR", "剔除_SR", "baseline_SR",
                                      "剔除后ΔSR", "仅用_AR", "剔除_AR"])
    sa4.attrs["base_sr"] = float(base_sr)
    return sa4


# ---------------------------------------------------------------------------
# Excel + 图表
# ---------------------------------------------------------------------------

def _format_ws(ws) -> None:
    """冻结首行 + 表头加粗 + 自动列宽（中文按 2 宽计）。"""
    ws.freeze_panes = "A2"
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center")
    for col in ws.columns:
        letter = get_column_letter(col[0].column)
        width = 0
        for cell in col:
            v = "" if cell.value is None else str(cell.value)
            w = sum(2 if ord(ch) > 127 else 1 for ch in v)
            width = max(width, w)
        ws.column_dimensions[letter].width = min(width + 3, 42)


def _mode_note() -> pd.DataFrame:
    return pd.DataFrame({
        "项目": ["复现模式", "本表性质", "SA1 子区间", "SA2 样本外", "SA3 参数网格",
                 "SA4 单因子", "数值对齐", "达标判定"],
        "说明": [
            "experimental（市场迁移复现）——SA 类为定性稳健性验证，不追平原文数值",
            "coder 冒烟端到端产出；方法是否跨子区间/参数稳健的判定归 verify/oos，本表不含结论",
            "A股测试段平移 3 个 4 年滚动窗 2019-2022/2020-2023/2021-2024（AS21/AS59），切已实现收益重定基",
            "A股测试段尾窗 2023-2024（AS34/AS60，与后续 oos 阶段『2024 之后延伸』区分）",
            "主市场 hs300：nc/w/λ 各 ±2 共 5 点（AS14 表头修正对称；λ={0.25,0.5,1,2,4}/AS3）",
            "hs300 CPMM-EW：仅用/剔除单因子（AS24 自算，仅用因子权重 γ̃≡1 隔离信号/AS62）",
            "豁免——实验模式：迁移市场定性验证",
            "§九 达标结论归 verify/diagnoser，本表只报数",
        ],
    })


def _sa3_legend_frame(sa3: pd.DataFrame) -> pd.DataFrame:
    """SA3 通用列名 → 各参数实际取值对照（列1..列5 分别是什么）。"""
    legend = sa3.attrs.get("legend", {})
    generic = [f"列{i+1}" for i in range(len(PARAM_OFFSETS))]
    rows = [{"参数": p, **{gc: h for gc, h in zip(generic, heads)}} for p, heads in legend.items()]
    return pd.DataFrame(rows, columns=["参数", *generic])


def save_excel(prep, ports, sel, clus, sig, sp, path: Path) -> dict:
    """写 robustness_summary.xlsx：说明 / SA1 / SA1 领先性 / SA2 / SA3 / SA3 列取值 / SA4。返回各表 DataFrame。"""
    path.parent.mkdir(parents=True, exist_ok=True)
    note = _mode_note()
    sa1_long, sa1_lead = build_sa1(prep, ports)
    sa2 = build_sa2(prep, ports)
    sa3 = build_sa3(prep, sel, clus, sig, sp)
    sa3_legend = _sa3_legend_frame(sa3)
    sa4 = build_sa4(prep, sig, sp, clus)

    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        note.to_excel(writer, sheet_name="说明", index=False)
        sa1_long.to_excel(writer, sheet_name="SA1_滚动子区间", index=False)
        sa1_lead.to_excel(writer, sheet_name="SA1_CPMM领先性", index=False)
        sa2.to_excel(writer, sheet_name="SA2_样本外尾窗", index=False)
        sa3.to_excel(writer, sheet_name="SA3_参数敏感性", index=False)
        sa3_legend.to_excel(writer, sheet_name="SA3_列取值对照", index=False)
        sa4.to_excel(writer, sheet_name="SA4_单因子敏感性", index=False)
        for name in ("说明", "SA1_滚动子区间", "SA1_CPMM领先性", "SA2_样本外尾窗",
                     "SA3_参数敏感性", "SA3_列取值对照", "SA4_单因子敏感性"):
            _format_ws(writer.sheets[name])
    print(f"[Excel] 已写 {path}")
    return {"sa1_long": sa1_long, "sa1_lead": sa1_lead, "sa2": sa2, "sa3": sa3, "sa4": sa4}


def _fig_sa1_leadership(sa1_lead: pd.DataFrame, path: Path) -> None:
    """SA1：三市场 CPMM-EW 在各滚动子区间的 SR（分组条形，看跨区间稳健领先）。"""
    markets = list(dict.fromkeys(sa1_lead["市场"]))
    labels = [p[0] for p in ROLLING_SUBPERIODS]
    fig, ax = plt.subplots(figsize=(12, 6))
    x = np.arange(len(labels))
    width = 0.25
    palette = [BLUE, RED, GREEN]
    for i, mk in enumerate(markets):
        sub = sa1_lead[sa1_lead["市场"] == mk].set_index("子区间")
        vals = [sub.loc[lb, "CPMM-EW_SR"] if lb in sub.index else np.nan for lb in labels]
        ax.bar(x + (i - 1) * width, vals, width, label=mk, color=palette[i % len(palette)], alpha=0.85)
    ax.set_title("SA1 CPMM-EW 跨滚动子区间 SR（4 年窗，3 市场）", fontsize=14)
    ax.set_xlabel("滚动子区间", fontsize=12)
    ax.set_ylabel("夏普比率 SR", fontsize=12)
    ax.set_xticks(x)
    ax.set_xticklabels(labels)
    ax.axhline(0, color="gray", lw=0.8, ls="--")
    ax.legend(fontsize=10)
    fig.tight_layout()
    fig.savefig(path, dpi=300, bbox_inches="tight")
    plt.close(fig)


def _fig_sa3_sensitivity(sa3: pd.DataFrame, path: Path) -> None:
    """SA3：nc/w/λ 三面板 SR 随参数变化（基准点标注），看方法稳健还是脆弱。"""
    legend = sa3.attrs.get("legend", {})
    base = sa3.attrs.get("base", {})
    generic = [f"列{i+1}" for i in range(len(PARAM_OFFSETS))]
    fig, axes = plt.subplots(1, 3, figsize=(15, 5))
    for ax, param, model in zip(axes, ("nc", "w", "λ"), ("CPMM-EW", "CPMM-EW", "CPMM-MV")):
        sr_row = sa3[(sa3["参数"] == param) & (sa3["指标"] == "SR")]
        if sr_row.empty:
            continue
        sr = [float(sr_row.iloc[0][gc]) for gc in generic]
        heads = legend.get(param, generic)
        xs = list(range(len(sr)))
        ax.plot(xs, sr, marker="o", color=BLUE, lw=2.0)
        # 基准点红色突出
        for xi, h in enumerate(heads):
            if "基准" in h:
                ax.plot(xi, sr[xi], marker="o", color=RED, markersize=11, zorder=5)
        ax.set_title(f"{param} 敏感性（{model}）", fontsize=13)
        ax.set_xlabel(f"{param} 取值", fontsize=11)
        ax.set_ylabel("SR", fontsize=11)
        ax.set_xticks(xs)
        ax.set_xticklabels([h.replace("(基准)", "\n(基准)") for h in heads], fontsize=9)
    fig.suptitle("SA3 参数敏感性（hs300；nc/w→CPMM-EW，λ→CPMM-MV；红=基准）", fontsize=14)
    fig.tight_layout()
    fig.savefig(path, dpi=300, bbox_inches="tight")
    plt.close(fig)


def _fig_sa4_factor(sa4: pd.DataFrame, path: Path) -> None:
    """SA4/图15：逐宏观因子『仅用』(蓝) 与『剔除』(红) 的 SR + baseline 参考线。"""
    base_sr = sa4.attrs.get("base_sr", np.nan)
    factors = list(sa4["宏观因子"])
    x = np.arange(len(factors))
    width = 0.38
    fig, ax = plt.subplots(figsize=(12, 6))
    ax.bar(x - width / 2, sa4["仅用_SR"].to_numpy(dtype=float), width,
           label="仅用该因子", color=BLUE, alpha=0.85)
    ax.bar(x + width / 2, sa4["剔除_SR"].to_numpy(dtype=float), width,
           label="剔除该因子(用其余6)", color=RED, alpha=0.85)
    if np.isfinite(base_sr):
        ax.axhline(base_sr, color="gray", lw=1.6, ls="--", label=f"baseline 全7因子 SR={base_sr:.3f}")
    ax.set_title("SA4 单因子纳入/剔除对 CPMM-EW 的 SR 影响（hs300，图15）", fontsize=14)
    ax.set_xlabel("宏观因子", fontsize=12)
    ax.set_ylabel("夏普比率 SR", fontsize=12)
    ax.set_xticks(x)
    ax.set_xticklabels(factors)
    ax.axhline(0, color="gray", lw=0.8)
    ax.legend(fontsize=10)
    fig.tight_layout()
    fig.savefig(path, dpi=300, bbox_inches="tight")
    plt.close(fig)


def save_charts(tables: dict) -> None:
    """产 SA1 领先性 / SA3 参数敏感性 / SA4 单因子（图15）三张 PNG。"""
    RESULTS_DIR.mkdir(parents=True, exist_ok=True)
    _fig_sa1_leadership(tables["sa1_lead"], RESULTS_DIR / "sa1_rolling_leadership.png")
    _fig_sa3_sensitivity(tables["sa3"], RESULTS_DIR / "sa3_parameter_sensitivity.png")
    _fig_sa4_factor(tables["sa4"], RESULTS_DIR / "sa4_factor_sensitivity.png")
    print(f"[图表] 已写 3 张 PNG 到 {RESULTS_DIR}")


# ---------------------------------------------------------------------------
# 主入口 + 冒烟（只报数，不宣布 §九达标结论——归 verify/oos）
# ---------------------------------------------------------------------------

def run_robustness(prep, sel, clus, sig, sp, ports) -> dict:
    """m7 主入口：在 base 链产物上产 SA1/SA2/SA3/SA4 表 + 图，返回各表 DataFrame。"""
    tables = save_excel(prep, ports, sel, clus, sig, sp,
                        RESULTS_DIR / "robustness_summary.xlsx")
    save_charts(tables)
    return tables


def _smoke() -> None:
    from src.ssrn_6115073.main import build_and_run

    print("=" * 88)
    print("m7 稳健性与敏感性分析冒烟 · ssrn_6115073（市场迁移；只报数不下 §九结论）")
    print("=" * 88)
    prep, sel, clus, sig, sp, ports = build_and_run()

    # base 复现校验（robustness 重跑路径应与 m6 基准逐值一致）
    m = SA3_MARKET
    md, rf = prep.markets[m], prep.rf
    base_ew = ports[m].models["CPMM-EW"].metrics.get("SR", np.nan)
    base_mv = ports[m].models["CPMM-MV"].metrics.get("SR", np.nan)
    # nc=base 复现
    labels = _forced_nc_labels(clus[m].su_matrix, clus[m].nc)
    rep_nc = _backtest_selected(_select_from_fs(sp[m].fs_panel, labels), md, rf, build="ew").get("SR", np.nan)
    # w=base 复现
    sig_w = _signals_forced_w(md, prep.macro.level_panel, list(sel[m].candidates),
                              list(sig[m].factors), prep.month_ends, list(sig[m].test_months),
                              int(sig[m].optimal_w))
    fs_w = _fs_from_signals(sig_w, pd.Series(sp[m].gamma_weight, dtype=float), list(sig[m].factors))
    rep_w = _backtest_selected(_select_from_fs(fs_w, dict(clus[m].cluster_labels)), md, rf, build="ew").get("SR", np.nan)
    # λ=base 复现
    base_selected = {t: list(v) for t, v in sp[m].selected.items()}
    rep_lam = _backtest_selected(base_selected, md, rf, build="mv", lam=MV_LAMBDA).get("SR", np.nan)
    print(f"\n[base 复现校验 {m}] m6 基准 vs robustness 重跑 SR（应≈一致）：")
    print(f"  CPMM-EW: 基准 {base_ew:.4f} | nc=base 重跑 {rep_nc:.4f}(Δ{rep_nc-base_ew:+.4f}) | "
          f"w=base 重跑 {rep_w:.4f}(Δ{rep_w-base_ew:+.4f})")
    print(f"  CPMM-MV: 基准 {base_mv:.4f} | λ=base 重跑 {rep_lam:.4f}(Δ{rep_lam-base_mv:+.4f})")

    tables = run_robustness(prep, sel, clus, sig, sp, ports)

    # SA1 领先性
    print("\n[SA1] CPMM-EW 跨滚动子区间 SR 排名（10 方法）：")
    print(tables["sa1_lead"].to_string(index=False))
    # SA2 SR 行
    print("\n[SA2] 样本外尾窗 SR（6 内部设置 × 3 市场）：")
    print(tables["sa2"][tables["sa2"]["指标"] == "SR"].to_string(index=False))
    # SA3 SR 行 + 列取值
    print("\n[SA3] 参数敏感性 SR 行（列取值见对照）：")
    print(tables["sa3"][tables["sa3"]["指标"] == "SR"].to_string(index=False))
    print("  列取值对照：")
    print(_sa3_legend_frame(tables["sa3"]).to_string(index=False))
    # SA4
    print("\n[SA4] 单因子纳入/剔除对 CPMM-EW SR 影响（hs300）：")
    print(tables["sa4"].to_string(index=False))
    print("\n[完成] m7 稳健性产物已落盘（SA1/SA2/SA3/SA4 表 + 3 图）。§九达标判定归 verify/oos。")


if __name__ == "__main__":
    _smoke()
